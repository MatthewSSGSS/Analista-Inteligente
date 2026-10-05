"""Plan de acción territorial (`core/territorio_plan.plan`).

Qué se protege:

- La estrategia de cada zona sale de reglas que un gerente puede explicar:
  cae y pesa (80% del resultado) → Rescatar; cae y pesa poco → Recuperar;
  vende menos de la mitad de lo normal para su población → Desarrollar;
  crece → Replicar; lo demás → Sostener.
- Lo que vale cada acción es AL MES y no se inventa: recuperar = volver al
  mes anterior completo; si el mes actual va a medias, el ritmo sale de la
  variación comparable al mismo día (un mes cargado al día 7 no se cuenta
  como una caída del 77%). Desarrollar y abrir = la mitad del camino a la
  penetración típica de la propia red.
- Las jugadas se ordenan por valor esperado (abrir es una apuesta: pesa un
  cuarto), y con tope de aperturas: antes «Abrir Medellín» tapaba todos los
  rescates solo por tamaño.
- El plan dice a QUIÉN visitar: el cliente que dejó de comprar en la zona.
- El plan por agente: perfil, posición y mediana sobre TODOS los agentes
  (regla del dominio), dónde perdió y su ruta solo por zonas donde es de los
  dos que más venden.
- Dónde abrir: municipios de 20.000+ habitantes sin presencia, en los
  departamentos donde se opera, con la zona más cercana como base.

PYTHONPATH=. python tests/territorio_plan_test.py
"""
import logging
import warnings

import pandas as pd

from core import territorio as T
from core import territorio_plan as P
from core.profile import profile_sheet

warnings.filterwarnings("ignore")
logging.disable(logging.WARNING)


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _perfil(df0):
    item = profile_sheet(df0, {"sheet_name": "x", "workbook_name": "x.xlsx"})
    return item["processed"], item["profile"]["schema"]


# Ventas diarias por municipio de Atlántico, proporcionales a la población
# (≈ 10 por cada 10.000 habitantes al día) salvo los casos a probar.
RED = {"BARRANQUILLA": 1_275_854, "SOLEDAD": 733_597, "MALAMBO": 153_223, "SABANALARGA": 106_209,
       "BARANOA": 76_917, "PUERTO COLOMBIA": 65_686}


def _red():
    filas = []
    for dia in pd.date_range("2026-07-01", "2026-09-07", freq="D"):
        sep = dia.month == 9
        for ciudad, pob in RED.items():
            diario = pob / 10_000 * 10
            if ciudad == "SOLEDAD":
                diario *= 0.2                       # rezagada: vende un 20% de lo normal
            if ciudad == "BARRANQUILLA":
                # Barranquilla cae a la mitad en septiembre: la Droguería Central deja de comprar.
                filas.append({"Fecha": dia, "Ciudad": ciudad, "Asesor": "Ana", "Cliente": "Droguería Central",
                              "Ventas": 0.0 if sep else diario / 2})
                filas.append({"Fecha": dia, "Ciudad": ciudad, "Asesor": "Ana", "Cliente": "Tienda Norte",
                              "Ventas": diario / 2})
                continue
            if ciudad == "PUERTO COLOMBIA" and sep:
                diario *= 1.5                       # crece
            asesor = "Beto" if ciudad in ("SOLEDAD", "MALAMBO") else "Carla"
            filas.append({"Fecha": dia, "Ciudad": ciudad, "Asesor": asesor, "Cliente": f"Cliente {ciudad.title()}",
                          "Ventas": diario})
    return pd.DataFrame(filas)


def _plan():
    df, schema = _perfil(_red())
    ub, _ = T.ubicar(df, schema)
    z = T.zonas(ub, "Ventas", "Suma", "municipio", "Fecha", None)
    agente = P.columna_agente(df, list(df.columns))
    cuenta = P.columna_cuenta(df, list(df.columns), excluir=(agente,))
    return df, ub, z, agente, cuenta, P.plan(ub, "Ventas", "Suma", "Fecha", z, ["Cliente"], {}, agente=agente, cuenta=cuenta)


def test_columnas():
    df, _, _, agente, cuenta, _ = _plan()
    check("detecta la columna de agente («Asesor») y la de cliente", (agente, cuenta) == ("Asesor", "Cliente"))
    check("sin columna de persona no inventa un agente", P.columna_agente(df, ["Ciudad", "Cliente"]) is None)


def test_estrategias_y_valores():
    _, _, z, _, _, pl = _plan()
    t = pl["zonas"].set_index("nombre")
    check("el mes actual va al día 7 y se compara al mismo día", pl["corte_dia"] == 7)
    check("Barranquilla cae y pesa → Rescatar", t.loc["Barranquilla", "estrategia"] == "rescatar")
    agosto = 1_275_854 / 10_000 * 10 * 31
    check("vale volver a agosto completo: la mitad de lo que vendía (la droguería que se perdió)",
          abs(t.loc["Barranquilla", "ref_mes"] - agosto) < 1 and abs(t.loc["Barranquilla", "valor_mes"] - agosto / 2) < 1)
    check("Soledad vende un 20% de lo normal → Desarrollar", t.loc["Soledad", "estrategia"] == "desarrollar")
    pen, tipica = t.loc["Soledad", "pen_mes"], pl["tipica"]
    check("desarrollar vale la mitad del camino a la penetración típica",
          abs(t.loc["Soledad", "valor_mes"] - (tipica - pen) * 733_597 / 10_000 * 0.5) < 1)
    check("Puerto Colombia crece → Replicar", t.loc["Puerto Colombia", "estrategia"] == "replicar")
    check("lo estable y sano → Sostener", t.loc["Baranoa", "estrategia"] == "sostener")
    check("un mes a medias no se lee como caída: Baranoa, igual todos los días, sigue a su ritmo de agosto",
          abs(t.loc["Baranoa", "ritmo"] - t.loc["Baranoa", "ref_mes"]) < 1)


def test_jugadas_y_a_quien_visitar():
    _, _, _, _, _, pl = _plan()
    j = pl["jugadas"]
    check("hay jugadas, numeradas desde 1", j and [x["n"] for x in j] == list(range(1, len(j) + 1)))
    check("la primera es rescatar Barranquilla, no abrir un municipio por tamaño",
          j[0]["estrategia"] == "rescatar" and j[0]["zona"] == "Barranquilla")
    check("como mucho 2 aperturas en la lista", sum(x["estrategia"] == "abrir" for x in j) <= P.TOPE_JUGADAS["abrir"])
    pasos = " ".join(j[0]["pasos"])
    check("dice a quién visitar: la Droguería Central dejó de comprar", "Droguería Central" in pasos)
    check("y con quién: Ana es la responsable", j[0]["responsable"] == "Ana" and "**Ana**" in pasos)
    check("el resumen dice cuánto hay en juego al mes", "al mes" in pl["resumen"][0] and "en juego" in pl["resumen"][0])


def test_agentes():
    _, _, _, _, _, pl = _plan()
    ag = {a["nombre"]: a for a in pl["agentes"]}
    check("un plan por cada agente, con posición sobre los 3", set(ag) == {"Ana", "Beto", "Carla"}
          and all(a["de"] == 3 for a in ag.values()))
    check("Ana cae: perfil «En caída», perdió en Barranquilla y perdió a la Droguería Central",
          ag["Ana"]["perfil"] == "caida" and ag["Ana"]["perdio"][0]["zona"] == "Barranquilla"
          and ag["Ana"]["cuentas_perdidas"][0]["nombre"] == "Droguería Central")
    check("su meta sugerida es su mejor nivel reciente (agosto completo)",
          abs(ag["Ana"]["meta"] - 1_275_854 / 10_000 * 10 * 31) < 1)
    check("la ruta de Beto pasa por Soledad (es su zona y hay que desarrollarla)",
          any(r["zona"] == "Soledad" for r in ag["Beto"]["ruta"]))
    check("y la de Ana no, porque Soledad no es su zona", not any(r["zona"] == "Soledad" for r in ag["Ana"]["ruta"]))


def test_donde_abrir():
    _, _, _, _, _, pl = _plan()
    ap = pl["aperturas"]
    presentes = {"Barranquilla", "Soledad", "Malambo", "Sabanalarga", "Baranoa", "Puerto Colombia"}
    check("propone municipios de Atlántico de 20.000+ habitantes sin presencia",
          len(ap) and set(ap["departamento"]) == {"Atlántico"} and (ap["poblacion"] >= 20_000).all()
          and not set(ap["municipio"]) & presentes)
    check("con la zona más cercana de la red como base y su distancia",
          set(ap["base"]) <= presentes and (ap["distancia_km"] > 0).all())
    fila = ap.iloc[0]
    check("valorado en la mitad de la penetración típica, al mes",
          abs(fila["potencial_mes"] - pl["tipica"] * fila["poblacion"] / 10_000 * 0.5) < 1)
    check("y el modelo de entrada depende de la distancia", fila["modelo"].startswith(("Ruta desde", "Agente", "Punto")))


def test_promedio_no_tiene_plan():
    df, schema = _perfil(_red())
    ub, _ = T.ubicar(df, schema)
    z = T.zonas(ub, "Ventas", "Promedio", "municipio", "Fecha", None)
    check("con un promedio no se arma un plan en plata", P.plan(ub, "Ventas", "Promedio", "Fecha", z) == {})


def test_informes_con_plan():
    import io
    from openpyxl import load_workbook
    from ui.report_territorial import build_territorial_excel, build_territorial_html
    _, ub, z, _, _, pl = _plan()
    inf = T.informe(ub, "Ventas", "Suma", "Fecha", z, metrica_label="Ventas")
    ctx = {"archivo": "x.xlsx", "hoja": "H", "registros": 1, "filtros": ""}
    libro = load_workbook(io.BytesIO(build_territorial_excel(inf, ctx, plan=pl)))
    check("el Excel trae «Plan de acción», «Agentes» y «Dónde abrir»",
          {"Plan de acción", "Agentes", "Dónde abrir"} <= set(libro.sheetnames))
    ws = libro["Plan de acción"]
    textos = [str(c.value) for fila in ws.iter_rows() for c in fila if c.value is not None]
    check("con estado «Pendiente» para el seguimiento y lista desplegable",
          "Pendiente" in textos and len(ws.data_validations.dataValidation) == 1)
    check("y sin marcas ** en el texto", not any("**" in s for s in textos))
    pagina = build_territorial_html(inf, ctx, plan=pl).decode("utf-8")
    check("el HTML trae el plan, las jugadas y el plan por agente",
          "🎯 Plan de acción" in pagina and "Droguería Central" in pagina and "Plan por agente comercial" in pagina)


def main():
    test_columnas()
    test_estrategias_y_valores()
    test_jugadas_y_a_quien_visitar()
    test_agentes()
    test_donde_abrir()
    test_promedio_no_tiene_plan()
    test_informes_con_plan()


if __name__ == "__main__":
    main()
