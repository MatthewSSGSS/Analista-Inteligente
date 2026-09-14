"""Hojas tipo informe: varias tablas, títulos, gráficos e imágenes en un mismo Excel.

Se construye un Excel con la forma de los informes comerciales reales, y se
exige que el panel lo sepa leer solo, sin nombres fijos:

- una tabla con los meses en columnas (y un hueco entre ellas) y las medidas
  en filas;
- bloques por región lado a lado, con Ejc / Ppto / Cum repetidos;
- una tabla resumen, y dos secciones con título (SIN y CON PYME) que no se
  deben mezclar;
- un ranking con los años combinados sobre meses sin año, porcentajes
  escritos como texto y fila de Total;
- una hoja normal con título, que tiene que seguir leyéndose como antes;
- un gráfico nativo y una imagen pegada.

PYTHONPATH=. python tests/informe_test.py
"""
import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference

from core.dashboard_engine import build_dashboard
from core.informe import leer_grafico, periodo_de
from core.loader import load_workbook

MESES = [dt.datetime(2026, m, 1) for m in range(1, 5)]


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


class _Subida:
    def __init__(self, nombre, libro):
        buf = io.BytesIO()
        libro.save(buf)
        self.name, self._datos = nombre, buf.getvalue()

    def getvalue(self):
        return self._datos


def _libro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws["B2"] = "ALTAS Y CUMPLIMIENTO SIN PYME"
    ws["B4"], ws["B5"] = "R1", "Mes"
    columnas = [3, 4, 6, 7]  # hueco en la columna 5, como en el informe real
    for c, m in zip(columnas, MESES):
        ws.cell(5, c, m)
    ws["B6"], ws["B8"], ws["B9"] = "Ejecución", "Presupuesto", "Cumplimiento"
    for c, e in zip(columnas, [100, 90, 95, 110]):
        ws.cell(6, c, e)
        ws.cell(8, c, 100)
        ws.cell(9, c, e / 100)
    # Bloques R2 y R3 lado a lado, meses bajando por la columna B.
    datos = {"R2": [50, 60, 70, 80], "R3": [120, 110, 100, 90]}
    for base, region in ((3, "R2"), (7, "R3")):
        ws.cell(11, base + 1, region)
        for k, nombre in enumerate(["Ejc", "Ppto", "Cum"]):
            ws.cell(12, base + k, nombre)
        for i, m in enumerate(MESES):
            ws.cell(13 + i, 2, m)
            ws.cell(13 + i, base, datos[region][i])
            ws.cell(13 + i, base + 1, 100)
            ws.cell(13 + i, base + 2, datos[region][i] / 100)
    ws["C19"] = "Prom Cum"
    for i, (region, v) in enumerate([("R1", 0.99), ("R2", 0.65), ("R3", 1.05)]):
        ws.cell(20 + i, 2, region)
        ws.cell(20 + i, 3, v)
    ws["B25"] = "ALTAS Y CUMPLIMIENTO CON PYME"
    ws["B27"], ws["B28"] = "R1", "Mes"
    for c, m in zip(columnas, MESES):
        ws.cell(28, c, m)
    ws["B29"], ws["B30"], ws["B32"] = "Ejecución", "Presupuesto", "GAP PYME"
    for c in columnas:
        ws.cell(29, c, 120)
        ws.cell(30, c, 100)
        ws.cell(32, c, 20)

    rk = wb.create_sheet("Ranking")
    rk["A1"], rk["A3"], rk["A4"] = "RANKING DE LOS JEFES", "Año", "Jefe"
    rk["B3"], rk["D3"], rk["F3"] = 2025, 2026, "Total"
    rk.merge_cells("B3:C3")
    rk.merge_cells("D3:E3")
    for i, mes in enumerate(["Nov", "Dic", "Ene", "Feb"]):
        rk.cell(4, 2 + i, mes)
    jefes = {"ANA PEREZ": [80, 90, 100, 110], "LUIS GOMEZ": [60, 70, 65, 50], "EVA DIAZ": [100, 100, 95, 90]}
    for f, (jefe, valores) in enumerate(jefes.items()):
        rk.cell(5 + f, 1, jefe)
        for i, v in enumerate(valores):
            rk.cell(5 + f, 2 + i, f"{v} %")
        rk.cell(5 + f, 6, f"{sum(valores) // 4} %")
    rk.cell(8, 1, "Total")
    for i in range(4):
        rk.cell(8, 2 + i, "80 %")

    vt = wb.create_sheet("Ventas")
    vt["A1"] = "INFORME DE VENTAS DIARIAS"
    for j, nombre in enumerate(["Fecha", "Ciudad", "Ventas"]):
        vt.cell(3, 1 + j, nombre)
    for d in range(1, 11):
        vt.cell(3 + d, 1, dt.datetime(2026, 1, d))
        vt.cell(3 + d, 2, ["Cali", "Bogota"][d % 2])
        vt.cell(3 + d, 3, d * 10)

    grafico = LineChart()
    grafico.title = "VENTAS POR DIA"
    grafico.add_data(Reference(vt, min_col=3, min_row=3, max_row=13), titles_from_data=True)
    grafico.set_categories(Reference(vt, min_col=1, min_row=4, max_row=13))
    ws.add_chart(grafico, "K2")
    try:
        from PIL import Image as PILImage
        from openpyxl.drawing.image import Image as XLImage
        imagen = io.BytesIO()
        PILImage.new("RGB", (40, 20), (200, 30, 60)).save(imagen, "PNG")
        imagen.seek(0)
        ws.add_image(XLImage(imagen), "K20")
    except ImportError:
        pass
    return wb


def test_los_meses_se_reconocen_en_cualquier_forma():
    check("«ene-26»", periodo_de("ene-26") == (2026, 1))
    check("«Oct» sin año", periodo_de("Oct") == (None, 10))
    check("«09. septiembre 2024»", periodo_de("09. septiembre 2024") == (2024, 9))
    check("una fecha de Excel", periodo_de(dt.datetime(2026, 3, 1)) == (2026, 3))
    check("un texto que empieza como mes no es un mes", periodo_de("Mar Azul") is None)
    check("un número no es un mes", periodo_de(5) is None)


def test_el_informe_se_convierte_en_tablas_ordenadas():
    libro = load_workbook(_Subida("informe.xlsx", _libro()))
    hojas = libro["sheets"]
    sin = hojas.get("ALTAS Y CUMPLIMIENTO SIN PYME")
    check("cada sección con título es su propia tabla", sin is not None and "ALTAS Y CUMPLIMIENTO CON PYME" in hojas)
    df = sin["processed"]
    check("la tabla sale con una fila por región y mes",
          list(df.columns) == ["Región", "Mes", "Ejecución", "Presupuesto", "Cumplimiento"] and len(df) == 12)
    check("la tabla con meses en columnas entra con su grupo (R1)", set(df["Región"]) == {"R1", "R2", "R3"})
    fila = df[(df["Región"] == "R2") & (pd.to_datetime(df["Mes"]) == "2026-02-01")]
    check("«Ejc» y «Ppto» de los bloques se unen con «Ejecución» y «Presupuesto»",
          len(fila) == 1 and float(fila["Ejecución"].iloc[0]) == 60 and float(fila["Presupuesto"].iloc[0]) == 100)
    fila = df[(df["Región"] == "R3") & (pd.to_datetime(df["Mes"]) == "2026-01-01")]
    check("el cumplimiento guardado como fracción se lee como porcentaje", float(fila["Cumplimiento"].iloc[0]) == 120)
    check("el mes es una fecha de verdad", "Mes" in sin["profile"]["schema"]["dates"])
    check("el título se conserva", sin["profile"].get("titulo") == "ALTAS Y CUMPLIMIENTO SIN PYME")

    tablero = build_dashboard(df, sin["profile"])
    check("la ejecución se suma entre regiones, no se promedia",
          tablero["performance"]["additive"]
          and round(tablero["executive"].get("current_period", 0)) == 110 + 80 + 90)
    base = tablero["performance"]["base"]
    check("la comparación entre regiones usa el presupuesto como meta", base["clave"] == "meta")
    check("y gana quien más cumplió (R3, 105%), no quien más vendió", base["mejor"]["nombre"] == "R3")

    con = hojas["ALTAS Y CUMPLIMIENTO CON PYME"]["processed"]
    check("la sección CON PYME trae su fila extra", "GAP PYME" in con.columns)
    check("y no se mezcla con SIN PYME (no se suma dos veces)", float(con["Ejecución"].sum()) == 480)

    resumen = hojas.get("ALTAS Y CUMPLIMIENTO SIN PYME · Resumen")
    check("la tabla resumen se conserva aparte", resumen is not None)
    check("con sus valores en porcentaje",
          float(resumen["processed"].set_index("Región").loc["R2", "Prom Cum"]) == 65)


def test_el_ranking_con_anios_arriba():
    libro = load_workbook(_Subida("informe.xlsx", _libro()))
    item = libro["sheets"].get("RANKING DE LOS JEFES")
    check("el ranking se lee con su título", item is not None and item["profile"].get("titulo") == "RANKING DE LOS JEFES")
    df = item["processed"]
    check("una fila por jefe y mes", list(df.columns) == ["Jefe", "Mes", "Valor (%)"] and len(df) == 12)
    check("la fila Total no es un jefe", "Total" not in set(df["Jefe"]))
    meses = pd.to_datetime(df["Mes"])
    check("el año sale de la fila de arriba (Nov y Dic de 2025, Ene y Feb de 2026)",
          meses.min() == pd.Timestamp("2025-11-01") and meses.max() == pd.Timestamp("2026-02-01"))
    fila = df[(df["Jefe"] == "ANA PEREZ") & (meses == "2026-02-01")]
    check("«110 %» escrito como texto se lee como 110", float(fila["Valor (%)"].iloc[0]) == 110)
    ejecutivo = build_dashboard(df, item["profile"])["executive"]
    check("un porcentaje se promedia entre jefes, no se suma",
          round(ejecutivo.get("current_period", 0), 1) == round((110 + 50 + 90) / 3, 1))


def test_una_hoja_normal_sigue_igual_y_guarda_su_titulo():
    libro = load_workbook(_Subida("informe.xlsx", _libro()))
    item = libro["sheets"].get("Ventas")
    check("la hoja normal conserva su nombre", item is not None)
    check("y sus columnas de siempre", {"Fecha", "Ciudad", "Ventas"} <= set(item["processed"].columns))
    check("guarda el título que tenía encima", item["profile"].get("titulo") == "INFORME DE VENTAS DIARIAS")


def test_graficos_e_imagenes():
    libro = load_workbook(_Subida("informe.xlsx", _libro()))
    grafico = libro["sheets"].get("Gráfico · VENTAS POR DIA")
    check("los datos de un gráfico nativo se pueden analizar", grafico is not None)
    check("con su serie y sus fechas", {"Mes", "Ventas"} <= set(grafico["processed"].columns)
          and len(grafico["processed"]) == 10)
    check("una imagen pegada se avisa en vez de ignorarse", any("imagen" in a for a in libro.get("avisos", [])))

    xml = (b'<?xml version="1.0" encoding="UTF-8"?>'
           b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
           b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>'
           b'<c:title><c:tx><c:rich><a:p><a:r><a:t>PRODUCTIVIDAD</a:t></a:r></a:p></c:rich></c:tx></c:title>'
           b'<c:plotArea><c:lineChart><c:ser>'
           b'<c:tx><c:strRef><c:f>H!$B$1</c:f><c:strCache><c:ptCount val="1"/>'
           b'<c:pt idx="0"><c:v>Productividad</c:v></c:pt></c:strCache></c:strRef></c:tx>'
           b'<c:cat><c:numRef><c:f>H!$A$2:$A$3</c:f><c:numCache><c:formatCode>mmm-yy</c:formatCode>'
           b'<c:ptCount val="2"/><c:pt idx="0"><c:v>45536</c:v></c:pt><c:pt idx="1"><c:v>45566</c:v></c:pt>'
           b'</c:numCache></c:numRef></c:cat>'
           b'<c:val><c:numRef><c:f>H!$B$2:$B$3</c:f><c:numCache><c:formatCode>General</c:formatCode>'
           b'<c:ptCount val="2"/><c:pt idx="0"><c:v>7.7</c:v></c:pt><c:pt idx="1"><c:v>7.6</c:v></c:pt>'
           b'</c:numCache></c:numRef></c:val></c:ser></c:lineChart></c:plotArea></c:chart></c:chartSpace>')
    leido = leer_grafico(xml)
    check("un gráfico real de Excel se lee desde la copia de datos que guarda",
          leido["titulo"] == "PRODUCTIVIDAD" and list(leido["datos"]["Productividad"]) == [7.7, 7.6])
    check("y sus categorías de fecha salen como meses",
          list(leido["datos"]["Mes"]) == [pd.Timestamp("2024-09-01"), pd.Timestamp("2024-10-01")])


if __name__ == "__main__":
    test_los_meses_se_reconocen_en_cualquier_forma()
    test_el_informe_se_convierte_en_tablas_ordenadas()
    test_el_ranking_con_anios_arriba()
    test_una_hoja_normal_sigue_igual_y_guarda_su_titulo()
    test_graficos_e_imagenes()
    print("\nInforme test completado sin errores.")
