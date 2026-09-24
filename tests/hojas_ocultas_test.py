"""Hojas ocultas en Excel: se cargan, pero no se muestran hasta que se piden.

El archivo real («Inf ventas - R1.xlsb») tiene 8 hojas y 6 están ocultas: la
base sin depurar, una copia de 63.711 filas y varias tablas de apoyo. El panel
las abría todas por igual, así que el selector mostraba 8 entradas y el análisis
arrancaba en «Hoja2» —siete filas de apoyo— en vez de en la base de 64.646
registros.

Se comprueba que el estado se lee de verdad del archivo (no se adivina por el
nombre), que la hoja oculta sigue cargada para poder mostrarla a pedido, y que
un archivo sin hojas ocultas no cambia en nada.

PYTHONPATH=. python tests/hojas_ocultas_test.py
"""
import io

import pandas as pd
from openpyxl import Workbook

from core.hojas_ocultas import hojas_ocultas
from core.loader import load_workbook


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


class _Subida:
    def __init__(self, nombre, libro):
        buf = io.BytesIO()
        libro.save(buf)
        self.name, self._datos = nombre, buf.getvalue()

    def getvalue(self):
        return self._datos


def _hoja(ws, prefijo):
    for j, nombre in enumerate(["Ciudad", "Ventas", "Costo"]):
        ws.cell(1, 1 + j, nombre)
    for i in range(1, 8):
        ws.cell(1 + i, 1, f"{prefijo}{i}")
        ws.cell(1 + i, 2, i * 10)
        ws.cell(1 + i, 3, i * 4)


def _libro(ocultar=True):
    wb = Workbook()
    base = wb.active
    base.title = "Base"
    _hoja(base, "C")
    for nombre, estado in (("Apoyo", "hidden"), ("Respaldo", "veryHidden")):
        ws = wb.create_sheet(nombre)
        _hoja(ws, nombre[0])
        if ocultar:
            ws.sheet_state = estado
    visible = wb.create_sheet("Resumen")
    _hoja(visible, "R")
    return wb


def test_se_leen_del_archivo_las_hojas_ocultas():
    datos = _Subida("informe.xlsx", _libro()).getvalue()
    check("se leen las ocultas y las muy ocultas, y solo esas",
          hojas_ocultas(datos, "informe.xlsx") == {"Apoyo", "Respaldo"})
    check("un archivo sin ninguna oculta devuelve conjunto vacío",
          hojas_ocultas(_Subida("x.xlsx", _libro(ocultar=False)).getvalue(), "x.xlsx") == set())
    check("un formato que no lo dice (CSV) no rompe nada",
          hojas_ocultas(b"a,b\n1,2\n", "datos.csv") == set())
    check("un archivo dañado tampoco", hojas_ocultas(b"no soy un excel", "roto.xlsx") == set())


def test_la_hoja_oculta_se_carga_pero_queda_marcada():
    libro = load_workbook(_Subida("informe.xlsx", _libro()))
    check("las cuatro hojas se cargan igual que antes", len(libro["sheets"]) == 4)
    check("las ocultas quedan señaladas", libro["ocultas"] == {"Apoyo", "Respaldo"})
    check("y sus datos siguen ahí para poder mostrarlas a pedido",
          len(libro["sheets"]["Apoyo"]["processed"]) == 7)
    check("se avisa, diciendo cuáles y dónde activarlas",
          any("ocultas en Excel" in a and "Apoyo" in a for a in libro["avisos"]))

    visibles = [h for h in libro["sheets"] if h not in libro["ocultas"]]
    check("el panel abriría en la primera hoja visible, no en una de apoyo",
          visibles[0] == "Base" and visibles == ["Base", "Resumen"])


def test_un_archivo_normal_no_cambia():
    libro = load_workbook(_Subida("informe.xlsx", _libro(ocultar=False)))
    check("sin hojas ocultas no hay nada marcado", libro["ocultas"] == set())
    check("y no se avisa de nada", not any("ocultas en Excel" in a for a in libro["avisos"]))

    csv = type("U", (), {"name": "datos.csv", "getvalue": lambda self: b"Ciudad,Ventas\nCali,10\nBogota,20\n"})()
    check("un CSV sigue igual", load_workbook(csv)["ocultas"] == set())


def test_si_no_queda_ninguna_visible_se_muestran_igual():
    """Excel no deja ocultar TODAS las hojas, pero sí puede pasar que la única
    visible no traiga tabla (una portada con texto suelto) y se descarte. Si
    esconder las ocultas dejara el panel en blanco, se muestran: es peor no
    tener nada que analizar que abrir en una hoja de trabajo."""
    wb = Workbook()
    portada = wb.active
    portada.title = "Portada"
    portada["A1"] = "Informe comercial"
    portada["A2"] = "Preparado por el equipo de ventas."
    for nombre in ("Datos", "Respaldo"):
        ws = wb.create_sheet(nombre)
        _hoja(ws, nombre[0])
        ws.sheet_state = "hidden"
    libro = load_workbook(_Subida("solo_portada.xlsx", wb))
    check("las hojas ocultas son lo único que quedó", set(libro["sheets"]) == {"Datos", "Respaldo"})
    check("no se avisa de esconderlas: no queda ninguna otra",
          not any("ocultas en Excel" in a for a in libro["avisos"]))
    check("y ninguna queda marcada, así que el selector las muestra", libro["ocultas"] == set())


if __name__ == "__main__":
    test_se_leen_del_archivo_las_hojas_ocultas()
    test_la_hoja_oculta_se_carga_pero_queda_marcada()
    test_un_archivo_normal_no_cambia()
    test_si_no_queda_ninguna_visible_se_muestran_igual()
    print("\nHojas ocultas test completado sin errores.")
