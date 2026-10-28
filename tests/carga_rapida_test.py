"""La carga rápida lee exactamente lo mismo que la lenta.

Un Excel de 6 MB tardaba ~36 s en cargar. Se aceleró (~13 s) sin cambiar
qué se lee:

- calamine en vez de openpyxl para .xlsx (core/loader._excel_engine), con
  el libro abierto una vez para todas las hojas;
- las celdas combinadas y las filas/columnas ocultas se leen sin construir
  el árbol XML de todas las celdas (pivot_flatten.xml_sin_celdas);
- las funciones celda a celda de core/informe se memorizan por valor;
- las transformaciones de texto de fechas, meses, zonas horarias y vacíos
  se calculan sobre los valores distintos y se reparten a cada fila.

Cada test compara la versión rápida contra la forma directa de antes.

PYTHONPATH=. python tests/carga_rapida_test.py
"""
import io
import logging
import warnings

import numpy as np
import pandas as pd
from openpyxl import Workbook

warnings.filterwarnings("ignore")
logging.disable(logging.WARNING)


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _libro_con_ocultos() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tablero"
    for r in range(1, 40):
        ws.append([f"t{r}", r, r * 2, r * 3, "x"])
    for r in (3, 7, 8, 21):
        ws.row_dimensions[r].hidden = True
    ws.column_dimensions["C"].hidden = True
    ws.column_dimensions["E"].width = 0.5
    ws.merge_cells("A1:D1")
    ws.merge_cells("B10:C12")
    otra = wb.create_sheet("Otra")
    otra.append(["a", "b"])
    otra.merge_cells("A5:B5")
    wb.create_sheet("Vacia")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_xml_sin_celdas():
    from core.celdas_ocultas import filas_y_columnas_ocultas
    from core.pivot_flatten import merged_ranges_by_sheet, xml_sin_celdas

    data = _libro_con_ocultos()
    check("las combinadas se leen sin las celdas",
          merged_ranges_by_sheet(data, "a.xlsx") == {"Tablero": [(0, 0, 0, 3), (9, 1, 11, 2)],
                                                     "Otra": [(4, 0, 4, 1)], "Vacia": []})
    ocultas = filas_y_columnas_ocultas(data, "a.xlsx")
    check("las filas ocultas se encuentran como texto", ocultas["Tablero"][0] == {2, 6, 7, 20})
    check("y las columnas ocultas o angostas, en el árbol sin celdas", ocultas["Tablero"][1] == {2, 4})
    xml = b'<worksheet><cols/><sheetData><row r="1"><c r="A1"/></row></sheetData><mergeCells count="0"/></worksheet>'
    check("xml_sin_celdas quita solo el bloque de celdas",
          xml_sin_celdas(xml) == b'<worksheet><cols/><mergeCells count="0"/></worksheet>')
    check("con <sheetData/> vacío no cambia nada", xml_sin_celdas(b"<w><sheetData/></w>") == b"<w><sheetData/></w>")
    check("con prefijo de espacio de nombres también",
          xml_sin_celdas(b"<x:w><x:sheetData><x:row/></x:sheetData></x:w>") == b"<x:w></x:w>")


def test_informe_memorizado():
    import datetime as dt
    from core import informe

    valores = ["Enero 2026", " Mar  ", "a\tb\n c", None, np.nan, 5, 3.2, dt.date(2026, 5, 1), ""]
    esperados = ["Enero 2026", "Mar", "a b c", "", "", "5", "3.2", "2026-05-01", ""]
    for vuelta in ("primera", "segunda (desde la memoria)"):
        check(f"_txt da lo mismo en la {vuelta} vuelta", [informe._txt(v) for v in valores] == esperados)
    check("periodo_de reconoce meses con y sin año",
          informe.periodo_de("Enero 2026") == (2026, 1) and informe.periodo_de("ene-26") == (2026, 1)
          and informe.periodo_de("Mar") == (None, 3) and informe.periodo_de("2026-03") == (2026, 3))
    check("y fechas, sin confundir números ni texto",
          informe.periodo_de(pd.Timestamp("2026-04-01")) == (2026, 4) and informe.periodo_de(5) is None
          and informe.periodo_de("ventas") is None)
    check("_norm quita tildes y pasa a minúsculas", informe._norm("Ñandú  Árbol") == "nandu arbol")


def test_valores_distintos():
    from core.dates import month_number_series
    from core.filter_engine import es_vacio

    s = pd.Series(["Enero", " FEB ", "marzo", None, "Enero", "x", 3, np.nan, "Dic"], index=range(10, 19))
    esperado = [1, 2, 3, pd.NA, 1, pd.NA, pd.NA, pd.NA, 12]
    obtenido = month_number_series(s)
    check("month_number_series da el mes de cada fila", [None if pd.isna(v) else int(v) for v in obtenido]
          == [None if v is pd.NA else v for v in esperado])
    check("con el índice y el tipo de siempre", list(obtenido.index) == list(s.index) and str(obtenido.dtype) == "Int64")

    viejo = lambda x: x.isna() | x.astype(str).str.strip().str.lower().isin({"", "nan", "nat", "none", "<na>"})
    for caso in (pd.Series(["a", None, "", "  ", "NaN", "None", "x", np.nan, "<NA>", "nat"]),
                 pd.Series([1, 2, np.nan, 4.0]), pd.Series(pd.to_datetime(["2026-01-01", None])),
                 pd.Series([None, None]), pd.Series([], dtype=object),
                 pd.Series(["a", pd.NA, "b"], dtype="string"), pd.Series([1, "1", 1.0, "  nan "])):
        check(f"es_vacio igual que antes en {caso.tolist()!r}", list(es_vacio(caso)) == list(viejo(caso)))


def test_calamine_lee_lo_mismo():
    from core import loader
    if not loader._CALAMINE:
        print("SKIP calamine no está instalado: la app usa openpyxl")
        return
    wb = Workbook()
    ws = wb.active
    ws.append(["Fecha", "Codigo", "Cedula", "Pct", "Texto", "Bool"])
    for i in range(1, 30):
        ws.append([pd.Timestamp(2026, 1 + i % 12, 1).to_pydatetime(), f"00{i}", 1047382910 + i, i / 7,
                   f"Ñandú {i}", i % 2 == 0])
    b = io.BytesIO()
    wb.save(b)
    data = b.getvalue()
    a = pd.read_excel(io.BytesIO(data), header=None, engine="openpyxl")
    c, _ = loader._grid(data, 0, None, "calamine", pd.ExcelFile(io.BytesIO(data), engine="calamine"))
    check("calamine lee la misma cuadrícula que openpyxl",
          a.shape == c.shape and all((a[k].astype(str) == c[k].astype(str)).all() for k in a.columns))
    check("el lector elegido para .xlsx es calamine", loader._excel_engine("x.xlsx") == "calamine")


def _base_grande(n=100_000):
    rng = np.random.default_rng(3)
    return pd.DataFrame({
        "Asesor": rng.choice([f"Asesor {i}" for i in range(80)], n),
        "Ciudad": rng.choice(["Bogotá", " Cali ", "Medellín\n", "N/A"], n),
        "Ventas": rng.integers(1, 10**6, n).astype(object),
        "Cédula": rng.integers(10**7, 10**9, n).astype(object),
    })


def test_filas_basura_no_va_fila_por_fila():
    """core/filas_basura iba fila por fila: más de 4 minutos en una base de
    120.000 registros. Ahora es por columnas y debe seguir encontrando el
    encabezado repetido."""
    import time
    from core.filas_basura import quitar_filas_basura
    df = _base_grande()
    df.iloc[50_000] = ["Asesor", "Ciudad", "Ventas", "Cédula"]
    t = time.time()
    limpio, log = quitar_filas_basura(df)
    check("filas basura en 100.000 filas tarda menos de 3 s", time.time() - t < 3)
    check("y sigue quitando el encabezado repetido", len(limpio) == len(df) - 1 and log)


def test_base_plana_no_pasa_por_el_lector_de_informes():
    """Una base grande y llena no es un informe: recorrerla celda por celda
    buscando títulos y meses costaba ~12 s. Un informe con meses en columnas,
    o una hoja chica, sí pasan por core/informe."""
    from core.loader import _es_base_plana
    base = _base_grande(5000)
    grid = pd.concat([pd.DataFrame([list(base.columns)], columns=base.columns), base], ignore_index=True)
    grid.columns = range(grid.shape[1])
    check("una base de 5.000 filas llena se reconoce", _es_base_plana(grid))
    con_meses = grid.copy(); con_meses.iloc[0, 2] = "Enero"
    check("con meses en el encabezado, no (puede ser un informe)", not _es_base_plana(con_meses))
    check("una hoja chica siempre pasa por el lector de informes", not _es_base_plana(grid.head(500)))


def test_limpieza_por_valores_distintos_da_lo_mismo():
    """core/cleaner limpia las columnas de solo texto sobre sus valores
    distintos y convierte directo las de solo números: debe dar lo mismo que
    hacerlo celda por celda."""
    from core.cleaner import clean
    from core.numeric import normalize_missing_series
    df = _base_grande(3000)
    limpio, _ = clean(df)
    esperado = normalize_missing_series(df["Ciudad"].astype("string").str.replace(r"[\r\n\t]", " ", regex=True).str.strip())
    check("el texto queda igual que celda por celda",
          limpio["Ciudad"].fillna("<vacío>").tolist() == esperado.fillna("<vacío>").tolist())
    check("«N/A» sigue siendo vacío", limpio["Ciudad"].isna().sum() == (df["Ciudad"] == "N/A").sum())
    check("las cifras se convierten igual",
          limpio["Ventas"].tolist() == pd.to_numeric(df["Ventas"].astype("string")).tolist())


def main():
    test_xml_sin_celdas()
    test_informe_memorizado()
    test_valores_distintos()
    test_calamine_lee_lo_mismo()
    test_filas_basura_no_va_fila_por_fila()
    test_base_plana_no_pasa_por_el_lector_de_informes()
    test_limpieza_por_valores_distintos_da_lo_mismo()


if __name__ == "__main__":
    main()
