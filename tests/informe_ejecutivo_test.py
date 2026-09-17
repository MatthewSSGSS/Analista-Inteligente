"""Informes comerciales con formas que antes se perdían o se leían mal.

Se reproduce con un Excel sintético cada forma del informe real "INFORMACION
PARA DASH BOARD NUEVO EJECUTIVO DE CALLE", y se exige que nada se pierda:

- meses combinados sobre Altas / Meta / Cum (antes: solo las altas del primer mes);
- meses intercalados con "% Participación" (antes: se perdía la participación);
- un listado normal dentro del informe (antes: desaparecía entero);
- un título de una palabra ("FWA") con R1…R5 debajo (antes: solo enero, y
  bajo el título de una imagen);
- agentes sin encabezado ("BMS", "RYL") que NO son títulos;
- un mes sin reportar que no es cero (antes: "cayó 100%");
- NIT repetido que no se suma, tablas copiadas en otra hoja que no se
  duplican, e imágenes avisadas por el nombre de su sección.

PYTHONPATH=. python tests/informe_ejecutivo_test.py
"""
import datetime as dt
import io
import warnings

import pandas as pd
from openpyxl import Workbook

from core.dashboard_engine import build_dashboard
from core.loader import load_workbook

warnings.filterwarnings("ignore")
MESES = [dt.datetime(2026, m, 1) for m in (4, 5, 6)]


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


class _Subida:
    def __init__(self, nombre, libro):
        buf = io.BytesIO()
        libro.save(buf)
        self.name, self._datos = nombre, buf.getvalue()

    def getvalue(self):
        return self._datos


def _bloque_otts(ws, fila=2):
    ws.cell(fila, 2, "CUMPLIMIENTO OTTS CALLE")
    h = fila + 2
    ws.cell(h, 2, "Región")
    ws.merge_cells(start_row=h, start_column=2, end_row=h + 1, end_column=2)
    for i, mes in enumerate(MESES):
        c = 3 + 3 * i
        ws.cell(h, c, mes)
        ws.merge_cells(start_row=h, start_column=c, end_row=h, end_column=c + 2)
        for k, nombre in enumerate(["Altas", "Meta", "Cum"]):
            ws.cell(h + 1, c + k, nombre)
    datos = {"R1": (5196, 4372), "R2": (6634, 5399), "R3": (6098, 6500)}
    for j, (region, (altas, meta)) in enumerate(datos.items()):
        ws.cell(h + 2 + j, 2, region)
        for i in range(3):
            c = 3 + 3 * i
            ws.cell(h + 2 + j, c, altas + 10 * i)
            ws.cell(h + 2 + j, c + 1, meta)
            ws.cell(h + 2 + j, c + 2, (altas + 10 * i) / meta)
    return h + 2 + len(datos)


def _libro():
    wb = Workbook()
    ws = wb.active
    ws.title = "BUENAS NOTICIAS"
    fin = _bloque_otts(ws)

    # Productos con el mes y su participación intercalados, y fila de Total.
    p = fin + 2
    ws.cell(p, 2, "Productos")
    for i, mes in enumerate(MESES):
        ws.cell(p, 3 + 2 * i, mes)
        ws.cell(p, 4 + 2 * i, "% Participación")
    productos = {"Netflix": [1330, 1498, 1474], "Win+": [1035, 1139, 1325], "Amazon Prime": [874, 842, 942]}
    totales = [sum(v[i] for v in productos.values()) for i in range(3)]
    for j, (prod, valores) in enumerate(productos.items()):
        ws.cell(p + 1 + j, 2, prod)
        for i, v in enumerate(valores):
            ws.cell(p + 1 + j, 3 + 2 * i, v)
            ws.cell(p + 1 + j, 4 + 2 * i, v / totales[i])
    ws.cell(p + 4, 2, "Total")
    for i, t in enumerate(totales):
        ws.cell(p + 4, 3 + 2 * i, t)

    # Listado normal: agentes con su jefe y NIT repetido.
    d = p + 7
    ws.cell(d, 2, "DISTRIBUCIÓN DE LA ZONIFICACION")
    for k, nombre in enumerate(["TIPO", "REGION", "AGENTE", "NIT", "JEFE", "@"]):
        ws.cell(d + 2, 1 + k, nombre)
    filas = [("UNICO", "COSTA", "GERA SAS", 9000236859, "JIMENEZ", 333),
             ("MIXTO", "COSTA", "GERA SAS", 9000236859, "ORTIZ", 261),
             ("MIXTO", "COSTA", "INVERCELL", 9007380074, "CASTILLO", 396),
             ("UNICO", "COSTA", "COMTEL SAS", 8190056771, "JIMENEZ", 245),
             ("MIXTO", "COSTA", "COMTEL SAS", 8190056771, "ROJAS", 119)]
    for j, f in enumerate(filas):
        for k, v in enumerate(f):
            ws.cell(d + 3 + j, 1 + k, v)

    # Casos de éxito: dos agentes, cada uno con su tabla; el último mes sin ingresos.
    e = d + 10
    ws.cell(e, 2, "CASOS DE ÉXITO AGENTES CALLE COSTA")
    for bloque, (agente, ejec, ingresos) in enumerate([("BMS", [592, 615, 570], [190.4, 189.9, None]),
                                                        ("RYL", [15, 32, 47], [2.5, 8.0, None])]):
        b = e + 2 + bloque * 7
        ws.cell(b, 2, agente)
        ws.cell(b + 1, 2, "Mes")
        for i, mes in enumerate(MESES):
            ws.cell(b + 1, 3 + i, mes)
            ws.cell(b + 2, 3 + i, ejec[i])
            if ingresos[i] is not None:
                ws.cell(b + 3, 3 + i, ingresos[i])
            ws.cell(b + 4, 3 + i, 500 if agente == "BMS" else 20)
        ws.cell(b + 2, 2, "Ejecución @")
        ws.cell(b + 3, 2, "Ingresos Comisión (Millones)")
        ws.cell(b + 4, 2, "Presupuesto @")

    # Un ranking que es una imagen, y debajo una tabla con título de una palabra.
    r = e + 18
    ws.cell(r, 2, "RANKING DE LOS JEFES")
    from PIL import Image as PILImage
    from openpyxl.drawing.image import Image as XLImage
    imagen = io.BytesIO()
    PILImage.new("RGB", (40, 20), (200, 30, 60)).save(imagen, "PNG")
    imagen.seek(0)
    ws.add_image(XLImage(imagen), f"B{r + 1}")
    f = r + 12
    ws.cell(f, 2, "FWA")
    ws.cell(f + 2, 2, "Zona")
    for i, mes in enumerate(MESES):
        ws.cell(f + 2, 3 + i, mes)
    for j, (zona, valores) in enumerate({"R1": [382, 288, 356], "R2": [57, 57, 79], "R3": [10, 5, 3]}.items()):
        ws.cell(f + 3 + j, 2, zona)
        for i, v in enumerate(valores):
            ws.cell(f + 3 + j, 3 + i, v)

    # Otra hoja que copia la tabla de OTTS tal cual.
    _bloque_otts(wb.create_sheet("INDICADORES CLAVE"))
    return wb


_CACHE = {}


def _leido():
    if "libro" not in _CACHE:
        _CACHE["libro"] = load_workbook(_Subida("informe_ejecutivo.xlsx", _libro()))
    return _CACHE["libro"]


def _tabla(parte):
    hojas = _leido()["sheets"]
    nombre = next((n for n in hojas if parte in n), None)
    return nombre, (hojas[nombre] if nombre else None)


def test_meses_combinados_sobre_medidas():
    nombre, item = _tabla("OTTS CALLE · por Región")
    check("la tabla por región existe y se separa de la de productos", item is not None)
    df = item["processed"]
    check("trae Altas, Meta y Cumplimiento", {"Región", "Mes", "Altas", "Meta", "Cumplimiento"} <= set(df.columns))
    check("de los tres meses y las tres regiones, no solo el primero", len(df) == 9 and df["Mes"].nunique() == 3)
    r1 = df[(df["Región"] == "R1")].sort_values("Mes").iloc[0]
    check("el cumplimiento queda en porcentaje", abs(r1["Cumplimiento"] - 5196 / 4372 * 100) < 0.05)


def test_mes_intercalado_con_participacion():
    _, item = _tabla("por Productos")
    df = item["processed"]
    check("los productos tienen su tabla", set(df["Productos"]) == {"Netflix", "Win+", "Amazon Prime"})
    check("la fila de Total no es un producto", "Total" not in set(df["Productos"]))
    check("la medida sin encabezado toma el nombre de la tabla hermana", "Altas" in df.columns)
    check("y la participación no se pierde", "% Participación" in df.columns)
    check("la participación queda en porcentaje", df["% Participación"].between(0, 100).all()
          and df["% Participación"].max() > 1)


def test_listado_normal_dentro_del_informe():
    nombre, item = _tabla("DISTRIBUCIÓN")
    check("el listado ya no desaparece", item is not None and len(item["processed"]) == 5)
    schema = item["profile"]["schema"]
    check("el NIT repetido es un identificador, no una métrica", "NIT" in schema["ids"] and "NIT" not in schema["metrics"])
    check("la columna @ sí es métrica", "@" in schema["metrics"])


def test_agentes_sin_encabezado_no_son_titulos():
    nombres = list(_leido()["sheets"])
    check("BMS y RYL no se vuelven tablas sueltas", not any(n in {"BMS", "RYL"} for n in nombres))
    _, item = _tabla("CASOS DE ÉXITO")
    df = item["processed"]
    check("los dos agentes quedan en la misma tabla", set(df["Agente"]) == {"BMS", "RYL"})
    ultimo = df[df["Mes"] == df["Mes"].max()]
    check("un mes sin reportar queda vacío, no en cero", ultimo["Ingresos Comisión (Millones)"].isna().all())
    tablero = build_dashboard(df, {**item["profile"], "schema": {**item["profile"]["schema"],
                                                               "metrica_preferida": "Ingresos Comisión (Millones)"}})
    check("y el panel no dice que los ingresos cayeron 100%",
          tablero["growth"] is None or tablero["growth"] > -99)


def test_titulo_de_una_palabra_y_secciones_con_imagen():
    nombres = list(_leido()["sheets"])
    check("FWA tiene su propia tabla", "FWA" in nombres)
    check("ninguna tabla queda bajo el título del ranking que es imagen",
          not any("RANKING" in n for n in nombres))
    df = _leido()["sheets"]["FWA"]["processed"]
    check("con los tres meses de cada zona", len(df) == 9 and "Zona" in df.columns and "FWA" in df.columns)
    avisos = _leido()["avisos"]
    check("la imagen se avisa con el nombre de su sección",
          any("RANKING DE LOS JEFES" in a and "imagen" in a for a in avisos))


def test_tabla_copiada_en_otra_hoja():
    nombres = [n for n in _leido()["sheets"] if "OTTS" in n]
    check("la copia idéntica no se duplica", len(nombres) == 2 and not any("(2)" in n for n in nombres))
    check("y se avisa", any("idéntica" in a for a in _leido()["avisos"]))


if __name__ == "__main__":
    test_meses_combinados_sobre_medidas()
    test_mes_intercalado_con_participacion()
    test_listado_normal_dentro_del_informe()
    test_agentes_sin_encabezado_no_son_titulos()
    test_titulo_de_una_palabra_y_secciones_con_imagen()
    test_tabla_copiada_en_otra_hoja()
    print("\nInforme ejecutivo test completado sin errores.")
