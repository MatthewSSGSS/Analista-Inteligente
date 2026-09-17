"""Leer imágenes pegadas en el Excel con OCR local y convertirlas en tablas.

Se dibujan con Pillow una tabla y un gráfico con la forma de las capturas
reales (ranking de jefes por mes; productividad con dos líneas y sus
valores escritos) y se exige:

- que cada número caiga en su mes por POSICIÓN: una celda vacía queda vacía
  y no corre los valores;
- que la fila Total no sea una persona y la columna Total no sea un mes;
- que un nombre sin espacios se restaure con los nombres del libro;
- que en el gráfico cada valor vaya a la serie de su color;
- que las imágenes del .xlsx se encuentren con su hoja y su fila.

Las pruebas que necesitan el OCR se saltan si RapidOCR no está instalado.

PYTHONPATH=. python tests/imagen_ocr_test.py
"""
import io
import warnings

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont

import core.imagen_ocr as ocr

warnings.filterwarnings("ignore")


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _fuente(tam):
    for nombre in ("arialbd.ttf", "arial.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(nombre, tam)
        except OSError:
            continue
    return ImageFont.load_default(size=tam)


def _png(img):
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


MESES = ["Oct", "Nov", "Dic", "Ene", "Feb", "Mar"]
JEFES = {
    "CASTILLO NAAR JACK": [77, 81, 83, 123, 113, 122],
    "ROJAS IBAÑEZ DIANA": [101, 128, 126, 115, 103, 108],
    "JIMENEZ MESTRA DABEIS": [None, None, 99, 87, 96, 90],
    "LINERO DIAZ CARLOS": [59, 75, 63, 57, 62, 59],
}


def _tabla_png():
    f = _fuente(18)
    img = Image.new("RGB", (1100, 300), "white")
    d = ImageDraw.Draw(img)
    x_mes = [420 + 95 * i for i in range(6)]
    d.text((150, 10), "Año", font=f, fill="black")
    d.text((500, 10), "2025", font=f, fill="black")
    d.text((780, 10), "2026", font=f, fill="black")
    d.text((1000, 10), "Total", font=f, fill="black")
    d.text((150, 45), "Jefe", font=f, fill="black")
    for x, m in zip(x_mes, MESES):
        d.text((x, 45), m, font=f, fill="black")
    for fila, (jefe, valores) in enumerate(JEFES.items()):
        y = 90 + 38 * fila
        # El OCR suele perder los espacios: se dibuja sin ellos a propósito.
        d.text((10, y), jefe.replace(" ", ""), font=f, fill="black")
        for x, v in zip(x_mes, valores):
            if v is not None:
                d.text((x - 5, y), f"{v} %", font=f, fill="black")
        validos = [v for v in valores if v is not None]
        d.text((1000, y), f"{round(sum(validos) / len(validos))} %", font=f, fill="black")
    d.text((10, 90 + 38 * 4), "Total", font=f, fill="black")
    for x in x_mes:
        d.text((x - 5, 90 + 38 * 4), "90 %", font=f, fill="black")
    return _png(img)


PRODUCTIVIDAD = [7.7, 7.6, 8.9, 9.5, 11.3, 13.8, 15.9, 18.8]
ASESORES = [568, 592, 643, 557, 446, 345, 283, 255]


def _grafico_png():
    f, fp = _fuente(16), _fuente(22)
    img = Image.new("RGB", (1000, 460), "white")
    d = ImageDraw.Draw(img)
    rojo, gris = (228, 60, 80), (60, 60, 72)
    d.text((380, 8), "PRODUCTIVIDAD MES A MES", font=fp, fill="black")
    d.ellipse((330, 52, 344, 66), fill=gris)
    d.text((350, 49), "Productividad", font=f, fill="black")
    d.ellipse((520, 52, 534, 66), fill=rojo)
    d.text((540, 49), "Asesores", font=f, fill="black")
    xs = [80 + 115 * i for i in range(8)]
    # Asesores arriba y productividad abajo, sin cruzarse: una etiqueta dibujada
    # encima de la otra línea es ambigua hasta para una persona.
    y_prod = [360 - (v - 7) * 8 for v in PRODUCTIVIDAD]
    y_ase = [110 + (650 - v) * 0.3 for v in ASESORES]
    d.line(list(zip(xs, y_ase)), fill=rojo, width=4)
    d.line(list(zip(xs, y_prod)), fill=gris, width=4)
    for x, yp, ya, p, a in zip(xs, y_prod, y_ase, PRODUCTIVIDAD, ASESORES):
        d.text((x - 12, yp - 28), f"{p}".replace(".", ","), font=f, fill="black")
        d.text((x - 14, ya - 28), str(a), font=f, fill="black")
    for i, x in enumerate(xs):
        d.text((x - 12, 390), f"{(i + 5) % 12 + 1:02d}.", font=f, fill="black")
    d.text((xs[1], 420), "2025", font=f, fill="black")
    d.text((xs[7], 420), "2026", font=f, fill="black")
    return _png(img)


def test_numeros_y_nombres_sin_ocr():
    check("porcentaje con espacio", ocr._numero("77 %") == 77)
    check("coma decimal", ocr._numero("7,7") == 7.7)
    check("miles con punto", ocr._numero("1.234") == 1234)
    check("texto no es número", ocr._numero("Oct") is None)
    vocab = ocr.vocabulario_de([pd.DataFrame({"JEFE": ["ROJAS IBAÑEZ DIANA CAROLINA", "CASTILLO NAAR JACK"]})])
    check("un nombre sin espacios ni Ñ se restaura", ocr.restaurar_nombre("ROJASIBANEZDIANACAROLINA", vocab)
          == "ROJAS IBAÑEZ DIANA CAROLINA")
    check("un nombre que no está en el libro se deja como se leyó", ocr.restaurar_nombre("OTRAPERSONA", vocab) == "OTRAPERSONA")
    check("los años siguen a los meses", ocr._asignar_anios([10, 11, 12, 1, 2], [2025], None) == [2025, 2025, 2025, 2026, 2026])


def test_imagenes_del_xlsx():
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    wb = Workbook()
    ws = wb.active
    ws.title = "INDICADORES"
    ws["B2"] = "RANKING DE LOS JEFES"
    ws.add_image(XLImage(io.BytesIO(_tabla_png())), "B4")
    buf = io.BytesIO()
    wb.save(buf)
    imagenes = ocr.imagenes_del_libro(buf.getvalue())
    check("se encuentra la imagen con su hoja y su fila", len(imagenes) == 1
          and imagenes[0]["hoja"] == "INDICADORES" and imagenes[0]["fila"] == 3)
    check("con sus bytes", imagenes[0]["bytes"][:4] == b"\x89PNG")

    from core.loader import load_workbook

    class Subida:
        name = "informe 2026.xlsx"

        def getvalue(self):
            return buf.getvalue()

    libro = load_workbook(Subida())
    check("el libro guarda la imagen con el título de su sección",
          libro["imagenes"] and libro["imagenes"][0]["seccion"] == "RANKING DE LOS JEFES")


def test_tabla_en_imagen():
    if not ocr.disponible():
        print("SKIP tabla en imagen: RapidOCR no está instalado")
        return
    vocab = ocr.vocabulario_de([pd.DataFrame({"JEFE": list(JEFES)})])
    r = ocr.interpretar_imagen(_tabla_png(), vocab)
    check("se reconoce una tabla", r["tipo"] == "tabla")
    a = r["ancha"].set_index(r["dimension"])
    check("una fila por jefe, sin la fila Total", sorted(a.index) == sorted(JEFES))
    check("con los nombres restaurados", "ROJAS IBAÑEZ DIANA" in a.index)
    meses = [c for c in a.columns if c in r["periodos"]]
    check("seis meses con su año", meses == ["oct 2025", "nov 2025", "dic 2025", "ene 2026", "feb 2026", "mar 2026"])
    check("la columna Total no es un mes", "Total" in a.columns and "Total" not in r["periodos"])
    for jefe, valores in JEFES.items():
        leidos = [None if pd.isna(v) else int(v) for v in a.loc[jefe, meses]]
        check(f"los valores de {jefe} caen en su mes", leidos == valores)
    check("las celdas vacías se avisan", any("vacías" in x for x in r["avisos"]))
    largo = ocr.a_tabla_larga(r, r["ancha"], "Cumplimiento %")
    check("la tabla larga tiene una fila por jefe y mes con dato", len(largo) == 22
          and largo["Mes"].min() == pd.Timestamp(2025, 10, 1))


def test_grafico_en_imagen():
    if not ocr.disponible():
        print("SKIP gráfico en imagen: RapidOCR no está instalado")
        return
    r = ocr.interpretar_imagen(_grafico_png())
    check("se reconoce un gráfico", r["tipo"] == "grafico")
    check("con sus dos series por la leyenda", sorted(r["series"]) == ["Asesores", "Productividad"])
    a = r["ancha"]
    check("ocho meses, de jun 2025 a ene 2026", list(a["Mes"]) == ["jun 2025", "jul 2025", "ago 2025", "sep 2025",
                                                                     "oct 2025", "nov 2025", "dic 2025", "ene 2026"])
    check("cada valor de productividad en su serie", list(a["Productividad"]) == PRODUCTIVIDAD)
    check("cada valor de asesores en su serie", list(a["Asesores"]) == ASESORES)
    largo = ocr.a_tabla_larga(r, a)
    check("el mes queda como fecha", isinstance(largo["Mes"].iloc[0], pd.Timestamp))


if __name__ == "__main__":
    test_numeros_y_nombres_sin_ocr()
    test_imagenes_del_xlsx()
    test_tabla_en_imagen()
    test_grafico_en_imagen()
    print("\nImagen OCR test completado sin errores.")
