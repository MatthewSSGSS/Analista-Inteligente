"""Regresión de los mensajes cuando un archivo no se puede analizar.

Un archivo que no es ni tabla plana ni dinámica reconocible tiene que
terminar en un mensaje que diga qué revisar. Los dos finales malos son:

- **silencioso**: el archivo "carga" y el panel muestra indicadores de una
  tabla que no existe (tres líneas de texto suelto se leían como una columna
  con dos registros);
- **confuso**: sale el error crudo de pandas, en inglés y en su jerga
  ("No columns to parse from file", "you must specify an engine manually"),
  que le pide al usuario final algo que no puede hacer.

Se corre igual que tests/smoke_test.py: PYTHONPATH=. python tests/errores_carga_test.py
"""
import io

from openpyxl import Workbook

from core.loader import load_workbook


def _upload(name, data):
    return type("Upload", (), {"getvalue": lambda self, d=data: d, "name": name})()


def _xlsx(filas, titulo="Hoja1"):
    wb = Workbook()
    ws = wb.active
    ws.title = titulo
    for f in filas:
        ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _error_de(nombre, data) -> str:
    try:
        load_workbook(_upload(nombre, data))
        return ""
    except ValueError as e:
        return str(e)


# ── Lo que antes pasaba en silencio: una hoja de texto suelto no es una
# tabla, y se leía como una columna con dos registros. ──
def texto_suelto_no_pasa_por_tabla():
    for etiqueta, nombre, data in [
        ("Excel", "nota.xlsx", _xlsx([["Estimado equipo:"], ["Adjunto el informe."], ["Saludos."]])),
        ("CSV", "nota.csv", "Estimado equipo:\nAdjunto el informe.\nSaludos.\n".encode("utf-8")),
    ]:
        error = _error_de(nombre, data)
        check(f"{etiqueta} con solo texto: ya no carga en silencio", bool(error))
        check(f"{etiqueta}: el mensaje dice que se abrió pero no había tabla",
              "no se reconoció ninguna tabla" in error)
        check(f"{etiqueta}: y dice qué revisar", "Revisa el archivo" in error)


# ── Lo que antes salía en inglés, crudo de pandas. ──
def errores_de_pandas_traducidos():
    vacio = _error_de("vacio.csv", b"")
    check("un CSV ilegible da un mensaje en español", "No se pudo leer el CSV" in vacio)
    check("y conserva el detalle técnico para diagnosticar", "detalle:" in vacio)

    roto = _error_de("roto.xlsx", b"esto no es un excel")
    check("un Excel dañado dice que puede estar dañado o mal nombrado",
          "puede estar dañado" in roto and "extensión" in roto)
    check("y ya no le pide al usuario que elija un motor de lectura",
          "specify an engine" not in roto.split("detalle:")[0])

    pdf = _error_de("informe.xlsx", b"%PDF-1.4\n%contenido falso")
    check("un PDF renombrado a .xlsx cae en el mismo mensaje claro", "puede estar dañado" in pdf)


def formato_no_soportado_dice_cuales_si():
    error = _error_de("datos.txt", b"cualquier cosa")
    check("un formato que no se admite lo dice con el nombre del archivo", "datos.txt" in error)
    check("y enumera los que sí sirven", ".xlsx" in error and ".csv" in error)


def archivo_vacio_se_distingue_de_archivo_sin_tabla():
    vacio = _error_de("vacio.xlsx", _xlsx([]))
    check("un Excel vacío dice que no hay datos", "no tiene datos que leer" in vacio)
    check("y no dice que encontró contenido sin forma de tabla",
          "no se reconoció ninguna tabla" not in vacio)


# ── Controles: lo que SÍ tiene que seguir cargando. ──
def control_lista_de_una_columna_sigue_cargando():
    filas = [["Nombre"]] + [[n] for n in ["Ana", "Luis", "Sara", "Juan", "Marta", "Pedro"]]
    libro = load_workbook(_upload("lista.xlsx", _xlsx(filas, "Lista")))
    df = libro["sheets"]["Lista"]["processed"]
    check("una lista de una sola columna con varias filas sigue siendo una tabla",
          list(df.columns) == ["Nombre"] and len(df) == 6)


def control_hoja_sin_tabla_no_tumba_el_resto_del_libro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    for f in [["Ciudad", "Ventas"], ["Bogotá", 10], ["Cali", 20]]:
        ws.append(f)
    portada = wb.create_sheet("Portada")
    portada.append(["Estimado equipo:"])
    portada.append(["Adjunto el informe."])
    b = io.BytesIO()
    wb.save(b)
    libro = load_workbook(_upload("mixto.xlsx", b.getvalue()))
    check("la hoja con datos se analiza igual", list(libro["sheets"]) == ["Ventas"])
    check("y se avisa de la que quedó fuera, sin tumbar la carga",
          any("Sin tabla reconocible" in a and "Portada" in a for a in libro["avisos"]))


def control_libro_que_solo_trae_imagenes_sigue_cargando():
    """Una hoja que es casi toda una imagen pegada (un ranking pegado como
    captura) no tiene tabla, pero el archivo SÍ es utilizable: sus cifras se
    leen con el OCR de «Imágenes del archivo». Si esto tronara, la
    herramienta que existe justo para ese caso no se abriría nunca."""
    from openpyxl.drawing.image import Image as XLImage
    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06"
           b"\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01\x00\x00\x05"
           b"\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82")
    wb = Workbook()
    ws = wb.active
    ws.title = "INDICADORES"
    ws["B2"] = "RANKING DE LOS JEFES"
    ws.add_image(XLImage(io.BytesIO(png)), "B4")
    b = io.BytesIO()
    wb.save(b)
    libro = load_workbook(_upload("ranking.xlsx", b.getvalue()))
    check("un libro cuyos datos están en imágenes carga igual", bool(libro["sheets"]))
    check("y se avisa de dónde están esas cifras",
          any("Imágenes del archivo" in a for a in libro["avisos"]))


if __name__ == "__main__":
    texto_suelto_no_pasa_por_tabla()
    errores_de_pandas_traducidos()
    formato_no_soportado_dice_cuales_si()
    archivo_vacio_se_distingue_de_archivo_sin_tabla()
    control_lista_de_una_columna_sigue_cargando()
    control_hoja_sin_tabla_no_tumba_el_resto_del_libro()
    control_libro_que_solo_trae_imagenes_sigue_cargando()
    print("\nErrores de carga test completado sin errores.")
