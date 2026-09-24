"""Hojas tipo tablero: bloques apilados con cabecera de dos niveles y sin meses.

Se arma un Excel con la forma del informe comercial real (tabla dinámica con
segmentadores):

- una fila de grupos combinados (Prepago / Pospago / Accesos) sobre una
  cabecera cuyos nombres se repiten bajo cada grupo (Ppto, Act, Pry, % Cump);
- la fila Total ENTRE los grupos y la cabecera;
- tres bloques apilados con la misma estructura y otra dimensión en la
  primera columna (Canal, Jefe, Territorio);
- una entidad con casi todo vacío (RETAIL) y celdas con guion;
- columnas al final sin grupo (Met Dig@) y un cumplimiento guardado como fracción.

Antes, el bloque de Jefe y el de Territorio quedaban metidos como filas de datos
del de Canal, y las columnas se llamaban "Prepago", "Unnamed"… sin nada que
dijera cuál era Ppto y cuál Act.

PYTHONPATH=. python tests/informe_bloques_test.py
"""
import io

import pandas as pd
from openpyxl import Workbook

from core.cuadro_comparativo import cuadro_comparativo
from core.dashboard_engine import build_dashboard
from core.loader import load_workbook


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


class _Subida:
    def __init__(self, nombre, libro):
        buf = io.BytesIO()
        libro.save(buf)
        self.name, self._datos = nombre, buf.getvalue()

    def getvalue(self):
        return self._datos


MEDIDAS = ["Ppto", "Act", "Pry", "% Cump"]
ACCESOS = ["Ppto Dig", "Dig", "Pry Dig", "% C Dig", "Ppto Ins", "Ins", "Pry Ins", "% C Ins"]
COLAS = ["Met Dig@", "Met Ins@"]

CANALES = {
    "CAV": [3111, 2155, 3079, 0.99, 9744, 5946, 8649, 0.888, 1515, 934, 1349, 0.891, 2000, 1700, 2100, 0.85],
    "DIGITAL": [None, 2, 3, None, 4489, 2501, 3638, 0.81, 1200, 800, 1100, 0.9, 1000, 900, 1200, 0.95],
    "RETAIL": [None, None, 1, None, None, None, None, None, None, None, None, None, None, None, None, None],
    "TAT": [147842, 74078, 105826, 0.716, "-", "-", "-", None, 500, 400, 450, 0.88, 400, 300, 380, 0.79],
    "AGENTES PDV": [7433, 2701, 3859, 0.519, 5966, 3847, 5596, 0.938, 1207, 749, 1082, 0.896, 2654, 1455, 2102, 0.792],
}


def _bloque(ws, fila, dimension, entidades, ocultar_fila=False, combinar=True):
    """Grupos combinados en `fila`, Total en fila+1, cabecera en fila+3, datos debajo."""
    grupos = [("Prepago", 3, 6), ("Pospago", 7, 10), ("Accesos", 11, 18)]
    for nombre, c1, c2 in grupos:
        ws.cell(fila, c1, nombre)
        if combinar:  # si no, el rótulo queda solo en su primera celda ("centrar en la selección")
            ws.merge_cells(start_row=fila, start_column=c1, end_row=fila, end_column=c2)
    # Indicadores sueltos al borde de la fila de grupos, como en el informe
    # real: no son datos del bloque, pero conviven con los rótulos.
    for i, v in enumerate((1.15, 26, 0.93)):
        ws.cell(fila, 21 + i, v)
    ws.cell(fila + 1, 1, "Total")
    for c in range(3, 19):
        ws.cell(fila + 1, c, 1000)  # el Total no debe entrar en los datos
    ws.row_dimensions[fila + 2].hidden = ocultar_fila
    ws.cell(fila + 3, 1, dimension)
    cabecera = MEDIDAS + MEDIDAS + ACCESOS[:4]
    for i, nombre in enumerate(cabecera):
        ws.cell(fila + 3, 3 + i, nombre)
    for i, nombre in enumerate(ACCESOS[4:]):
        ws.cell(fila + 3, 15 + i, nombre)
    for i, nombre in enumerate(COLAS):
        ws.cell(fila + 3, 19 + i, nombre)
    for k, (entidad, valores) in enumerate(entidades.items()):
        ws.cell(fila + 4 + k, 1, entidad)
        for i, v in enumerate(valores):
            if v is not None:
                ws.cell(fila + 4 + k, 3 + i, v)
        ws.cell(fila + 4 + k, 19, 10 + k)
        ws.cell(fila + 4 + k, 20, 20 + k)
    return fila + 4 + len(entidades)


def _libro_sin_combinar():
    wb = Workbook()
    ws = wb.active
    ws.title = "Tablero"
    _bloque(ws, 1, "Canal", CANALES, combinar=False)
    return wb


def _libro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Tablero"
    fin = _bloque(ws, 1, "Canal", CANALES)
    jefes = {f"Jefe {i}": [0, 10 * i, 12 * i, 0.5 + i / 20] * 3 + [100, 80 * i, 90 * i, 0.9 - i / 50] * 2
             for i in range(1, 6)}
    fin = _bloque(ws, fin + 3, "Jefe", jefes, ocultar_fila=True)
    _bloque(ws, fin + 3, "Territorio",
            {"Atlantico": [None] * 4 + [None] * 4 + [1575, 1148, 1658, 1.053] + [2626, 1796, 2594, 0.988] + [1] * 4,
             "Bolivar": [None] * 8 + [952, 637, 920, 0.967] + [1754, 1048, 1514, 0.863] + [2] * 4,
             "Cesar": [None] * 8 + [483, 244, 352, 0.73] + [891, 468, 676, 0.759] + [3] * 4})

    normal = wb.create_sheet("Normal")
    for j, nombre in enumerate(["Ciudad", "Ventas", "Costo"]):
        normal.cell(1, 1 + j, nombre)
    for i in range(1, 8):
        normal.cell(1 + i, 1, f"C{i}")
        normal.cell(1 + i, 2, i * 10)
        normal.cell(1 + i, 3, i * 4)
    meta = wb.create_sheet("Metas")
    for j, nombre in enumerate(["Region", "Meta", "Ventas", "% Cumplimiento"]):
        meta.cell(1, 1 + j, nombre)
    for i, (region, m, v) in enumerate([("Norte", 100, 90), ("Sur", 200, 210), ("Este", 150, 120), ("Oeste", 80, 80)]):
        meta.cell(2 + i, 1, region)
        meta.cell(2 + i, 2, m)
        meta.cell(2 + i, 3, v)
        meta.cell(2 + i, 4, v / m)  # fracción: 0,9 = 90 %
    return wb


def test_cada_bloque_es_una_tabla_por_su_dimension():
    libro = load_workbook(_Subida("tablero.xlsx", _libro()))
    hojas = libro["sheets"]
    check("un bloque por dimensión, más la hoja normal",
          {"Tablero · por Canal", "Tablero · por Jefe", "Tablero · por Territorio", "Normal"} <= set(hojas))
    check("el bloque de Jefe no queda metido como filas del de Canal",
          not any(str(v).startswith("Jefe") for v in hojas["Tablero · por Canal"]["processed"]["Canal"]))

    canal = hojas["Tablero · por Canal"]["processed"]
    check("una fila por canal y sin la fila Total", list(canal["Canal"]) == list(CANALES))
    check("las columnas dicen su grupo y su medida",
          {"Prepago · Presupuesto", "Prepago · Actual", "Prepago · Proyección", "Prepago · % Cumplimiento",
           "Pospago · Actual", "Accesos · Presupuesto Dig", "Accesos · % Cumplimiento Ins"} <= set(canal.columns))
    check("las columnas sin grupo conservan su nombre", {"Met Dig@", "Met Ins@"} <= set(canal.columns))
    check("ningún nombre de columna genérico", not any(str(c).startswith(("Unnamed", "Columna")) for c in canal.columns))
    cav = canal.set_index("Canal").loc["CAV"]
    check("el cumplimiento guardado como fracción se lee como porcentaje", cav["Prepago · % Cumplimiento"] == 99)
    check("las cifras absolutas no se tocan", cav["Prepago · Presupuesto"] == 3111 and cav["Prepago · Actual"] == 2155)
    check("el Total no se cuela (habría 1.000 en el presupuesto)", 1000 not in set(canal["Prepago · Presupuesto"].dropna()))


def test_vacio_no_es_cero_y_el_guion_tampoco():
    canal = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]["Tablero · por Canal"]["processed"].set_index("Canal")
    check("un canal sin cifra queda vacío, no en cero", pd.isna(canal.loc["RETAIL", "Prepago · Actual"]))
    check("un guion es «sin dato», no cero", pd.isna(canal.loc["TAT", "Pospago · Actual"]))
    check("el promedio ignora los vacíos", round(canal["Prepago · Actual"].mean(), 1) == round((2155 + 2 + 74078 + 2701) / 4, 1))


def test_otros_bloques_y_fila_oculta():
    hojas = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]
    jefe = hojas["Tablero · por Jefe"]["processed"]
    check("el bloque de Jefe tiene sus 5 jefes", list(jefe["Jefe"]) == [f"Jefe {i}" for i in range(1, 6)])
    terr = hojas["Tablero · por Territorio"]["processed"]
    check("el bloque de Territorio tiene sus 3 territorios", list(terr["Territorio"]) == ["Atlantico", "Bolivar", "Cesar"])
    check("y no arrastra columnas vacías del grupo que no reportan (Prepago)",
          not any(c.startswith("Prepago") for c in terr.columns))


def test_se_puede_analizar():
    item = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]["Tablero · por Canal"]
    tablero = build_dashboard(item["processed"], item["profile"])
    check("el panel arma el tablero sobre el bloque", tablero is not None and tablero.get("kpis") is not None)
    schema = item["profile"]["schema"]
    check("«Canal» se reconoce como categoría", "Canal" in schema.get("categorical", []) or "Canal" in schema.get("text", []))
    check("los cumplimientos se reconocen como porcentaje",
          any("Cumplimiento" in c for c in schema.get("metrics", [])))


def test_la_metrica_principal_es_la_del_cumplimiento():
    item = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]["Tablero · por Canal"]
    tablero = build_dashboard(item["processed"], item["profile"])
    check("la métrica principal es la que el archivo usa para su % (Proyección), no el presupuesto",
          tablero["primary_metric"] == "Prepago · Proyección")
    check("y se deja dicho en el registro de carga",
          any("equivale a «Prepago · Proyección» ÷ «Prepago · Presupuesto»" in l for l in item["profile"]["cleaning_log"]))
    cuadro = cuadro_comparativo(item["processed"], item["profile"]["schema"], "Canal", tablero["primary_metric"])
    check("el cuadro compara contra el presupuesto del mismo grupo", cuadro["base"] == "meta"
          and cuadro["meta_col"] == "Prepago · Presupuesto")
    por_canal = {f["nombre"]: f["cumplimiento"] for f in cuadro["filas"]}
    check("el cumplimiento recalculado coincide con el «% Cump» de la hoja",
          all(abs(por_canal[c] - v) < 0.6 for c, v in {"CAV": 99.0, "TAT": 71.6, "AGENTES PDV": 51.9}.items()))
    check("la referencia es el grupo completo (5 canales), no los elegidos", cuadro["total_grupo"] == 5)
    hallazgos = " ".join(i["finding"] for i in tablero["insights"])
    check("ya no dice que el archivo no trae meta", "no trae una meta" not in hallazgos)


def test_sin_celdas_combinadas_tambien():
    canal = load_workbook(_Subida("tablero.xlsx", _libro_sin_combinar()))["sheets"]["Tablero"]["processed"]
    check("el rótulo del grupo llega a todas sus columnas aunque solo esté en la primera celda",
          {"Prepago · Presupuesto", "Prepago · % Cumplimiento", "Pospago · Presupuesto",
           "Pospago · % Cumplimiento", "Accesos · Presupuesto Dig"} <= set(canal.columns))
    check("sin nada mezclado entre grupos", not ({"Presupuesto", "Actual"} & set(canal.columns)))


def test_la_fila_de_grupos_se_reconoce_aunque_traiga_cifras_al_borde():
    """La fila de los grupos del informe real lleva indicadores sueltos a la
    derecha (1,15 · 26 · 0,93). Con la regla vieja —descartar la fila si traía
    UN número— los bloques de Canal y de Jefe salían sin el nombre de su grupo
    («Presupuesto_2» en vez de «Pospago · Presupuesto»), mientras los de más
    abajo, sin esas cifras al lado, sí lo traían."""
    canal = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]["Tablero · por Canal"]["processed"]
    check("el grupo llega a sus columnas pese a las cifras sueltas de esa fila",
          {"Prepago · Presupuesto", "Pospago · Presupuesto", "Accesos · Presupuesto Dig"} <= set(canal.columns))
    check("y no queda ninguna columna numerada por no saber su grupo",
          not any(str(c).endswith(("_2", "_3", "_4")) for c in canal.columns))
    check("la fila Total sigue sin colarse como si fuera el grupo",
          not any("Total" in str(c) for c in canal.columns))


def test_una_columna_estrechada_hasta_desaparecer_no_entra():
    """La otra forma de esconder una columna: dejarla en 0,14 caracteres de
    ancho sin marcarla como oculta. En el informe real era la columna B, con
    un «Act» de más que salía como una columna repetida."""
    from openpyxl.utils import get_column_letter

    wb = _libro()
    ws = wb["Tablero"]
    # Se inserta una columna invisible justo antes de las medidas del bloque.
    ws.insert_cols(2)
    ws.cell(4, 2, "Act")
    for k in range(len(CANALES)):
        ws.cell(5 + k, 2, 999)
    ws.column_dimensions[get_column_letter(2)].width = 0.14  # openpyxl marca customWidth solo
    canal = load_workbook(_Subida("tablero.xlsx", wb))["sheets"]["Tablero · por Canal"]["processed"]
    check("la columna estrechada hasta no verse queda fuera",
          999 not in set(pd.to_numeric(canal.select_dtypes("number").stack(), errors="coerce").dropna()))
    check("y no aparece una medida repetida por su culpa",
          not any(str(c).endswith("_2") for c in canal.columns))


def test_una_hoja_normal_no_se_toca():
    normal = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]["Normal"]["processed"]
    check("la hoja normal sigue igual", list(normal.columns) == ["Ciudad", "Ventas", "Costo"] and len(normal) == 7)
    normal = load_workbook(_Subida("tablero.xlsx", _libro()))["sheets"]["Metas"]
    check("una hoja con Meta y Ventas sigue su camino: sin cambios de forma",
          list(normal["processed"].columns) == ["Region", "Meta", "Ventas", "% Cumplimiento"])
    check("Ventas (lo que el % mide) pasa a ser la principal, no la Meta",
          normal["profile"]["schema"]["metrics"][0] == "Ventas")


if __name__ == "__main__":
    test_cada_bloque_es_una_tabla_por_su_dimension()
    test_vacio_no_es_cero_y_el_guion_tampoco()
    test_otros_bloques_y_fila_oculta()
    test_se_puede_analizar()
    test_la_metrica_principal_es_la_del_cumplimiento()
    test_sin_celdas_combinadas_tambien()
    test_la_fila_de_grupos_se_reconoce_aunque_traiga_cifras_al_borde()
    test_una_columna_estrechada_hasta_desaparecer_no_entra()
    test_una_hoja_normal_no_se_toca()
    print("\nInforme de bloques test completado sin errores.")
