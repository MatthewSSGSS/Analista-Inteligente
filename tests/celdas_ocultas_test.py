"""El andamiaje oculto de un tablero no se lee como si fueran datos.

Un tablero armado sobre tablas dinámicas trae debajo sus rótulos de trabajo
—«Periodo», «Etiquetas de columna», «Suma de Activaciones»— y columnas de un
grupo que ya no se reporta. Quien lo arma los esconde en Excel en vez de
borrarlos. En la hoja «Dash» del archivo real son 20 filas y 25 columnas: leídas
como datos, los rótulos salían como categorías y las columnas como «Columna 7».

Se comprueba que lo oculto se descarta SOLO en la ruta de tablero, que una hoja
de datos se sigue leyendo entera (ahí un filtro esconde filas y perderlas sería
perder registros), y que si al quitar lo oculto no se reconoce nada, se
reintenta con la hoja completa.

PYTHONPATH=. python tests/celdas_ocultas_test.py
"""
import io

import pandas as pd
from openpyxl import Workbook

from core.celdas_ocultas import filas_y_columnas_ocultas, sin_lo_oculto
from core.loader import load_workbook


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


class _Subida:
    def __init__(self, nombre, libro):
        buf = io.BytesIO()
        libro.save(buf)
        self.name, self._datos = nombre, buf.getvalue()

    def getvalue(self):
        return self._datos


MEDIDAS = ["Ppto", "Act", "Pry", "% Cump"]
CANALES = {"CAV": [3111, 2155, 3079, 0.99], "DIGITAL": [1500, 1200, 1400, 0.93],
           "TAT": [147842, 74078, 105826, 0.716], "AGENTES PDV": [7433, 2701, 3859, 0.519]}


def _tablero():
    """Como el archivo real: la fila «Periodo» y la de «Suma de…» ocultas, y
    un grupo entero («Power») escondido entre Prepago y Pospago."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dash"
    # Fila 1: los grupos. Power (columnas 7-10) está oculto.
    for nombre, c1, c2 in (("Prepago", 3, 6), ("Power", 7, 10), ("Pospago", 11, 14)):
        ws.cell(1, c1, nombre)
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
    ws.cell(2, 1, "Periodo")           # andamiaje: se oculta
    ws.cell(2, 2, "Mes Act")
    ws.cell(3, 1, "Etiquetas de fila")  # andamiaje: se oculta
    for c in range(3, 15):
        ws.cell(3, c, "Suma de Activaciones")
    ws.cell(4, 1, "Canal")
    for i, nombre in enumerate(MEDIDAS * 3):
        ws.cell(4, 3 + i, nombre)
    for k, (canal, valores) in enumerate(CANALES.items()):
        ws.cell(5 + k, 1, canal)
        for i, v in enumerate(valores * 3):
            ws.cell(5 + k, 3 + i, v)
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[3].hidden = True
    for letra in ("G", "H", "I", "J"):          # el grupo Power
        ws.column_dimensions[letra].hidden = True
    # La otra forma de esconder una columna: estrecharla hasta que no se vea.
    ws.cell(4, 15, "Sobra")
    for k in range(len(CANALES)):
        ws.cell(5 + k, 15, 999)
    ws.column_dimensions["O"].width = 0.14

    datos = wb.create_sheet("Datos")
    for j, nombre in enumerate(["Ciudad", "Ventas", "Costo"]):
        datos.cell(1, 1 + j, nombre)
    for i in range(1, 9):
        datos.cell(1 + i, 1, f"C{i}")
        datos.cell(1 + i, 2, i * 10)
        datos.cell(1 + i, 3, i * 4)
    datos.row_dimensions[4].hidden = True       # como si un filtro la escondiera
    return wb


def test_se_leen_del_archivo_las_filas_y_columnas_ocultas():
    datos = _Subida("tablero.xlsx", _tablero()).getvalue()
    mapa = filas_y_columnas_ocultas(datos, "tablero.xlsx")
    filas, columnas = mapa["Dash"]
    check("las filas ocultas salen contando desde 0", filas == {1, 2})
    check("y las columnas del grupo escondido también", {6, 7, 8, 9} <= columnas)
    check("una columna estrechada hasta no verse cuenta como oculta, aunque Excel no la marque",
          14 in columnas)
    check("la hoja de datos trae su fila oculta aparte", mapa["Datos"][0] == {3})
    check("un formato que no lo dice (CSV) no rompe nada", filas_y_columnas_ocultas(b"a,b\n1,2\n", "x.csv") == {})
    check("un archivo dañado tampoco", filas_y_columnas_ocultas(b"no soy excel", "roto.xlsx") == {})


def test_quitar_lo_oculto_reindexa_y_se_protege():
    grid = pd.DataFrame([[1, 2, 3], [4, 5, 6], [7, 8, 9]])
    visible = sin_lo_oculto(grid, {1}, {0})
    check("queda solo lo visible, reindexado desde 0",
          visible.shape == (2, 2) and list(visible.columns) == [0, 1] and visible.iat[1, 0] == 8)
    check("sin nada oculto devuelve la misma cuadrícula", sin_lo_oculto(grid, set(), set()) is grid)
    check("si se quedara sin filas, no se toca", sin_lo_oculto(grid, {0, 1, 2}, set()) is grid)
    check("si se quedara sin columnas, tampoco", sin_lo_oculto(grid, set(), {0, 1, 2}) is grid)


def test_el_tablero_se_lee_sin_su_andamiaje():
    libro = load_workbook(_Subida("tablero.xlsx", _tablero()))
    hoja = next(n for n in libro["sheets"] if n.startswith("Dash"))
    df = libro["sheets"][hoja]["processed"]
    check("el bloque sale por su dimensión real, no por «Etiquetas de fila»", df.columns[0] == "Canal")
    check("una fila por canal", list(df["Canal"]) == list(CANALES))
    check("los rótulos de la tabla dinámica no entran como datos",
          not any("Suma de" in str(v) for v in df.values.ravel()))
    check("ni «Periodo» ni «Mes Act» se cuelan", "Mes Act" not in set(df.astype(str).values.ravel()))
    check("ninguna columna genérica", not any(str(c).startswith(("Unnamed", "Columna")) for c in df.columns))
    check("el grupo oculto queda fuera: 2 grupos × 4 medidas, no 3",
          len([c for c in df.columns if c != "Canal"]) == 8)
    check("y la columna estrechada hasta no verse tampoco entra", "Sobra" not in df.columns)
    cav = df.set_index("Canal").loc["CAV"]
    check("y las cifras son las que se ven en Excel", cav.iloc[0] == 3111 and cav.iloc[1] == 2155)
    check("se deja dicho en el registro de carga",
          any("oculto" in l for l in libro["sheets"][hoja]["profile"]["cleaning_log"]))


def test_una_hoja_de_datos_se_lee_entera():
    libro = load_workbook(_Subida("tablero.xlsx", _tablero()))
    datos = libro["sheets"]["Datos"]["processed"]
    check("la fila escondida por un filtro NO se descarta: sería perder registros",
          len(datos) == 8 and list(datos.columns) == ["Ciudad", "Ventas", "Costo"])
    check("y sigue estando la que Excel ocultaba", "C3" in set(datos["Ciudad"]))


def test_si_sin_lo_oculto_no_se_reconoce_nada_se_usa_la_hoja_completa():
    """Una hoja cuyo tablero vive justo en las filas ocultas: quitarlas dejaría
    la hoja sin nada, así que se reintenta con todo y se lee como antes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dash"
    for nombre, c1, c2 in (("Prepago", 3, 6), ("Pospago", 7, 10)):
        ws.cell(1, c1, nombre)
        ws.merge_cells(start_row=1, start_column=c1, end_row=1, end_column=c2)
    ws.cell(2, 1, "Canal")
    for i, nombre in enumerate(MEDIDAS * 2):
        ws.cell(2, 3 + i, nombre)
    for k, (canal, valores) in enumerate(CANALES.items()):
        ws.cell(3 + k, 1, canal)
        for i, v in enumerate(valores * 2):
            ws.cell(3 + k, 3 + i, v)
        ws.row_dimensions[3 + k].hidden = True   # TODOS los datos, ocultos
    libro = load_workbook(_Subida("raro.xlsx", wb))
    hoja = next(n for n in libro["sheets"] if n.startswith("Dash"))
    df = libro["sheets"][hoja]["processed"]
    check("se reintenta con la hoja completa y los canales siguen ahí",
          set(CANALES) <= set(df.iloc[:, 0].astype(str)))


if __name__ == "__main__":
    test_se_leen_del_archivo_las_filas_y_columnas_ocultas()
    test_quitar_lo_oculto_reindexa_y_se_protege()
    test_el_tablero_se_lee_sin_su_andamiaje()
    test_una_hoja_de_datos_se_lee_entera()
    test_si_sin_lo_oculto_no_se_reconoce_nada_se_usa_la_hoja_completa()
    print("\nCeldas ocultas test completado sin errores.")
