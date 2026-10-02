"""El Excel de Exportar es un informe para enviar, no un volcado de datos.

Antes el botón «Excel» traía tres hojas sin formato (datos, estadística y
hallazgos). Aquí se exige que el libro traiga la portada con indicadores y
gráficos, el plan de acción, la tendencia, «cómo va cada uno» con la
referencia sobre el grupo completo, tablas dinámicas reales y los datos como
tabla de Excel; que se adapte cuando el archivo no da para una sección (sin
fechas, tabla ancha de meses, una sola fila) y que un texto que empieza con
«=» no se vuelva fórmula.

PYTHONPATH=. python tests/excel_informe_test.py
"""
import io
import logging
import warnings

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from core.dashboard_engine import build_dashboard
from core.profile import profile_sheet
from ui.report_excel import build_excel_report

warnings.filterwarnings("ignore")
logging.disable(logging.WARNING)


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _libro(df0, nombre="Ventas"):
    item = profile_sheet(df0, {"sheet_name": nombre, "workbook_name": f"{nombre}.xlsx"})
    df, profile = item["processed"], item["profile"]
    dashboard = build_dashboard(df, profile)
    contenido = build_excel_report(df, profile["schema"], dashboard, f"{nombre}.xlsx", nombre, "Sin filtros aplicados")
    return df, load_workbook(io.BytesIO(contenido))


def _ventas():
    rng = np.random.default_rng(7)
    filas = []
    for mes in pd.date_range("2026-01-01", periods=6, freq="MS"):
        for region in ["Norte", "Sur", "Centro", "Oriente"]:
            for canal in ["Tienda", "Online"]:
                filas.append({"Fecha": mes + pd.Timedelta(days=int(rng.integers(0, 25))), "Región": region,
                              "Canal": canal, "Ventas": int(rng.integers(800, 3000)) * 1000, "Meta": 2_000_000,
                              "Nota": "=SUMA(A1:A9)" if region == "Norte" and mes.month == 1 else "ok"})
    return pd.DataFrame(filas)


def test_libro_completo_con_fechas():
    df, wb = _libro(_ventas())
    nombres = wb.sheetnames
    check("la portada va primero", nombres[0] == "Resumen")
    for hoja in ("Plan de acción", "Tendencia", "Por Región", "Matriz", "Datos"):
        check(f"trae la hoja «{hoja}»", hoja in nombres)
    check("la portada trae gráficos", len(wb["Resumen"]._charts) >= 2)
    check("la tendencia trae gráficos", len(wb["Tendencia"]._charts) >= 1)
    pivotes = [ws for ws in wb.worksheets if ws._pivots]
    check("trae tablas dinámicas reales de Excel", len(pivotes) >= 1)
    check("las dinámicas se recalculan al abrir", all(p.cache.refreshOnLoad for ws in pivotes for p in ws._pivots))
    datos = wb["Datos"]
    check("los datos van como tabla de Excel", len(datos.tables) == 1)
    check("todas las filas están en Datos", datos.max_row == len(df) + 1)
    check("se agrega el periodo por mes para las dinámicas", datos.cell(row=1, column=datos.max_column).value == "Periodo (mes)")
    # La línea de inmovilizar partía los gráficos y escondía filas en las hojas de informe.
    check("en Datos queda fijo solo el encabezado", datos.freeze_panes == "A2")
    check("las hojas de informe no inmovilizan paneles",
          all(wb[n].freeze_panes is None for n in nombres if n != "Datos" and not wb[n]._pivots))


def test_referencia_sobre_el_grupo_completo():
    _, wb = _libro(_ventas())
    ws = wb["Por Región"]
    encabezados = [c.value for c in next(r for r in ws.iter_rows() if any(v == "Posición" for v in (x.value for x in r)))]
    check("la participación dice sobre cuántos se calcula", any(str(h).startswith("Participación (de 4)") for h in encabezados if h))
    check("el promedio de referencia es el de los 4", any(h == "vs promedio de los 4" for h in encabezados if h))


def test_texto_con_igual_no_es_formula():
    _, wb = _libro(_ventas())
    ws = wb["Datos"]
    col = [c.value for c in ws[1]].index("Nota") + 1
    celdas = [ws.cell(row=r, column=col) for r in range(2, ws.max_row + 1)]
    sospechosas = [c for c in celdas if isinstance(c.value, str) and c.value.startswith("=")]
    check("el texto «=SUMA(...)» se conserva", len(sospechosas) > 0)
    check("y queda como texto, no como fórmula", all(c.data_type == "s" for c in sospechosas))


def test_sin_fechas_tabla_ancha_y_una_fila():
    rng = np.random.default_rng(3)
    _, wb = _libro(pd.DataFrame({"Producto": [f"P{i}" for i in range(9)], "Categoría": ["A", "B", "C"] * 3,
                                 "Stock": rng.integers(0, 50, 9)}), "Inventario")
    check("sin fechas no hay hoja de tendencia", "Tendencia" not in wb.sheetnames)
    check("pero sí el ranking por dimensión", any(n.startswith("Por ") for n in wb.sheetnames))

    ancha = pd.DataFrame({"Región": ["Antioquia", "Cundinamarca", "Valle"],
                          **{m: [100 + i * 10, 90 + i * 5, 80 + i * 7]
                             for i, m in enumerate(["Enero", "Febrero", "Marzo", "Abril"])}})
    _, wb = _libro(ancha, "Resumen 2026")
    check("una tabla ancha de meses también tiene tendencia", "Tendencia" in wb.sheetnames)
    check("y sus datos por mes para la dinámica", "Datos por mes" in wb.sheetnames)
    t = wb["Tendencia"]
    valores = [t.cell(row=r, column=3).value for r in range(1, t.max_row + 1)]
    check("la tendencia suma cada mes (270 en enero)", 270 in valores)

    _, wb = _libro(pd.DataFrame({"Nombre": ["X"], "Valor": [3]}), "Mini")
    check("con una sola fila el libro igual se arma", wb.sheetnames[0] == "Resumen" and "Datos" in wb.sheetnames)


def test_vacios_de_pandas():
    """Columnas Int64/string con `pd.NA` (como llegan de un informe tipo tabla):
    el libro se arma y esas celdas quedan vacías, no con el texto «<NA>»."""
    from core.dashboard_engine import build_dashboard
    item = profile_sheet(_ventas(), {"sheet_name": "Ventas", "workbook_name": "v.xlsx"})
    df = item["processed"].copy()
    df["Ventas"] = df["Ventas"].astype("Int64")
    df.loc[[1, 4], "Ventas"] = pd.NA
    df["Región"] = df["Región"].astype("string")
    df.loc[3, "Región"] = pd.NA
    contenido = build_excel_report(df, item["profile"]["schema"], build_dashboard(df, item["profile"]),
                                   "v.xlsx", "Ventas", "")
    ws = load_workbook(io.BytesIO(contenido))["Datos"]
    textos = [c.value for fila in ws.iter_rows() for c in fila if c.value == "<NA>"]
    check("con pd.NA el Excel se arma y no escribe «<NA>»", not textos)


def main():
    test_vacios_de_pandas()
    test_libro_completo_con_fechas()
    test_referencia_sobre_el_grupo_completo()
    test_texto_con_igual_no_es_formula()
    test_sin_fechas_tabla_ancha_y_una_fila()


if __name__ == "__main__":
    main()
