"""Regresión: números escritos como texto, con su formato, se leen como números.

Un Excel armado a mano trae "$ 1.234.567", "(500.000)", "2.500.000,50",
"1,200,000" o "45%" escritos como texto. La limpieza solo convertía lo que
pd.to_numeric entendía, así que la columna «Valor» entera quedaba como texto:
el panel no tenía métrica, ni KPIs, aunque el nombre fuera perfecto. Y lo que
sí convertía lo leía con la convención inglesa: "800.000" eran 800.

Lo delicado, y lo que fija este archivo:
- el punto y la coma se deciden por COLUMNA cuando el valor solo no lo dice
  ("1.500" junto a "1.5" es decimal; junto a "1.234.567", de miles);
- en una misma columna pueden convivir "2.500.000,50" y "1,200,000";
- una celda que en Excel ya era un número no se reinterpreta;
- fechas, teléfonos y códigos NO se vuelven números.

Se corre igual que tests/smoke_test.py: PYTHONPATH=. python tests/numeros_formato_test.py
"""
import io

import pandas as pd
from openpyxl import Workbook

from core.loader import load_workbook
from core.numeric import numeros_con_formato


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _leer(valores):
    r = numeros_con_formato(pd.Series(valores, dtype=object))
    return None if r is None else (r[0].tolist(), r[1])


def cada_formato_se_lee_bien():
    check("moneda, paréntesis, coma decimal y estilo inglés en la misma columna",
          _leer(["$ 1.234.567", "(500.000)", "2.500.000,50", "800000", "1,200,000"])[0]
          == [1234567, -500000, 2500000.5, 800000, 1200000])
    check("porcentajes en escala 0-100, con su formato",
          _leer(["45%", "12,5%", "30%", "20 %"]) == ([45, 12.5, 30, 20], "porcentaje"))
    check("«800.000» son ochocientos mil, no 800", _leer(["800.000", "1.500", "2.300"])[0] == [800000, 1500, 2300])
    check("con un «1.5» en la columna, el punto es decimal", _leer(["1.5", "2.25", "3.234"])[0] == [1.5, 2.25, 3.234])
    check("con un «0.125», también", _leer(["0.125", "1.500", "2.750"])[0] == [0.125, 1.5, 2.75])
    check("el menos al final (contabilidad) es negativo", _leer(["1.234-", "-2.000", "3.000"])[0] == [-1234, -2000, 3000])
    check("moneda antes o después del número", _leer(["1.234 COP", "US$ 2.000", "€ 15,50"]) == ([1234, 2000, 15.5], "moneda"))
    check("una celda que ya era número no se reinterpreta",
          _leer([1234567.0, "1.234.567", "2.000.000"])[0] == [1234567, 1234567, 2000000])


def lo_que_no_es_numero_no_se_toca():
    check("fechas", _leer(["01/02/2025", "03/02/2025", "04/02/2025"]) is None)
    check("teléfonos con espacios", _leer(["300 123 4567", "310 555 1234", "320 111 2222"]) is None)
    check("códigos", _leer(["A-102", "B-203", "C-301"]) is None)
    check("enteros planos siguen por el camino de siempre", _leer(["100", "200", "300"]) is None)
    check("una columna de texto con un número suelto", _leer(["Bogotá", "Cali", "12", "Medellín"]) is None)


def el_archivo_completo_tiene_metricas():
    wb = Workbook(); ws = wb.active; ws.title = "Ventas"
    ws.append(["Fecha", "Vendedor", "Valor", "Col8"])
    filas = [("2025-02-01", "Ana", "$ 1.234.567", "45%"), ("2025-02-03", "Luis", "(500.000)", "12,5%"),
             ("2025-03-02", "Ana", "2.500.000,50", "30%"), ("2025-03-04", "Juan", 800000, "20%"),
             ("2025-04-05", "Luis", "1,200,000", "18%"), ("2025-04-06", "Juan", "$ 950.000", "22%")]
    for f in filas:
        ws.append(list(f))
    b = io.BytesIO(); wb.save(b)

    class _Subida:
        name = "ventas.xlsx"
        def getvalue(self):
            return b.getvalue()

    hoja = load_workbook(_Subida())["sheets"]["Ventas"]
    df, schema = hoja["processed"], hoja["profile"]["schema"]
    check("«Valor» es una métrica de dinero", "Valor" in schema["metrics"] and schema["types"]["Valor"] == "Moneda")
    check("y suma lo que debe", abs(float(df["Valor"].sum()) - 6184567.5) < 0.01)
    check("«Col8» (escrita con %) es porcentaje aunque el nombre no lo diga",
          schema["types"]["Col8"] == "Porcentaje" and "Col8" in schema.get("por_valores", {}))
    check("el registro de limpieza lo cuenta",
          any("escritos como texto" in x for x in hoja["profile"]["cleaning_log"]))


if __name__ == "__main__":
    cada_formato_se_lee_bien()
    lo_que_no_es_numero_no_se_toca()
    el_archivo_completo_tiene_metricas()
    print("\nNúmeros con formato test completado sin errores.")
