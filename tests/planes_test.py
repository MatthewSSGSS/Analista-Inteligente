"""Que el análisis termine en algo que alguien pueda ejecutar.

Dos pestañas se prueban aquí:

- **Planes de mejora**: cada hallazgo tiene que salir convertido en pasos
  concretos que mencionen a los casos por su nombre. Un plan que dice
  "revisar los segmentos afectados" no es un plan, es el hallazgo otra vez.
- **Análisis de seguimiento**: las observaciones de UN caso tienen que estar
  comparadas contra algo —el grupo, la meta o su propio pasado—. "Vendió
  3.2M" no dice si va bien.

PYTHONPATH=. python tests/planes_test.py
"""
import pandas as pd

from core.dashboard_engine import build_dashboard
from core.planes import generar
from core.profile import profile_sheet
from core.seguimiento import observaciones
from ui.person_profile import has_entity, resolve_entity


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


def _archivo(con_problemas=True):
    """Seis meses por punto. Con problemas: RIOHACHA se hunde y MAICAO se va."""
    puntos = ("RIOHACHA", "CIENAGA", "SOLEDAD", "CARTAGENA", "MONTERIA", "MAICAO")
    filas = []
    for mes in range(1, 7):
        for punto in puntos:
            base = 900
            if con_problemas and punto == "RIOHACHA":
                base = max(0, 1200 - mes * 180)
            if con_problemas and punto == "MAICAO" and mes >= 5:
                base = 0
            for i in range(6):
                filas.append({"Fecha": f"2026-{mes:02d}-{(i * 4) + 2:02d}", "PUNTO": punto,
                              "Ventas": int(base * (0.9 + i * 0.04)), "Meta": 1000})
    item = profile_sheet(pd.DataFrame(filas), {"sheet_name": "V", "workbook_name": "v.xlsx"})
    return item["processed"], item["profile"]["schema"], build_dashboard(item["processed"], item["profile"])


def _texto_de(plan):
    return " ".join([plan["situacion"], plan["por_que"], plan["medir"]] + plan["pasos"])


def test_los_planes_nombran_a_los_casos():
    df, schema, dashboard = _archivo()
    resultado = generar(df, schema, dashboard)
    planes = resultado["planes"]
    check("el análisis produce planes", bool(planes))
    check("hay al menos un frente crítico", resultado["criticos"] >= 1)
    check("el resumen dice por dónde empezar", "atención" in resultado["resumen"].lower()
          or "crítico" in resultado["resumen"].lower())

    caida = next((p for p in planes if p["titulo"].startswith("Dónde se concentra")), None)
    check("la caída se convierte en plan", caida is not None)
    check("y el plan menciona al punto que cae", "RIOHACHA" in _texto_de(caida))
    check("con pasos en orden, no una sola frase", len(caida["pasos"]) >= 3)
    check("y con una forma de saber si funcionó", bool(caida["medir"]))
    check("lo crítico va primero", planes[0]["estado"] == "critico")

    # Ningún paso puede quedarse en la plantilla sin rellenar.
    for plan in planes:
        check(f"el plan '{plan['titulo'][:28]}' no deja huecos sin rellenar",
              "{nombres}" not in _texto_de(plan) and "{primero}" not in _texto_de(plan))


def test_si_todo_va_bien_propone_mejoras():
    """Una pestaña vacía es la forma más rápida de que nadie la vuelva a abrir."""
    filas = [{"Fecha": f"2026-{m:02d}-1{i}", "Zona": z, "Ventas": 400 + m * 20, "Meta": 300}
             for m in range(1, 7) for z in ("N", "S", "E", "O") for i in range(5)]
    item = profile_sheet(pd.DataFrame(filas), {"sheet_name": "V", "workbook_name": "v.xlsx"})
    df, schema = item["processed"], item["profile"]["schema"]
    resultado = generar(df, schema, build_dashboard(df, item["profile"]))
    check("sin problemas no se inventa ninguno", resultado["criticos"] == 0)
    check("pero igual se propone algo", bool(resultado["planes"]))
    check("y todo lo propuesto es una mejora",
          all(p["estado"] == "mejora" for p in resultado["planes"]))
    check("el resumen lo dice sin alarmar", "rojo" in resultado["resumen"])


def test_el_seguimiento_compara_contra_algo():
    df, schema, _ = _archivo()
    check("el archivo de puntos habilita el seguimiento", has_entity(df, schema))
    entidad = resolve_entity(df, schema)
    check("y se sigue por el punto de venta", entidad["column"] == "PUNTO")

    hallazgos = observaciones(df, schema, "PUNTO", "RIOHACHA", "Ventas")
    check("el seguimiento produce observaciones", bool(hallazgos))
    titulos = " ".join(h["title"] for h in hallazgos)

    posicion = next((h for h in hallazgos if h["title"].startswith("Posición")), None)
    check("dice en qué puesto está", posicion is not None)
    check("contra cuántos comparables", "de 6" in posicion["title"])
    check("y contra qué se le está midiendo", "meta" in posicion["implication"].lower())
    check("con la mediana del grupo como referencia",
          any("Mediana" in str(e["nombre"]) for e in posicion["evidence"]))
    check("y nombrando al mejor del grupo",
          any("CARTAGENA" in str(e["nombre"]) for e in posicion["evidence"]))

    check("detecta que viene cayendo", "cayendo" in titulos.lower())
    check("y dice cuánto pesa en el total", "peso" in titulos.lower())

    # Un caso sano no debe recibir las mismas alertas que uno que se hunde.
    sano = observaciones(df, schema, "PUNTO", "CARTAGENA", "Ventas")
    check("un caso sano no se marca como caído",
          not any("cayendo" in h["title"].lower() for h in sano))
    check("y su posición se reporta como favorable",
          any(h["kind"] == "positive" for h in sano))


def test_no_rompe_sin_datos():
    vacio = pd.DataFrame({"A": [], "B": []})
    check("sin filas no truena", observaciones(vacio, {}, "A", "x") == [])
    check("con una columna inexistente tampoco",
          observaciones(pd.DataFrame({"A": [1, 2]}), {}, "NoExiste", "x") == [])


def test_el_plan_se_puede_ejecutar():
    """Un plan sin dueño ni fecha no se ejecuta, y uno que se queda en la app
    se pierde al cerrar la sesión."""
    import io as _io
    from datetime import date, timedelta

    import openpyxl

    from ui.planes import (_clave_plan, alerta_plan, estado_plan, excel_plan_de_accion,
                           filas_plan_de_accion, resumen_para_compartir)

    df, schema, dashboard = _archivo()
    planes = generar(df, schema, dashboard)["planes"]
    check("hay al menos tres planes para probar", len(planes) >= 3)
    hoy = date(2026, 9, 13)
    primero, segundo = planes[0], planes[1]
    c1, c2 = _clave_plan(primero), _clave_plan(segundo)
    asignaciones = {c1: {"responsable": "Ana Gómez", "fecha": hoy - timedelta(days=2)},
                    c2: {"responsable": "", "fecha": hoy + timedelta(days=5)}}
    pasos = {c1: [True] + [False] * (len(primero["pasos"]) - 1),
             c2: [True] * len(segundo["pasos"])}

    check("sin pasos marcados, el plan está pendiente", estado_plan(0, 4) == "Pendiente")
    check("con algunos, está en curso", estado_plan(2, 4) == "En curso")
    check("con todos, está hecho", estado_plan(4, 4) == "Hecho")
    check("una fecha pasada con pasos pendientes es un vencido",
          alerta_plan({"responsable": "X", "fecha": hoy - timedelta(days=1)}, 1, 3, hoy) == "Vencido")

    filas = filas_plan_de_accion(planes, asignaciones, pasos, hoy)
    check("hay una fila por plan", len(filas) == len(planes))
    check("el plan con fecha vencida se señala", filas[0]["Alerta"] == "Vencido")
    check("el responsable asignado llega a la fila", filas[0]["Responsable"] == "Ana Gómez")
    check("uno terminado no tiene alerta aunque le falte responsable",
          filas[1]["Estado"] == "Hecho" and filas[1]["Alerta"] == "")
    check("uno sin asignar avisa que no tiene responsable", filas[2]["Alerta"] == "Sin responsable")

    libro = openpyxl.load_workbook(_io.BytesIO(excel_plan_de_accion(planes, asignaciones, pasos, hoy)))
    check("el Excel trae una hoja de frentes y otra de pasos", libro.sheetnames == ["Plan de acción", "Pasos"])
    hoja = libro["Plan de acción"]
    encabezados = [c.value for c in hoja[1]]
    check("con responsable, fecha y cómo se mide",
          {"Responsable", "Fecha compromiso", "Cómo se mide"} <= set(encabezados))
    check("una fila por plan", hoja.max_row == len(planes) + 1)
    check("y el nombre del responsable llega al archivo", any(c.value == "Ana Gómez" for c in hoja[2]))
    total_pasos = sum(len(p["pasos"]) for p in planes)
    check("la hoja de pasos trae cada paso", libro["Pasos"].max_row == total_pasos + 1)

    texto = resumen_para_compartir(planes, asignaciones, pasos, hoy)
    check("el resumen para compartir nombra al responsable", "Ana Gómez" in texto)
    check("marca lo vencido", "Vencido" in texto)
    check("y avisa lo que no tiene dueño", "sin asignar" in texto)


def test_el_tablero_muestra_cada_plan_donde_va():
    """El tablero reemplaza a la tabla como vista principal: cada plan en la
    columna de su urgencia, con su avance y a quién afecta."""
    from datetime import date

    import streamlit as st

    from ui.planes import _clave_plan, _hero, _tablero, filas_plan_de_accion

    df, schema, dashboard = _archivo()
    resultado = generar(df, schema, dashboard)
    planes = resultado["planes"]
    hoy = date(2026, 9, 14)
    pasos = {_clave_plan(planes[0]): [True, True] + [False] * (len(planes[0]["pasos"]) - 2)}
    filas = filas_plan_de_accion(planes, {}, pasos, hoy)

    capturado = []
    original = st.markdown
    st.markdown = lambda html, **k: (capturado.append(str(html)), original(html, **k))[1]
    try:
        _hero(resultado, planes, filas)
        _tablero(planes, filas)
    finally:
        st.markdown = original
    html = "".join(capturado)
    total_pasos = sum(len(p["pasos"]) for p in planes)
    avance = round(2 / total_pasos * 100)
    check("el anillo muestra el avance real del plan", f"--p:{avance};" in html)
    check("y cuántos pasos lleva", f"2 de {total_pasos} pasos" in html)
    check("cada plan tiene su tarjeta en el tablero", html.count('class="tarjeta-titulo"') == len(planes))
    for estado, etiqueta in (("critico", "Crítico"), ("atencion", "En observación"), ("mejora", "Oportunidad")):
        n = sum(1 for p in planes if p["estado"] == estado)
        check(f"la columna {etiqueta} cuenta sus planes", f"{etiqueta}</span><b>{n}</b>" in html)
    check("sin responsable asignado, la tarjeta lo avisa", "Sin responsable" in html)


if __name__ == "__main__":
    test_los_planes_nombran_a_los_casos()
    test_si_todo_va_bien_propone_mejoras()
    test_el_seguimiento_compara_contra_algo()
    test_no_rompe_sin_datos()
    test_el_plan_se_puede_ejecutar()
    test_el_tablero_muestra_cada_plan_donde_va()
    print("\nPlanes test completado sin errores.")
