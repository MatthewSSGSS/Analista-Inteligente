"""Regresión: encabezados repetidos y notas al pie no son registros.

El archivo de la prueba es el «Excel desordenado» con el que se encontró el
problema: título arriba, dos bloques pegados con su fila de encabezados
repetida, notas al pie. Antes, la fila repetida metía la palabra «Valor»
dentro de la columna Valor y la dejaba entera como texto (sin KPIs), y las
notas al pie se contaban como registros. Ver core/filas_basura.py.

También fija lo que NO se quita: una fila con cifras, un registro corto con
los números en blanco al final de la tabla, una celda que por casualidad se
llama igual que su columna.

Se corre igual que tests/smoke_test.py: PYTHONPATH=. python tests/filas_basura_test.py
"""
import io

import pandas as pd
from openpyxl import Workbook

from core.filas_basura import quitar_filas_basura
from core.loader import load_workbook


def check(label, condition):
    if not condition:
        raise AssertionError(label)
    print("OK  ", label)


class _Subida:
    def __init__(self, nombre, datos):
        self.name, self._datos = nombre, datos

    def getvalue(self):
        return self._datos


def _excel(filas):
    wb = Workbook(); ws = wb.active; ws.title = "Ventas"
    for f in filas:
        ws.append(f)
    b = io.BytesIO(); wb.save(b)
    return _Subida("ventas.xlsx", b.getvalue())


ENCABEZADO = ["Fecha", "Vendedor", "Cédula", "Ciudad", "Valor", "Margen"]


def el_excel_desordenado_se_lee_completo():
    hoja = load_workbook(_excel([
        ["INFORME DE VENTAS"], [],
        ENCABEZADO,
        ["01/02/2025", "Juan Pérez", "1.234.567.890", "Bogotá", "$ 1.234.567", "45%"],
        ["2025-02-03", "Juan Pérez", "1234567890", "Bogotá", "(500.000)", "12,5%"],
        ["VENTAS SEGUNDA QUINCENA"],
        ENCABEZADO,
        ["16/02/2025", "Ana Gómez", "52.111.222", "Medellín", "2.500.000,50", "30%"],
        ["17/02/2025", "Luis Díaz", "79.000.111", "Cali", "1,200,000", "20%"],
        [],
        ["Fuente: sistema comercial. Cifras en pesos."],
        ["Elaboró: área comercial"],
    ]))["sheets"]["Ventas"]
    df, schema = hoja["processed"], hoja["profile"]["schema"]
    check("quedan los 4 registros, nada más", len(df) == 4)
    check("sin la palabra «Valor» ni el título del bloque entre los datos",
          not df.astype(str).isin(["Valor", "VENTAS SEGUNDA QUINCENA", "Fecha"]).any().any())
    check("«Valor» es dinero y suma lo que debe",
          schema["types"]["Valor"] == "Moneda" and abs(float(df["Valor"].sum()) - 4434567.5) < 0.01)
    check("«Margen» es porcentaje", schema["types"]["Margen"] == "Porcentaje")
    check("«Fecha» es fecha", schema["dates"] == ["Fecha"])
    check("la misma cédula con y sin puntos es un solo valor", df["Cédula"].nunique() == 3)
    log = " ".join(hoja["profile"]["cleaning_log"])
    check("el registro de limpieza cuenta el encabezado repetido", "repetían el encabezado" in log)
    check("y las notas al pie", "nota(s) al pie" in log)


def el_csv_tambien():
    texto = ("Vendedor;Zona;Valor\nAna;Norte;100\nLuis;Sur;200\nVendedor;Zona;Valor\n"
             "Juan;Norte;300\nNota: datos preliminares del mes\n")
    hoja = load_workbook(_Subida("v.csv", texto.encode("utf-8")))["sheets"]["CSV"]
    df = hoja["processed"]
    check("CSV: encabezado repetido y nota fuera", len(df) == 3 and float(df["Valor"].sum()) == 600)


def lo_que_parece_basura_pero_es_dato():
    cols = ["Vendedor", "Zona", "Valor", "Estado"]
    df = pd.DataFrame([["Ana", "Norte", 100, "Estado"],       # una celda igual a su columna, nada más
                       ["Luis", "Sur", 200, "Activo"],
                       ["Juan", None, None, None]],           # registro real con las cifras en blanco
                      columns=cols)
    limpio, log = quitar_filas_basura(df)
    check("una celda que se llama como su columna no hace encabezado", len(limpio) == 3)
    check("un registro corto al final no es una nota", limpio.iloc[-1]["Vendedor"] == "Juan" and not log)
    df = pd.DataFrame([["Ana", "Norte", 100, "Activo"],
                       ["Total sucursal centro norte del mes", None, 300, None]], columns=cols)
    check("una fila con cifras nunca se quita aunque tenga una frase", len(quitar_filas_basura(df)[0]) == 2)


if __name__ == "__main__":
    el_excel_desordenado_se_lee_completo()
    el_csv_tambien()
    lo_que_parece_basura_pero_es_dato()
    print("\nFilas basura test completado sin errores.")
