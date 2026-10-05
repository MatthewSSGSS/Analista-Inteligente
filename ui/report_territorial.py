"""Informes descargables del Análisis Territorial: Excel ejecutivo y HTML.

Los dos salen del mismo diccionario, `core/territorio.informe`, que es lo
que dibuja el tablero del semáforo en la pantalla: el archivo dice
exactamente lo mismo que la app, sin cálculos propios.

- **Excel** (`build_territorial_excel`): «Resumen» (indicadores, semáforo,
  por dónde y por qué, lectura, gráficos), «Semáforo» (todas las zonas con
  su estado y su razón, como tabla de Excel con filtros), «Departamentos»,
  «Mes a mes» (zona × mes con escala de color y gráfico del total), «Dónde
  crecer» y «Notas» (cómo se calcula cada cosa).
- **HTML** (`build_territorial_html`): una página que se abre en cualquier
  navegador y se puede reenviar o imprimir a PDF, con el mapa interactivo
  incrustado (el fondo del mapa necesita internet), el semáforo, las
  razones, la tabla completa con búsqueda y orden, y la evolución.

Sin Streamlit: reciben datos y devuelven bytes.
"""
from __future__ import annotations

import html
import io
import re
from datetime import datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from core import territorio as T
from core import territorio_plan as P

# Tinta y grises del tema claro; el semáforo con los mismos tonos que el mapa.
TINTA, MUTED, LINEA, ZEBRA, PANEL, TEAL = "131826", "5B6473", "D8DCE6", "F6F7F9", "F1F3F7", "0F8A85"
SEM = {"subio": ("15803D", "DCFCE7", "Subió"), "estable": ("A16207", "FEF9C3", "Estable"),
       "bajo": ("B91C1C", "FEE2E2", "Bajó")}
SEM_GRAFICO = {"subio": "22C55E", "estable": "EAB308", "bajo": "EF4444"}
FMT_NUM, FMT_PCT, FMT_PCT_SIGNO, FMT_SIGNO = "#,##0", "0.0%", "+0.0%;-0.0%;0.0%", "+#,##0;-#,##0;0"

_fino = Side(style="thin", color=LINEA)
_CAJA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


def _plano(texto) -> str:
    """La frase sin las marcas de negrita (**)."""
    return str(texto or "").replace("**", "")


def _negritas_html(texto) -> str:
    partes = html.escape(str(texto or "")).split("**")
    return "".join(f"<b>{p}</b>" if i % 2 else p for i, p in enumerate(partes))


def _num(v):
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if np.isfinite(f) else None


def _var_txt(v, td) -> str:
    if _num(v) is not None:
        return f"{float(v):+.0%}"
    return "nuevo" if td == 1 else "sin actividad" if td == -1 else ""


def _zona_txt(nivel: str) -> str:
    return "departamentos" if nivel == "departamento" else "municipios" if nivel == "municipio" else "zonas"


# ── Excel ─────────────────────────────────────────────────────────────────

def _encabezado(ws, titulo: str, subtitulo: str, ancho: int = 12) -> int:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ancho)
    c = ws.cell(1, 1, titulo)
    c.font = Font(size=18, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=TINTA)
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ancho)
    s = ws.cell(2, 1, subtitulo)
    s.font = Font(size=10, color=MUTED, italic=True)
    s.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 30
    for col in range(1, ancho + 1):
        ws.cell(1, col).fill = PatternFill("solid", fgColor=TINTA)
    return 4


def _seccion(ws, fila: int, texto: str, ancho: int = 12) -> int:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    c = ws.cell(fila, 1, texto.upper())
    c.font = Font(size=10.5, bold=True, color=TEAL)
    c.border = Border(bottom=Side(style="medium", color=TEAL))
    for col in range(1, ancho + 1):
        ws.cell(fila, col).border = Border(bottom=Side(style="medium", color=TEAL))
    ws.row_dimensions[fila].height = 22
    return fila + 1


def _parrafo(ws, fila: int, texto: str, ancho: int = 12, negrita: bool = False, color: str = TINTA,
             alto: float | None = None) -> int:
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ancho)
    c = ws.cell(fila, 1, _plano(texto))
    c.data_type = "s"
    c.font = Font(size=11, bold=negrita, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.row_dimensions[fila].height = alto or max(18, 15 * (len(_plano(texto)) // 150 + 1))
    return fila + 1


def _kpi(ws, fila: int, col: int, titulo: str, valor, formato: str | None, detalle: str, color: str = TINTA) -> None:
    """Una tarjeta de 2 columnas × 3 filas."""
    ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + 1)
    ws.merge_cells(start_row=fila + 1, start_column=col, end_row=fila + 1, end_column=col + 1)
    ws.merge_cells(start_row=fila + 2, start_column=col, end_row=fila + 2, end_column=col + 1)
    t = ws.cell(fila, col, titulo.upper())
    t.font = Font(size=8.5, bold=True, color=MUTED)
    v = ws.cell(fila + 1, col, valor)
    v.font = Font(size=18, bold=True, color=color)
    if formato:
        v.number_format = formato
    d = ws.cell(fila + 2, col, detalle)
    d.font = Font(size=9, color=MUTED)
    d.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(fila, fila + 3):
        for c in (col, col + 1):
            celda = ws.cell(r, c)
            celda.fill = PatternFill("solid", fgColor=PANEL)
            celda.alignment = Alignment(wrap_text=True, vertical="center" if r < fila + 2 else "top", horizontal="left", indent=1)
    ws.cell(fila, col).border = Border(top=Side(style="thick", color=color if color != TINTA else TEAL))
    ws.cell(fila, col + 1).border = Border(top=Side(style="thick", color=color if color != TINTA else TEAL))


def _tabla(ws, fila: int, encabezados: list, filas: list, formatos: dict | None = None, nombre: str | None = None,
           anchos: dict | None = None) -> tuple[int, int]:
    """Encabezado + filas; si hay `nombre`, una tabla de Excel con filtros. Devuelve (primera, última) fila de datos."""
    formatos = formatos or {}
    for j, e in enumerate(encabezados, start=1):
        c = ws.cell(fila, j, e)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=TINTA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 30
    for i, valores in enumerate(filas, start=1):
        for j, v in enumerate(valores, start=1):
            c = ws.cell(fila + i, j, v)
            if isinstance(v, str):
                c.data_type = "s"
            if j in formatos:
                c.number_format = formatos[j]
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 30)
            if not nombre and i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
    if nombre and filas:
        ref = f"A{fila}:{get_column_letter(len(encabezados))}{fila + len(filas)}"
        tabla = Table(displayName=nombre, ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(tabla)
    for col, ancho in (anchos or {}).items():
        ws.column_dimensions[get_column_letter(col)].width = ancho
    return fila + 1, fila + len(filas)


def _pintar_estado(celda, estado) -> None:
    if estado in SEM:
        texto, fondo, rotulo = SEM[estado]
        celda.value = rotulo
        celda.font = Font(bold=True, color=texto)
        celda.fill = PatternFill("solid", fgColor=fondo)
        celda.alignment = Alignment(horizontal="center", vertical="top")


def _hoja_resumen(wb, inf: dict, ctx: dict, serie_total: pd.DataFrame, plan: dict | None = None) -> int:
    ws = wb.active
    ws.title = "Resumen"
    for col in range(1, 13):
        ws.column_dimensions[get_column_letter(col)].width = 13
    fila = _encabezado(ws, f"Informe territorial · {inf['metrica']}",
                       f"{ctx['archivo']} · hoja {ctx['hoja']} · {ctx['registros']:,} registros{ctx['filtros']}"
                       f" · generado el {ctx['generado']}" + (f"\nSemáforo: {inf['comparacion']}" if inf["comparacion"] else ""))
    ct, sem, cob = inf["cambio_total"], inf["semaforo"], inf["cobertura"]
    _kpi(ws, fila, 1, ("Total" if inf["sumable"] else "Promedio") + f" · {inf['metrica']}", _num(inf["total"]), FMT_NUM,
         "todo el periodo, zonas ubicadas")
    if ct:
        color = SEM["subio"][0] if ct["delta"] >= 0 else SEM["bajo"][0]
        _kpi(ws, fila, 3, "Cambio vs mes anterior", ct["pct"], FMT_PCT_SIGNO,
             f"{T.cifra_signo(ct['delta'])} · {T.cifra(ct['antes'])} → {T.cifra(ct['ahora'])}", color)
    _kpi(ws, fila, 5, "Zonas con actividad", inf["n"], FMT_NUM, _zona_txt(inf["nivel"]))
    if sem:
        _kpi(ws, fila, 7, "Subieron", sem["subio"]["n"], FMT_NUM, f"{T.cifra_signo(sem['subio']['cambio'])} entre todas",
             SEM["subio"][0])
        _kpi(ws, fila, 9, "Bajaron", sem["bajo"]["n"], FMT_NUM, f"{T.cifra_signo(sem['bajo']['cambio'])} entre todas",
             SEM["bajo"][0])
    if cob and cob.get("pct_deptos") is not None:
        _kpi(ws, fila, 11, "Cobertura de población", cob["pct_deptos"], "0%", "de la gente de tus departamentos")
    ws.row_dimensions[fila + 1].height = 30
    ws.row_dimensions[fila + 2].height = 30
    fila += 4

    if plan and plan.get("resumen"):
        fila = _seccion(ws, fila, "Plan de acción · lo primero de la semana")
        for frase in plan["resumen"]:
            fila = _parrafo(ws, fila, "• " + frase)
        for j in plan["jugadas"][:5]:
            valor = f"vale ≈ {T.cifra(j['valor'])} al mes" if j["valor"] else "buena práctica"
            fila = _parrafo(ws, fila, f"{j['n']}. {j['icono']} {j['titulo']} {j['zona']} ({j['departamento']}) · {valor}"
                            + (f" · responsable: {j['responsable']}" if j["responsable"] else "") + f" · {j['plazo']}",
                            negrita=True, color=SEM["bajo"][0] if j["tono"] == "bajo" else TINTA)
            if j["pasos"]:
                fila = _parrafo(ws, fila, "     → " + _plano(j["pasos"][0]), color=MUTED)
        c = ws.cell(fila, 1, "→ El plan completo, para hacerle seguimiento, está en la hoja «Plan de acción»")
        c.hyperlink = "#'Plan de acción'!A1"
        c.font = Font(color=TEAL, bold=True, underline="single")
        fila += 2

    if inf.get("por_donde") or inf.get("por_que"):
        fila = _seccion(ws, fila, "Qué pasó y por qué")
        if ct:
            fila = _parrafo(ws, fila, f"El total {'subió' if ct['delta'] >= 0 else 'bajó'} {T.cifra(abs(ct['delta']))}"
                            + (f" ({ct['pct']:+.1%})" if ct["pct"] is not None else "") + f": {inf['comparacion']}.",
                            negrita=True)
        if inf.get("por_donde"):
            fila = _parrafo(ws, fila, "📍 Por dónde: " + inf["por_donde"]["frase"])
        if inf.get("por_que"):
            fila = _parrafo(ws, fila, "🔎 Por qué: " + inf["por_que"]["frase"])
        fila += 1

    if inf["lectura"]:
        fila = _seccion(ws, fila, "Lectura del territorio")
        for frase in inf["lectura"]:
            fila = _parrafo(ws, fila, "• " + frase)
        fila += 1

    # Datos de apoyo para los gráficos (a la derecha, fuera de la vista).
    tabla = inf["tabla"]
    top = tabla.head(10)
    base = 16
    ws.cell(fila, base, "Zona").font = Font(color=MUTED, size=8)
    ws.cell(fila, base + 1, inf["metrica"]).font = Font(color=MUTED, size=8)
    for i, r in enumerate(top.itertuples(), start=1):
        ws.cell(fila + i, base, str(r.nombre))
        ws.cell(fila + i, base + 1, _num(r.valor))
    graficos = fila
    fila = _seccion(ws, fila, "Las 10 zonas más fuertes y el semáforo")
    if len(top):
        bar = BarChart()
        bar.type, bar.style, bar.title = "bar", 10, f"Top 10 · {inf['metrica']}"
        bar.add_data(Reference(ws, min_col=base + 1, min_row=graficos, max_row=graficos + len(top)), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=base, min_row=graficos + 1, max_row=graficos + len(top)))
        bar.y_axis.scaling.orientation = "minMax"
        bar.x_axis.scaling.orientation = "maxMin"   # la más fuerte arriba
        bar.y_axis.crosses = "max"                   # y el eje de valores abajo, no encima del título
        bar.y_axis.majorGridlines = None
        bar.legend = None
        bar.x_axis.delete = False
        bar.y_axis.delete = False
        bar.y_axis.numFmt = '#,##0'
        bar.series[0].graphicalProperties.solidFill = TEAL
        bar.series[0].invertIfNegative = False
        bar.width, bar.height = 17, 9
        bar.visible_cells_only = False   # los datos van en columnas ocultas
        ws.add_chart(bar, f"A{fila}")
    if sem:
        ws.cell(graficos, base + 3, "Estado").font = Font(color=MUTED, size=8)
        ws.cell(graficos, base + 4, "Zonas").font = Font(color=MUTED, size=8)
        for i, k in enumerate(("subio", "estable", "bajo"), start=1):
            ws.cell(graficos + i, base + 3, {"subio": "Subieron", "estable": "Se mantuvieron", "bajo": "Bajaron"}[k])
            ws.cell(graficos + i, base + 4, sem[k]["n"])
        dona = DoughnutChart()
        dona.title = "Semáforo"
        dona.add_data(Reference(ws, min_col=base + 4, min_row=graficos, max_row=graficos + 3), titles_from_data=True)
        dona.set_categories(Reference(ws, min_col=base + 3, min_row=graficos + 1, max_row=graficos + 3))
        for i, k in enumerate(("subio", "estable", "bajo")):
            punto = DataPoint(idx=i)
            punto.graphicalProperties.solidFill = SEM_GRAFICO[k]
            dona.series[0].dPt.append(punto)
        dona.dataLabels = DataLabelList()
        dona.dataLabels.showVal = True
        dona.dataLabels.showCatName = False
        dona.dataLabels.showSerName = False
        dona.dataLabels.showPercent = False
        dona.dataLabels.showLeaderLines = False
        dona.holeSize = 55
        dona.width, dona.height = 10, 9
        dona.visible_cells_only = False
        ws.add_chart(dona, f"H{fila}")
    fila += 20
    if serie_total is not None and len(serie_total) >= 2:
        ws.cell(graficos, base + 6, "Mes").font = Font(color=MUTED, size=8)
        ws.cell(graficos, base + 7, inf["metrica"]).font = Font(color=MUTED, size=8)
        for i, r in enumerate(serie_total.itertuples(), start=1):
            ws.cell(graficos + i, base + 6, T.etiqueta_mes(r.mes, True)
                    + (f" (al día {inf['corte_dia']})" if inf["corte_dia"] and r.mes == inf["mes_b"] else ""))
            ws.cell(graficos + i, base + 7, _num(r.valor))
        fila = _seccion(ws, fila, "Evolución del total")
        linea = LineChart()
        linea.title, linea.style = f"{inf['metrica']} mes a mes", 12
        linea.add_data(Reference(ws, min_col=base + 7, min_row=graficos, max_row=graficos + len(serie_total)),
                       titles_from_data=True)
        linea.set_categories(Reference(ws, min_col=base + 6, min_row=graficos + 1, max_row=graficos + len(serie_total)))
        linea.legend = None
        linea.x_axis.delete = False
        linea.y_axis.delete = False
        linea.y_axis.numFmt = '#,##0'
        linea.series[0].graphicalProperties.line.solidFill = TEAL
        linea.series[0].graphicalProperties.line.width = 28000
        linea.series[0].smooth = False
        linea.width, linea.height = 27, 8
        linea.visible_cells_only = False
        ws.add_chart(linea, f"A{fila}")
        fila += 17
    for col in range(base, base + 8):
        ws.column_dimensions[get_column_letter(col)].hidden = True

    return fila


_DESCRIPCION_HOJAS = {"Plan de acción": "qué hacer en cada zona, responsable, plazo, meta y estado para el seguimiento",
                      "Agentes": "el plan de cada agente: dónde perdió, clientes perdidos, ruta, meta y estrategia",
                      "Dónde abrir": "municipios sin presencia, desde dónde atenderlos y cuánto valen",
                      "Semáforo": "todas las zonas con su estado, su cambio y por qué se movieron",
                      "Departamentos": "cada departamento y qué municipios lo empujaron",
                      "Mes a mes": "cada zona por mes, con escala de color",
                      "Dónde crecer": "municipios grandes sin presencia y presentes por debajo de lo normal",
                      "Notas": "cómo se calcula cada cosa"}


def _indice(wb, fila: int) -> None:
    """Vínculos a las hojas que sí se armaron (cada una se omite si el archivo no da para ella)."""
    ws = wb["Resumen"]
    fila = _seccion(ws, fila, "Hojas de este libro")
    for nombre in wb.sheetnames[1:]:
        c = ws.cell(fila, 1, f"→ {nombre}")
        c.hyperlink = f"#'{nombre}'!A1"
        c.font = Font(color=TEAL, bold=True, underline="single")
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=12)
        ws.cell(fila, 3, _DESCRIPCION_HOJAS.get(nombre, "")).font = Font(color=MUTED)
        fila += 1


def _hoja_semaforo(wb, inf: dict) -> None:
    tabla = inf["tabla"]
    if tabla.empty:
        return
    ws = wb.create_sheet("Semáforo")
    mes_a = T.etiqueta_mes(inf["mes_a"], True) if inf["mes_a"] else ""
    mes_b = T.etiqueta_mes(inf["mes_b"], True) + (f" al día {inf['corte_dia']}" if inf["corte_dia"] else "") if inf["mes_b"] else ""
    corte_a = f" al día {inf['corte_dia']}" if inf["corte_dia"] else ""
    fila = _encabezado(ws, f"Semáforo por zona · {inf['metrica']}",
                       (f"{inf['comparacion']}. " if inf["comparacion"] else "")
                       + f"Posición y participación sobre las {inf['n']} zonas. Verde: subió más de 5% · amarillo: ±5% · "
                         "rojo: bajó más de 5%. Usa los filtros del encabezado.", ancho=14)
    con_mes = "estado" in tabla.columns
    encabezados = ["#", "Zona", "Departamento", f"{inf['metrica']} (periodo)", "Participación"]
    if con_mes:
        encabezados += [f"{mes_a}{corte_a}", mes_b, "Cambio", "Variación", "Estado", "Por qué"]
    extras = [(c, e) for c, e in (("cumplimiento", "Cumplimiento meta"), ("poblacion", "Población"),
                                  ("por_10k", "Por 10.000 hab.")) if c in tabla.columns and tabla[c].notna().any()]
    encabezados += [e for _, e in extras]
    filas = []
    for r in tabla.itertuples():
        f = [int(r.posicion), str(r.nombre), str(r.departamento), _num(r.valor), _num(getattr(r, "participacion", None))]
        if con_mes:
            razon = inf["razones"].get(str(r.zona), {}).get("frase_larga", "")
            f += [_num(r.mes_a), _num(r.mes_b), _num(r.cambio), _num(r.variacion), r.estado or "",
                  _plano(razon)]
        f += [_num(getattr(r, c)) for c, _ in extras]
        filas.append(f)
    formatos = {1: "0", 4: FMT_NUM, 5: FMT_PCT}
    anchos = {1: 6, 2: 26, 3: 20, 4: 16, 5: 13}
    if con_mes:
        formatos.update({6: FMT_NUM, 7: FMT_NUM, 8: FMT_SIGNO, 9: FMT_PCT_SIGNO})
        anchos.update({6: 15, 7: 15, 8: 13, 9: 12, 10: 11, 11: 60})
    for k, (c, _) in enumerate(extras, start=len(encabezados) - len(extras) + 1):
        formatos[k] = "0%" if c == "cumplimiento" else FMT_NUM if c == "poblacion" else "#,##0.0"
        anchos[k] = 14
    ini, fin = _tabla(ws, fila, encabezados, filas, formatos, nombre="TablaSemaforo", anchos=anchos)
    if con_mes:
        for i in range(ini, fin + 1):
            _pintar_estado(ws.cell(i, 10), ws.cell(i, 10).value)
            ws.cell(i, 11).alignment = Alignment(wrap_text=True, vertical="top")
        ws.conditional_formatting.add(f"I{ini}:I{fin}", ColorScaleRule(
            start_type="num", start_value=-0.3, start_color="F8B4B4", mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=0.3, end_color="A7F3D0"))
    ws.conditional_formatting.add(f"D{ini}:D{fin}", DataBarRule(start_type="min", end_type="max", color="7FC8C4"))


def _hoja_departamentos(wb, inf: dict) -> None:
    deps = inf["departamentos"]
    if not deps:
        return
    ws = wb.create_sheet("Departamentos")
    fila = _encabezado(ws, "Por departamento · qué municipios lo empujaron",
                       f"{inf['comparacion']}. Ordenados por el tamaño del cambio.", ancho=8)
    filas = [[d["nombre"], d["valor"], d["antes"], d["ahora"], d["cambio"], d["variacion"], d["estado"] or "",
              _plano(d["frase"])] for d in deps]
    ini, fin = _tabla(ws, fila, ["Departamento", f"{inf['metrica']} (periodo)", "Mes anterior", "Último mes", "Cambio",
                                 "Variación", "Estado", "Lo empujaron"], filas,
                      {2: FMT_NUM, 3: FMT_NUM, 4: FMT_NUM, 5: FMT_SIGNO, 6: FMT_PCT_SIGNO}, nombre="TablaDepartamentos",
                      anchos={1: 22, 2: 16, 3: 14, 4: 14, 5: 13, 6: 12, 7: 11, 8: 80})
    for i in range(ini, fin + 1):
        _pintar_estado(ws.cell(i, 7), ws.cell(i, 7).value)
        ws.cell(i, 8).alignment = Alignment(wrap_text=True, vertical="top")


def _hoja_mes_a_mes(wb, inf: dict, mensual: pd.DataFrame) -> None:
    if mensual is None or mensual.empty or mensual.shape[1] < 2:
        return
    ws = wb.create_sheet("Mes a mes")
    meses = list(mensual.columns)
    fila = _encabezado(ws, f"{inf['metrica']} por zona y mes",
                       "Cada zona por mes, ordenadas por su total. La escala de color va de lo más bajo (claro) a lo más "
                       "alto (intenso) de toda la tabla."
                       + (f" Ojo: {T.etiqueta_mes(inf['mes_b'])} va hasta el día {inf['corte_dia']}." if inf["corte_dia"] else ""),
                       ancho=min(len(meses) + 2, 16))
    encabezados = ["Zona"] + [T.etiqueta_mes(m, True) for m in meses] + ["Total"]
    filas = [[str(z)] + [_num(v) for v in fila_v] + [_num(np.nansum(fila_v.astype(float)))]
             for z, fila_v in zip(mensual.index, mensual.to_numpy())]
    filas.append(["Total"] + [_num(mensual[m].sum()) for m in meses] + [_num(np.nansum(mensual.to_numpy(dtype=float)))])
    formatos = {j: FMT_NUM for j in range(2, len(encabezados) + 1)}
    ini, fin = _tabla(ws, fila, encabezados, filas, formatos, anchos={1: 34, **{j: 12 for j in range(2, len(encabezados) + 1)}})
    ultima = get_column_letter(len(meses) + 1)
    ws.conditional_formatting.add(f"B{ini}:{ultima}{fin - 1}", ColorScaleRule(
        start_type="min", start_color="FFFFFF", end_type="max", end_color="5EC4BE"))
    for j in range(1, len(encabezados) + 1):
        c = ws.cell(fin, j)
        c.font = Font(bold=True)
        c.border = Border(top=Side(style="medium", color=TINTA))


def _alto_texto(texto, ancho_col: float) -> float:
    """Alto de fila para que un texto con varias líneas y ajuste se vea completo."""
    caracteres = int(max(ancho_col * 1.15, 10))
    renglones = sum(max(1, -(-len(linea) // caracteres)) for linea in str(texto or "").split("\n"))
    return min(409, 15.5 * max(2, renglones) + 4)


_ESTADOS_SEGUIMIENTO ='"Pendiente,En curso,Hecho,Descartado"'


def _seguimiento(ws, ini: int, fin: int, col_estado: int) -> None:
    """Lista desplegable de estado y color según cómo va cada acción."""
    if fin < ini:
        return
    dv = DataValidation(type="list", formula1=_ESTADOS_SEGUIMIENTO, allow_blank=True)
    ws.add_data_validation(dv)
    letra = get_column_letter(col_estado)
    dv.add(f"{letra}{ini}:{letra}{fin}")
    from openpyxl.formatting.rule import FormulaRule
    for texto, fondo, color in (("Hecho", "DCFCE7", "15803D"), ("En curso", "FEF9C3", "A16207"),
                                ("Pendiente", "FEE2E2", "B91C1C")):
        ws.conditional_formatting.add(f"{letra}{ini}:{letra}{fin}", FormulaRule(
            formula=[f'{letra}{ini}="{texto}"'], fill=PatternFill("solid", fgColor=fondo), font=Font(color=color, bold=True)))


def _hoja_plan(wb, plan: dict) -> None:
    """El plan de acción para repartir y hacerle seguimiento: una fila por acción."""
    zonas = plan["zonas"]
    ws = wb.create_sheet("Plan de acción")
    fila = _encabezado(ws, "Plan de acción territorial",
                       "Una fila por zona con algo que hacer, ordenadas por urgencia y por lo que valen al mes. Las columnas "
                       "«Estado» y «Comentarios» son para el seguimiento semanal: elige el estado en la lista.", ancho=13)
    filas = []
    acciones = zonas[zonas["estrategia"] != "sostener"].copy()
    acciones["_orden"] = acciones["estrategia"].map(lambda e: P.ESTRATEGIAS[e]["orden"])
    acciones = acciones.sort_values(["_orden", "valor_mes"], ascending=[True, False])
    for r in acciones.itertuples():
        e = P.ESTRATEGIAS[r.estrategia]
        filas.append([e["orden"], f"{e['icono']} {e['titulo']}", str(r.nombre), str(getattr(r, "departamento", "")),
                      "\n".join(f"{k}. {_plano(x)}" for k, x in enumerate(r.pasos, start=1)), r.responsable or "",
                      r.plazo, _num(r.ritmo), _num(r.meta_mes), _num(r.valor_mes), r.kpi, "Pendiente", ""])
    for r in (plan["aperturas"].itertuples() if len(plan["aperturas"]) else []):
        filas.append([P.ESTRATEGIAS["abrir"]["orden"], "🎯 Abrir", str(r.municipio), str(r.departamento),
                      "\n".join(f"{k}. {_plano(x)}" for k, x in enumerate(r.pasos, start=1)), r.agente_sugerido or "",
                      P.ESTRATEGIAS["abrir"]["plazo"], 0, _num(r.potencial_mes), _num(r.potencial_mes),
                      f"{T.cifra(r.potencial_mes)} al mes en 90 días", "Pendiente", ""])
    ini, fin = _tabla(ws, fila, ["Prioridad", "Estrategia", "Zona", "Departamento", "Qué hacer", "Responsable", "Plazo",
                                 "Ritmo al mes", "Meta al mes", "Vale al mes", "Cómo se mide", "Estado", "Comentarios"],
                      filas, {8: FMT_NUM, 9: FMT_NUM, 10: FMT_NUM}, nombre="TablaPlan",
                      anchos={1: 9, 2: 15, 3: 22, 4: 18, 5: 80, 6: 18, 7: 14, 8: 13, 9: 13, 10: 13, 11: 28, 12: 13, 13: 30})
    for i in range(ini, fin + 1):
        ws.cell(i, 5).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 11).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = _alto_texto(ws.cell(i, 5).value, 80)
        estrategia = str(ws.cell(i, 2).value)
        tono = "bajo" if ("Rescatar" in estrategia or "Recuperar" in estrategia) else \
            "subio" if "Replicar" in estrategia else "estable"
        ws.cell(i, 2).font = Font(bold=True, color=SEM[tono][0])
    _seguimiento(ws, ini, fin, 12)
    ws.freeze_panes = None
    if fin >= ini:
        ws.conditional_formatting.add(f"J{ini}:J{fin}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                    color="F4A3A3"))


def _hoja_agentes(wb, plan: dict) -> None:
    agentes = plan["agentes"]
    if not agentes:
        return
    ws = wb.create_sheet("Agentes")
    n = len(agentes)
    fila = _encabezado(ws, "Plan por agente comercial",
                       f"Posición y mediana sobre los {n} agentes. Ritmo = lo que vende al mes hoy; meta sugerida = su mejor "
                       "nivel reciente (promedio de sus últimos 3 meses completos o el último, el mayor).", ancho=12)
    filas = []
    for a in agentes:
        filas.append([a["posicion"], a["nombre"], f"{a['perfil_icono']} {a['perfil_titulo']}", _num(a["ritmo"]), _num(a["var"]),
                      _num(a["vs_mediana"]), _num(a["meta"]), a["zonas"],
                      ", ".join(f"{p['zona']} ({T.cifra_signo(p['delta'])})" for p in a["perdio"]),
                      ", ".join(c["nombre"] for c in a["cuentas_perdidas"]),
                      " → ".join(r["zona"] for r in a["ruta"]),
                      "\n".join(f"{k}. {_plano(x)}" for k, x in enumerate(a["estrategia"], start=1))])
    ini, fin = _tabla(ws, fila, ["#", "Agente", "Perfil", "Ritmo al mes", "Vs mes anterior", f"Vs mediana de {n}",
                                 "Meta sugerida", "Zonas", "Dónde perdió", "Clientes que dejaron de comprar", "Ruta sugerida",
                                 "Estrategia"], filas,
                      {4: FMT_NUM, 5: FMT_PCT_SIGNO, 6: FMT_PCT_SIGNO, 7: FMT_NUM}, nombre="TablaAgentes",
                      anchos={1: 5, 2: 22, 3: 17, 4: 13, 5: 12, 6: 13, 7: 13, 8: 8, 9: 34, 10: 34, 11: 34, 12: 80})
    for i in range(ini, fin + 1):
        for c in (9, 10, 11, 12):
            ws.cell(i, c).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = max(_alto_texto(ws.cell(i, 12).value, 80), _alto_texto(ws.cell(i, 9).value, 34),
                                          _alto_texto(ws.cell(i, 10).value, 34))
        perfil = str(ws.cell(i, 3).value)
        tono = "bajo" if "caída" in perfil else "subio" if "Referente" in perfil else "estable"
        ws.cell(i, 3).font = Font(bold=True, color=SEM[tono][0])
    if fin >= ini:
        ws.conditional_formatting.add(f"E{ini}:E{fin}", ColorScaleRule(
            start_type="num", start_value=-0.3, start_color="F8B4B4", mid_type="num", mid_value=0, mid_color="FFFFFF",
            end_type="num", end_value=0.3, end_color="A7F3D0"))


def _hoja_abrir(wb, plan: dict) -> None:
    ap = plan["aperturas"]
    if ap is None or ap.empty:
        return
    ws = wb.create_sheet("Dónde abrir")
    fila = _encabezado(ws, "Dónde abrir · municipios sin presencia",
                       f"Municipios de 20.000+ habitantes sin presencia, en los departamentos donde ya operas. Potencial = la mitad "
                       f"de la penetración típica de tu red ({T.cifra(plan['tipica'])} por cada 10.000 habitantes al mes). "
                       "Población: DANE 2026.", ancho=10)
    filas = [[r.municipio, r.departamento, _num(r.poblacion), _num(r.poblacion_cabecera), _num(r.crecimiento_2030), r.base,
              _num(round(r.distancia_km)), r.agente_sugerido or "", _num(r.potencial_mes), r.modelo] for r in ap.itertuples()]
    ini, fin = _tabla(ws, fila, ["Municipio", "Departamento", "Población 2026", "Urbana", "Crec. 2030", "Atender desde",
                                 "Km", "Agente sugerido", "Potencial al mes", "Modelo de entrada"], filas,
                      {3: FMT_NUM, 4: FMT_NUM, 5: FMT_PCT_SIGNO, 7: FMT_NUM, 9: FMT_NUM}, nombre="TablaAbrir",
                      anchos={1: 24, 2: 20, 3: 14, 4: 12, 5: 11, 6: 22, 7: 7, 8: 18, 9: 15, 10: 70})
    for i in range(ini, fin + 1):
        ws.cell(i, 10).alignment = Alignment(wrap_text=True, vertical="top")
    if fin >= ini:
        ws.conditional_formatting.add(f"I{ini}:I{fin}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                                                    color="C4B5FD"))


def _hoja_crecer(wb, inf: dict) -> None:
    crecer = inf["crecer"] or {}
    blancos, rezagados = crecer.get("blancos"), crecer.get("rezagados")
    if (blancos is None or blancos.empty) and (rezagados is None or rezagados.empty):
        return
    ws = wb.create_sheet("Dónde crecer")
    tipica = crecer.get("penetracion_tipica")
    fila = _encabezado(ws, "Dónde crecer",
                       f"Referencia: en un municipio típico de tu red hay {T.cifra(tipica)} de {inf['metrica'].lower()} por "
                       "cada 10.000 habitantes. El potencial es la MITAD del camino hasta esa referencia. Población: DANE 2026.",
                       ancho=6)
    if blancos is not None and len(blancos):
        fila = _seccion(ws, fila, "Municipios grandes sin presencia (en los departamentos donde ya operas)", ancho=6)
        _, fin = _tabla(ws, fila, ["Municipio", "Departamento", "Población", "Urbana", "Potencial estimado"],
                        [[r.municipio, r.departamento, _num(r.poblacion), _num(r.poblacion_cabecera), _num(r.potencial)]
                         for r in blancos.itertuples()], {3: FMT_NUM, 4: FMT_NUM, 5: FMT_NUM}, nombre="TablaBlancos",
                        anchos={1: 26, 2: 22, 3: 14, 4: 14, 5: 18})
        fila = fin + 2
    if rezagados is not None and len(rezagados):
        fila = _seccion(ws, fila, "Con presencia pero por debajo de lo normal", ancho=6)
        _tabla(ws, fila, ["Municipio", "Departamento", "Población", "Por 10.000 hab.", "Potencial estimado"],
               [[r.nombre, r.departamento, _num(r.poblacion), _num(r.por_10k), _num(r.potencial)]
                for r in rezagados.itertuples()], {3: FMT_NUM, 4: "#,##0.0", 5: FMT_NUM}, nombre="TablaRezagados")


def _hoja_notas(wb, inf: dict) -> None:
    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 120
    fila = _encabezado(ws, "Cómo se calcula", "Para leer el informe sin la app.", ancho=1)
    notas = [
        "Semáforo: compara el último mes con datos contra el anterior. Verde = subió más de 5%; amarillo = entre −5% y +5%; "
        "rojo = bajó más de 5%. Una zona que deja de tener actividad cuenta como roja y una que empieza, como verde.",
        "Mes a medias: si el último mes llega hasta un día bastante anterior al que suelen llegar los demás meses del archivo, "
        "se compara contra el mes anterior HASTA ESE MISMO DÍA, para no confundir un mes incompleto con una caída.",
        "Por qué subió o bajó: entre las columnas del archivo (canal, asesor, producto… o el municipio, dentro de un "
        "departamento) se elige la que concentra el movimiento NETO en menos nombres; las que solo rotan (mucho movimiento "
        "que se anula entre sí) no cuentan como razón. Solo con métricas que se suman.",
        "Posición y participación: siempre sobre TODAS las zonas visibles con los filtros aplicados, no solo las mostradas.",
        "Población, municipios y contornos: DANE (DIVIPOLA, Marco Geoestadístico Nacional y proyecciones de población 2026).",
        "Dónde crecer: la penetración típica es la mediana de valor por cada 10.000 habitantes en los municipios donde ya hay "
        "actividad; el potencial es la mitad del camino hasta ella.",
    ]
    for n in notas:
        fila = _parrafo(ws, fila, "• " + n, ancho=1, alto=48)


def build_territorial_excel(inf: dict, ctx: dict, serie_total: pd.DataFrame | None = None,
                            mensual: pd.DataFrame | None = None, plan: dict | None = None) -> bytes:
    """El Excel ejecutivo del territorio. `ctx` = {archivo, hoja, registros, filtros, generado}.
    Con `plan` (`core/territorio_plan.plan`), además: el plan en la portada y las hojas «Plan de acción»
    (para el seguimiento), «Agentes» y «Dónde abrir» (que reemplaza a «Dónde crecer»)."""
    ctx = _ctx(ctx)
    wb = Workbook()
    fila = _hoja_resumen(wb, inf, ctx, serie_total if serie_total is not None else pd.DataFrame(), plan)
    if plan:
        _hoja_plan(wb, plan)
        _hoja_agentes(wb, plan)
    _hoja_semaforo(wb, inf)
    _hoja_departamentos(wb, inf)
    _hoja_mes_a_mes(wb, inf, mensual)
    if plan and len(plan.get("aperturas", [])):
        _hoja_abrir(wb, plan)
    else:
        _hoja_crecer(wb, inf)
    _hoja_notas(wb, inf)
    _indice(wb, fila)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _ctx(ctx: dict) -> dict:
    ctx = dict(ctx or {})
    ctx.setdefault("archivo", "archivo")
    ctx.setdefault("hoja", "")
    ctx.setdefault("registros", 0)
    ctx.setdefault("generado", datetime.now().strftime("%d/%m/%Y %H:%M"))
    filtros = ctx.get("filtros") or ""
    ctx["filtros"] = f" · filtros: {filtros}" if filtros and not str(filtros).startswith(" · ") else filtros
    return ctx


def nombre_archivo(base: str, extension: str) -> str:
    limpio = re.sub(r"[^\w\-]+", "_", str(base)).strip("_")[:40] or "territorio"
    return f"informe_territorial_{limpio}_{datetime.now():%Y%m%d}.{extension}"


# ── HTML ──────────────────────────────────────────────────────────────────

_CSS = """
:root{--tinta:#131826;--muted:#5b6473;--linea:#d8dce6;--panel:#f6f7f9;--teal:#0f8a85;
  --verde:#16a34a;--verde-s:#dcfce7;--ambar:#ca8a04;--ambar-s:#fef9c3;--rojo:#dc2626;--rojo-s:#fee2e2;--violeta:#9333ea}
*{box-sizing:border-box}body{margin:0;font-family:Inter,"Segoe UI",system-ui,sans-serif;color:var(--tinta);background:#eef1f5}
.pagina{max-width:1280px;margin:0 auto;padding:28px 22px 60px}
header.portada{background:linear-gradient(135deg,#0f172a,#13324a 60%,#0f8a85);color:#fff;border-radius:20px;padding:28px 32px;
  box-shadow:0 18px 40px rgba(15,23,42,.25)}
header.portada small{letter-spacing:.16em;text-transform:uppercase;font-weight:800;font-size:11px;opacity:.8}
header.portada h1{margin:6px 0 6px;font-size:30px;letter-spacing:-.02em}
header.portada p{margin:0;opacity:.85;font-size:13.5px}
.aviso{margin-top:14px;background:rgba(250,204,21,.18);border:1px solid rgba(250,204,21,.5);color:#fde68a;border-radius:10px;
  padding:8px 12px;font-size:13px;font-weight:600}
section{background:#fff;border:1px solid var(--linea);border-radius:18px;padding:20px 22px;margin-top:18px;
  box-shadow:0 2px 10px rgba(20,26,43,.05)}
h2{margin:0 0 4px;font-size:18px}h2+p.sub{margin:0 0 14px;color:var(--muted);font-size:13px}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(180px,1fr));gap:12px;margin-top:18px}
.kpi{background:#fff;border:1px solid var(--linea);border-top:4px solid var(--teal);border-radius:14px;padding:14px 16px}
.kpi span{display:block;font-size:11px;font-weight:800;letter-spacing:.08em;text-transform:uppercase;color:var(--muted)}
.kpi b{display:block;font-size:28px;margin-top:4px}.kpi small{color:var(--muted);font-size:12.5px}
.kpi.verde{border-top-color:var(--verde)}.kpi.verde b{color:var(--verde)}
.kpi.rojo{border-top-color:var(--rojo)}.kpi.rojo b{color:var(--rojo)}
.barra{display:flex;height:14px;border-radius:99px;overflow:hidden;gap:3px;margin:8px 0 4px}
.barra i{display:block}.barra .subio{background:#22c55e}.barra .estable{background:#eab308}.barra .bajo{background:#ef4444}
.total{display:grid;grid-template-columns:minmax(240px,.8fr) 1.5fr 1.5fr;gap:20px;align-items:start}
.total .cifra{font-size:26px;font-weight:800}.total .cifra.subio{color:var(--verde)}.total .cifra.bajo{color:var(--rojo)}
.total h4{margin:0 0 4px;font-size:11px;letter-spacing:.1em;text-transform:uppercase;color:var(--muted)}
.total p{margin:0;font-size:14.5px;line-height:1.55}
.cols{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:14px}
.col{border:1px solid var(--linea);border-top:5px solid;border-radius:14px;padding:14px 16px}
.col.bajo{border-top-color:#ef4444}.col.estable{border-top-color:#eab308}.col.subio{border-top-color:#22c55e}
.col .cab{display:flex;justify-content:space-between;align-items:baseline}.col .cab b{font-size:30px}
.col.bajo .cab b{color:var(--rojo)}.col.estable .cab b{color:var(--ambar)}.col.subio .cab b{color:var(--verde)}
.fila{padding:8px 0;border-top:1px dashed var(--linea)}.fila .l1{display:flex;gap:8px;align-items:baseline;font-size:14px}
.fila .l1 b{flex:1}.fila small{display:block;color:var(--muted);font-size:12.5px;margin-top:2px;line-height:1.45}
.chip{display:inline-block;font-size:11.5px;font-weight:800;padding:2px 8px;border-radius:99px}
.chip.subio{color:var(--verde);background:var(--verde-s)}.chip.estable{color:var(--ambar);background:var(--ambar-s)}
.chip.bajo{color:var(--rojo);background:var(--rojo-s)}
table{width:100%;border-collapse:collapse;font-size:13.5px}th{position:sticky;top:0;background:var(--tinta);color:#fff;
  text-align:left;padding:9px 10px;font-size:12px;cursor:pointer;user-select:none;white-space:nowrap}
th:hover{background:#25304a}td{padding:8px 10px;border-bottom:1px solid #eceff4;vertical-align:top}
tr:nth-child(even) td{background:#fafbfc}td.n{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
td small{color:var(--muted)}.tabla-caja{max-height:640px;overflow:auto;border:1px solid var(--linea);border-radius:12px}
.buscar{width:100%;max-width:360px;padding:9px 12px;border:1px solid var(--linea);border-radius:10px;font-size:14px;margin-bottom:10px}
.mapa{width:100%;height:640px;border:0;border-radius:14px;background:#e5e7eb}
ul.lectura{margin:0;padding-left:20px}ul.lectura li{margin-bottom:7px;line-height:1.55;font-size:14.5px}
.dot{display:inline-block;width:10px;height:10px;border-radius:50%;margin-right:6px}
.dot.subio{background:#22c55e}.dot.estable{background:#eab308}.dot.bajo{background:#ef4444}
footer{color:var(--muted);font-size:12px;margin-top:24px;line-height:1.6}
.jugadas{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));gap:14px;margin-top:14px}
.jug{border:1px solid var(--linea);border-left:6px solid var(--muted);border-radius:14px;padding:14px 16px;break-inside:avoid}
.jug.bajo{border-left-color:#ef4444}.jug.estable{border-left-color:#eab308}.jug.subio{border-left-color:#22c55e}
.jug.oport{border-left-color:#a855f7}.chip.oport{color:var(--violeta);background:#f3e8ff}
.jcab{display:flex;align-items:center;gap:10px;flex-wrap:wrap}.jn{width:28px;height:28px;border-radius:8px;background:var(--tinta);
  color:#fff;display:grid;place-items:center;font-weight:800}.jv{margin-left:auto;color:var(--muted);font-size:13px}
.jv b{color:var(--tinta);font-size:16px}.jug h3{margin:10px 0 6px;font-size:20px}.jug h3 small{color:var(--muted);font-size:13px;font-weight:500}
.jug ol,.agente ol{margin:0;padding-left:20px}.jug li,.agente li{font-size:14px;line-height:1.5;margin-bottom:4px}
.jpie{margin-top:8px;padding-top:8px;border-top:1px dashed var(--linea);color:var(--muted);font-size:12.5px}
details.agente{border:1px solid var(--linea);border-left:6px solid var(--muted);border-radius:12px;padding:10px 14px;margin-top:8px}
details.agente.bajo{border-left-color:#ef4444}details.agente.subio{border-left-color:#22c55e}details.agente.estable{border-left-color:#eab308}
details.agente summary{cursor:pointer;display:flex;gap:12px;align-items:center;flex-wrap:wrap;font-size:14.5px}
details.agente summary span:last-child{color:var(--muted);font-size:13px}details.agente p{margin:8px 0 4px;font-size:13.5px}
@media(max-width:900px){.jugadas{grid-template-columns:1fr}}
@media print{details.agente{break-inside:avoid}details.agente>*{display:block}}
@media(max-width:900px){.total,.cols{grid-template-columns:1fr}}
@media print{body{background:#fff}.pagina{padding:0}section,header.portada{box-shadow:none;break-inside:avoid}
  .buscar,.mapa-caja{display:none}.tabla-caja{max-height:none;overflow:visible}}
"""

_JS = """
document.querySelectorAll('table.ordenable').forEach(function(t){
  t.querySelectorAll('th').forEach(function(th,i){th.addEventListener('click',function(){
    var cuerpo=t.tBodies[0],filas=Array.from(cuerpo.rows),asc=th.dataset.asc!=='1';th.dataset.asc=asc?'1':'0';
    filas.sort(function(a,b){var x=a.cells[i].dataset.v,y=b.cells[i].dataset.v;
      if(x!==undefined&&y!==undefined){return asc?(+x)-(+y):(+y)-(+x);}
      return asc?a.cells[i].innerText.localeCompare(b.cells[i].innerText):b.cells[i].innerText.localeCompare(a.cells[i].innerText);});
    filas.forEach(function(f){cuerpo.appendChild(f);});});});});
var b=document.getElementById('buscar');if(b){b.addEventListener('input',function(){var q=b.value.toLowerCase();
  document.querySelectorAll('#tabla-zonas tbody tr').forEach(function(f){f.style.display=f.innerText.toLowerCase().indexOf(q)>=0?'':'none';});});}
"""


def _svg_linea(serie: pd.DataFrame, inf: dict, ancho: int = 1180, alto: int = 240) -> str:
    """La evolución del total en SVG (sin librerías). El mes a medias va punteado."""
    if serie is None or len(serie) < 2:
        return ""
    v = [float(x) for x in serie["valor"]]
    mn, mx = min(min(v), 0), max(v)
    rango = (mx - mn) or 1.0
    izq, der, arr, aba = 70, 20, 20, 40
    xs = [izq + i / (len(v) - 1) * (ancho - izq - der) for i in range(len(v))]
    ys = [arr + (1 - (x - mn) / rango) * (alto - arr - aba) for x in v]
    medias = bool(inf.get("corte_dia")) and serie["mes"].iloc[-1] == inf.get("mes_b")
    n_ok = len(v) - 1 if medias else len(v)
    puntos = " ".join(f"{x:.1f},{y:.1f}" for x, y in zip(xs[:n_ok], ys[:n_ok]))
    partes = [f'<svg viewBox="0 0 {ancho} {alto}" width="100%" role="img">']
    for k in range(5):
        y = arr + k / 4 * (alto - arr - aba)
        valor = mx - k / 4 * rango
        partes.append(f'<line x1="{izq}" x2="{ancho - der}" y1="{y:.1f}" y2="{y:.1f}" stroke="#e5e7eb"/>'
                      f'<text x="{izq - 8}" y="{y + 4:.1f}" font-size="11" text-anchor="end" fill="#5b6473">{html.escape(T.cifra(valor))}</text>')
    partes.append(f'<polyline points="{puntos}" fill="none" stroke="#0f8a85" stroke-width="3" stroke-linejoin="round"/>')
    if medias:
        partes.append(f'<line x1="{xs[-2]:.1f}" y1="{ys[-2]:.1f}" x2="{xs[-1]:.1f}" y2="{ys[-1]:.1f}" stroke="#0f8a85" '
                      f'stroke-width="2" stroke-dasharray="5 5"/>')
    for i, (x, y) in enumerate(zip(xs, ys)):
        hueco = medias and i == len(v) - 1
        partes.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="5" fill="{"#fff" if hueco else "#0f8a85"}" stroke="#0f8a85" stroke-width="2">'
                      f'<title>{html.escape(T.etiqueta_mes(serie["mes"].iloc[i], True))}: {html.escape(T.cifra(v[i]))}</title></circle>'
                      f'<text x="{x:.1f}" y="{alto - 14}" font-size="11.5" text-anchor="middle" fill="#5b6473">'
                      f'{html.escape(T.etiqueta_mes(serie["mes"].iloc[i], True))}</text>')
    if medias:
        partes.append(f'<text x="{xs[-1]:.1f}" y="{ys[-1] - 12:.1f}" font-size="11.5" text-anchor="middle" fill="#ca8a04" '
                      f'font-weight="700">al día {inf["corte_dia"]}</text>')
    partes.append("</svg>")
    return "".join(partes)


def _plan_html(plan: dict) -> str:
    """🎯 El plan de acción en el informe HTML: jugadas, agentes y dónde abrir."""
    e = html.escape
    partes = ['<section><h2>🎯 Plan de acción</h2><p class="sub">Qué hacer, dónde, quién y cuánto vale al mes. Primero lo que más vale y '
              'antes se logra (recuperar es rápido; abrir es una apuesta a 90 días).</p><ul class="lectura">' + "".join(f"<li>{_negritas_html(f)}</li>" for f in plan["resumen"]) + "</ul>"]
    tarjetas = []
    for j in plan["jugadas"]:
        valor = f"vale ≈ <b>{e(T.cifra(j['valor']))}</b> al mes" if j["valor"] else "buena práctica"
        pie = " · ".join(x for x in (f"👤 {e(j['responsable'])}" if j["responsable"] else "", f"⏱ {e(j['plazo'])}",
                                     f"📏 {e(j['kpi'])}") if x)
        tarjetas.append(f'<div class="jug {j["tono"]}"><div class="jcab"><span class="jn">{j["n"]}</span>'
                        f'<span class="chip {j["tono"]}">{j["icono"]} {e(j["titulo"])}</span><span class="jv">{valor}</span></div>'
                        f'<h3>{e(j["zona"])} <small>{e(j["departamento"])}</small></h3>'
                        "<ol>" + "".join(f"<li>{_negritas_html(x)}</li>" for x in j["pasos"]) + f'</ol><div class="jpie">{pie}</div></div>')
    partes.append('<div class="jugadas">' + "".join(tarjetas) + "</div></section>")
    agentes = plan["agentes"]
    if agentes:
        n = len(agentes)
        fichas = []
        for a in agentes:
            cuerpo = "<ol>" + "".join(f"<li>{_negritas_html(x)}</li>" for x in a["estrategia"]) + "</ol>"
            extra = []
            if a["perdio"]:
                extra.append("<b>Dónde perdió:</b> " + ", ".join(f"{e(p['zona'])} ({e(T.cifra_signo(p['delta']))})" for p in a["perdio"]))
            if a["cuentas_perdidas"]:
                extra.append("<b>Clientes que dejaron de comprar:</b> " + ", ".join(e(c["nombre"]) for c in a["cuentas_perdidas"]))
            if a["ruta"]:
                extra.append("<b>Ruta sugerida:</b> " + " → ".join(e(r["zona"]) for r in a["ruta"]))
            var = "—" if a["var"] is None else f"{a['var']:+.0%}"
            fichas.append(f'<details class="agente {a["perfil_tono"]}"><summary><span class="chip {a["perfil_tono"]}">'
                          f'{a["perfil_icono"]} {e(a["perfil_titulo"])}</span><b>{e(a["nombre"])}</b>'
                          f'<span>{a["posicion"]}.º de {n} · ritmo {e(T.cifra(a["ritmo"]))}/mes · {var} · meta '
                          f'{e(T.cifra(a["meta"]))}</span></summary>'
                          + "".join(f"<p>{x}</p>" for x in extra) + cuerpo + "</details>")
        partes.append(f'<section><h2>👥 Plan por agente comercial</h2><p class="sub">{n} agentes. Posición y mediana sobre todos. '
                      'Clic en un agente para ver su plan; se puede imprimir o reenviar.</p>' + "".join(fichas) + "</section>")
    ap = plan["aperturas"]
    if ap is not None and len(ap):
        filas = "".join(f'<tr><td><b>{e(str(r.municipio))}</b><br><small>{e(str(r.departamento))}</small></td>'
                        f'<td class="n" data-v="{r.poblacion}">{e(T.cifra(r.poblacion))}</td>'
                        f'<td>{e(str(r.base))}<br><small>{r.distancia_km:.0f} km</small></td><td>{e(r.agente_sugerido or "—")}</td>'
                        f'<td class="n" data-v="{r.potencial_mes}">{e(T.cifra(r.potencial_mes))}</td><td><small>{e(r.modelo)}</small></td></tr>'
                        for r in ap.itertuples())
        partes.append('<section><h2>📍 Dónde abrir</h2><p class="sub">Municipios de 20.000+ habitantes sin presencia en los '
                      f'departamentos donde operas. Potencial al mes = la mitad de la penetración típica ({e(T.cifra(plan["tipica"]))} '
                      'por cada 10.000 hab.).</p><div class="tabla-caja"><table class="ordenable"><thead><tr><th>Municipio</th>'
                      '<th>Población</th><th>Atender desde</th><th>Agente sugerido</th><th>Potencial al mes</th><th>Modelo de entrada</th>'
                      f'</tr></thead><tbody>{filas}</tbody></table></div></section>')
    return "".join(partes)


def build_territorial_html(inf: dict, ctx: dict, serie_total: pd.DataFrame | None = None,
                           mapa_html: str | None = None, plan: dict | None = None) -> bytes:
    """Página autocontenida con el análisis territorial (y el mapa interactivo si se pasa `mapa_html`)."""
    ctx = _ctx(ctx)
    e = html.escape
    sem, ct, tabla = inf["semaforo"], inf["cambio_total"], inf["tabla"]
    zona_txt = _zona_txt(inf["nivel"])
    partes = [f'<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,'
              f'initial-scale=1"><title>Informe territorial · {e(inf["metrica"])}</title><style>{_CSS}</style></head><body>'
              '<div class="pagina">']
    partes.append(f'<header class="portada"><small>Análisis territorial · Colombia</small><h1>Informe territorial · '
                  f'{e(inf["metrica"])}</h1><p>{e(ctx["archivo"])} · hoja {e(str(ctx["hoja"]))} · {ctx["registros"]:,} registros'
                  f'{e(ctx["filtros"])} · generado el {e(ctx["generado"])}</p>'
                  + (f'<p style="margin-top:6px">Semáforo: <b>{e(inf["comparacion"])}</b></p>' if inf["comparacion"] else "")
                  + (f'<div class="aviso">⏱ {e(T.etiqueta_mes(inf["mes_b"]).capitalize())} va hasta el día {inf["corte_dia"]}: '
                     f'se compara contra {e(T.etiqueta_mes(inf["mes_a"]))} hasta ese mismo día.</div>' if inf["corte_dia"] else "")
                  + "</header>")

    # Indicadores.
    kpis = [("", ("Total" if inf["sumable"] else "Promedio") + f" · {inf['metrica']}", T.cifra(inf["total"]), "todo el periodo")]
    if ct:
        kpis.append(("verde" if ct["delta"] >= 0 else "rojo", "Cambio vs mes anterior",
                     f"{ct['pct']:+.1%}" if ct["pct"] is not None else T.cifra_signo(ct["delta"]),
                     f"{T.cifra_signo(ct['delta'])} · {T.cifra(ct['antes'])} → {T.cifra(ct['ahora'])}"))
    kpis.append(("", "Zonas con actividad", f"{inf['n']:,}", zona_txt))
    if sem:
        kpis.append(("verde", "Subieron", str(sem["subio"]["n"]), f"{T.cifra_signo(sem['subio']['cambio'])} entre todas"))
        kpis.append(("rojo", "Bajaron", str(sem["bajo"]["n"]), f"{T.cifra_signo(sem['bajo']['cambio'])} entre todas"))
    cob = inf["cobertura"]
    if cob and cob.get("pct_deptos") is not None:
        kpis.append(("", "Cobertura de población", f"{cob['pct_deptos']:.0%}", "de la gente de tus departamentos"))
    partes.append('<div class="kpis">' + "".join(f'<div class="kpi {c}"><span>{e(t)}</span><b>{e(v)}</b><small>{e(d)}</small></div>'
                                                 for c, t, v, d in kpis) + "</div>")

    if plan and (plan.get("jugadas") or plan.get("agentes")):
        partes.append(_plan_html(plan))

    # Qué pasó y por qué.
    if ct and sem:
        sube = ct["delta"] >= 0
        total_n = sum(sem[k]["n"] for k in ("subio", "estable", "bajo")) or 1
        barra = '<div class="barra">' + "".join(f'<i class="{k}" style="width:{sem[k]["n"] / total_n * 100:.1f}%"></i>'
                                                for k in ("bajo", "estable", "subio") if sem[k]["n"]) + "</div>"
        razones = "".join(f'<div><h4>{t}</h4><p>{_negritas_html(inf[c]["frase"])}</p></div>'
                          for t, c in (("📍 Por dónde", "por_donde"), ("🔎 Por qué", "por_que")) if inf.get(c))
        partes.append(f'<section><h2>🚦 Qué pasó y por qué</h2><p class="sub">{e(inf["comparacion"])}. Verde: subió más de 5% · '
                      f'amarillo: se mantuvo (±5%) · rojo: bajó más de 5%.</p><div class="total"><div><h4>Total de '
                      f'{e(inf["metrica"].lower())}</h4><div class="cifra {"subio" if sube else "bajo"}">{"▲ Subió" if sube else "▼ Bajó"} '
                      f'{e(T.cifra(abs(ct["delta"])))}' + (f' ({ct["pct"]:+.1%})' if ct["pct"] is not None else "")
                      + f'</div>{barra}<small>🟢 {sem["subio"]["n"]} subieron · 🟡 {sem["estable"]["n"]} estables · '
                      f'🔴 {sem["bajo"]["n"]} bajaron</small></div>{razones}</div></section>')

    if mapa_html:
        partes.append('<section class="mapa-caja"><h2>🗺️ Mapa</h2><p class="sub">Interactivo: acerca, aleja, gira con clic derecho y pasa '
                      'el mouse por una zona para ver su detalle. El fondo del mapa necesita internet.</p>'
                      f'<iframe class="mapa" srcdoc="{e(mapa_html, quote=True)}" loading="lazy"></iframe></section>')

    # Rojo · amarillo · verde.
    if sem:
        cols = []
        for clave, titulo in (("bajo", "🔴 Bajaron"), ("estable", "🟡 Se mantuvieron (±5%)"), ("subio", "🟢 Subieron")):
            parte = sem[clave]
            filas = []
            for r in parte["tabla"].head(10).itertuples():
                razon = inf["razones"].get(str(r.zona), {}).get("frase", "") if clave != "estable" else ""
                filas.append(f'<div class="fila"><div class="l1"><b>{e(str(r.nombre))}</b><span>{e(T.cifra_signo(r.cambio))}</span>'
                             f'<span class="chip {clave}">{e(_var_txt(r.variacion, r.tendencia))}</span></div>'
                             + (f"<small>{_negritas_html(razon)}</small>" if razon else "") + "</div>")
            resto = parte["n"] - len(filas)
            cols.append(f'<div class="col {clave}"><div class="cab"><span><b style="font-size:14px">{titulo}</b></span><b>{parte["n"]}</b></div>'
                        f'<small style="color:var(--muted)">{e(T.cifra_signo(parte["cambio"]))} entre los {parte["n"]} {zona_txt}</small>'
                        + ("".join(filas) or '<div class="fila"><small>Ninguna zona en este grupo.</small></div>')
                        + (f'<div class="fila"><small>y {resto} más en la tabla completa ↓</small></div>' if resto > 0 else "")
                        + "</div>")
        partes.append(f'<section><h2>Zona por zona</h2><p class="sub">Las 10 que más se movieron de cada grupo, con la razón.</p>'
                      f'<div class="cols">{"".join(cols)}</div></section>')

    if inf["departamentos"]:
        filas = "".join(
            f'<tr><td><span class="dot {d["estado"] or ""}"></span>{e(d["nombre"])}</td>'
            f'<td class="n" data-v="{d["valor"]}">{e(T.cifra(d["valor"]))}</td>'
            f'<td class="n" data-v="{d["cambio"]}">{e(T.cifra_signo(d["cambio"]))}</td>'
            f'<td class="n" data-v="{d["variacion"] if d["variacion"] is not None else 0}"><span class="chip {d["estado"] or ""}">'
            f'{e(_var_txt(d["variacion"], d["tendencia"]))}</span></td><td>{_negritas_html(d["frase"])}</td></tr>'
            for d in inf["departamentos"])
        partes.append('<section><h2>Por departamento</h2><p class="sub">Qué municipios empujaron el cambio de cada departamento. '
                      'Clic en un encabezado para ordenar.</p><div class="tabla-caja"><table class="ordenable"><thead><tr>'
                      f'<th>Departamento</th><th>{e(inf["metrica"])}</th><th>Cambio</th><th>Variación</th><th>Lo empujaron</th>'
                      f'</tr></thead><tbody>{filas}</tbody></table></div></section>')

    if not tabla.empty:
        con_mes = "estado" in tabla.columns
        cab = ["#", "Zona", "Departamento", inf["metrica"], "Participación"] + (["Cambio", "Variación", "Por qué"] if con_mes else [])
        filas = []
        for r in tabla.itertuples():
            celdas = [f'<td class="n" data-v="{int(r.posicion)}">{int(r.posicion)}</td>', f"<td><b>{e(str(r.nombre))}</b></td>",
                      f"<td>{e(str(r.departamento))}</td>", f'<td class="n" data-v="{_num(r.valor) or 0}">{e(T.cifra(r.valor))}</td>',
                      f'<td class="n" data-v="{_num(getattr(r, "participacion", None)) or 0}">'
                      + (f'{r.participacion:.1%}' if _num(getattr(r, "participacion", None)) is not None else "—") + "</td>"]
            if con_mes:
                razon = inf["razones"].get(str(r.zona), {}).get("frase_larga", "")
                celdas += [f'<td class="n" data-v="{_num(r.cambio) or 0}">{e(T.cifra_signo(r.cambio))}</td>',
                           f'<td class="n" data-v="{_num(r.tendencia) or 0}"><span class="chip {r.estado or ""}">'
                           f'{e(_var_txt(r.variacion, r.tendencia))}</span></td>', f"<td><small>{_negritas_html(razon)}</small></td>"]
            filas.append("<tr>" + "".join(celdas) + "</tr>")
        partes.append(f'<section><h2>Todas las zonas</h2><p class="sub">{inf["n"]} {zona_txt}. Posición y participación sobre todas. '
                      'Busca por nombre o clic en un encabezado para ordenar.</p><input id="buscar" class="buscar" '
                      'placeholder="Buscar zona o departamento…"><div class="tabla-caja"><table id="tabla-zonas" class="ordenable">'
                      "<thead><tr>" + "".join(f"<th>{e(c)}</th>" for c in cab) + "</tr></thead><tbody>" + "".join(filas)
                      + "</tbody></table></div></section>")

    svg = _svg_linea(serie_total, inf)
    if svg:
        partes.append(f'<section><h2>Evolución del total</h2><p class="sub">{e(inf["metrica"])} mes a mes.</p>{svg}</section>')

    crecer = inf["crecer"] or {}
    blancos = crecer.get("blancos")
    if blancos is not None and len(blancos):
        filas = "".join(f'<tr><td><b>{e(str(r.municipio))}</b></td><td>{e(str(r.departamento))}</td>'
                        f'<td class="n" data-v="{r.poblacion}">{e(T.cifra(r.poblacion))}</td>'
                        f'<td class="n" data-v="{r.potencial}">{e(T.cifra(r.potencial))}</td></tr>' for r in blancos.itertuples())
        partes.append('<section><h2>🎯 Dónde crecer</h2><p class="sub">Municipios grandes sin presencia en los departamentos donde ya '
                      f'operas. Potencial = la mitad del camino a la penetración típica de tu red ({e(T.cifra(crecer.get("penetracion_tipica")))} '
                      'por cada 10.000 habitantes).</p><div class="tabla-caja"><table class="ordenable"><thead><tr><th>Municipio</th>'
                      f'<th>Departamento</th><th>Población 2026</th><th>Potencial</th></tr></thead><tbody>{filas}</tbody></table></div></section>')

    if inf["lectura"]:
        partes.append('<section><h2>Lectura del territorio</h2><ul class="lectura">'
                      + "".join(f"<li>{_negritas_html(f)}</li>" for f in inf["lectura"]) + "</ul></section>")

    partes.append('<footer>Semáforo: verde subió más de 5%, amarillo entre −5% y +5%, rojo bajó más de 5%; una zona que deja de tener '
                  'actividad cuenta como roja. Si el último mes va a medias, se compara contra el anterior hasta el mismo día. Las '
                  'razones son la columna que concentra el movimiento neto. Municipios, contornos y población: DANE (DIVIPOLA, MGN y '
                  'proyecciones 2026).</footer>')
    partes.append(f"</div><script>{_JS}</script></body></html>")
    return "".join(partes).encode("utf-8")
