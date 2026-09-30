"""Excel ejecutivo descargable: el libro que se le manda a los jefes.

Antes el botón «Excel» de Exportar volcaba los datos, la tabla de
estadística y los hallazgos en tres hojas sin formato: nombres de columna y
números sueltos, nada que se pudiera reenviar. Este libro se arma para que
quien lo abra sin la app entienda cómo va el negocio y qué hacer:

1. **Resumen**: portada con los indicadores, el veredicto del último mes,
   por qué se movió, las prioridades y los dos gráficos principales, con un
   índice que lleva a cada hoja.
2. **Plan de acción**: los mismos planes de «🎯 Qué atacar» (core/planes),
   con pasos, cómo se mide, control semanal y el valor de cada frente.
3. **Tendencia**, **Por <dimensión>** y **Matriz**: evolución mensual, cómo
   va cada uno (core/cuadro_comparativo, con la referencia sobre el grupo
   COMPLETO) y cruces tipo dinámica con mapa de calor.
4. **Tabla dinámica**: tablas dinámicas REALES de Excel sobre la hoja Datos
   (se recalculan al abrir y se pueden reorganizar arrastrando campos).
5. **Hallazgos**, **Estadística** y **Datos**: el soporte, con los datos
   como tabla de Excel con filtros.

No hay cálculo de negocio nuevo: cada hoja usa el mismo motor que su
pestaña en la app. Cada sección se arma por separado y, si el archivo no da
para ella (sin fechas no hay tendencia, sin dimensiones no hay ranking) o
falla, se omite sin tumbar el libro.
"""
from __future__ import annotations

import io
import math
import re
from datetime import date, datetime

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.pivot.cache import (CacheDefinition, CacheField, CacheSource, Missing, Number, SharedItems, Text,
                                  WorksheetSource)
from openpyxl.pivot.table import (DataField, FieldItem, Location, PageField, PivotField, PivotTableStyle,
                                  RowColField, TableDefinition)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.cuadro_comparativo import CERCA_DE_META, cuadro_comparativo, opciones_de_comparacion
from core.dates import a_datetime
from core.explorador import PERIODO, calculo_automatico, columna_fecha, tabla_cruzada, ultimo_incompleto
from core.numeric import numeric_valid
from ui.labels import clean_display_text
from ui.report_base import fecha_larga, mes
from visualization.charts import dimension_candidates, metric_candidates, month_columns

# ── Paleta ────────────────────────────────────────────────────────────────
# Tinta y rojo de marca (ui/styles/theme.py, modo claro) para la portada; los
# gráficos usan el azul de la serie 1 y un gris de comparación, para que el
# rojo no se lea como "malo" en cada barra. Los colores de estado (verde,
# ámbar, rojo) quedan reservados para el semáforo y siempre van con texto.
TINTA, MUTED, LINEA, ZEBRA, PANEL = "131826", "5B6473", "D8DCE6", "F6F7F9", "F1F3F7"
MARCA = "E4002B"
SERIE, SERIE_SUAVE, GRIS = "2A78D6", "BFD6F3", "9AA4B2"
CATEGORICA = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7"]
ESTADO = {  # tono → (texto, fondo)
    "bueno": ("0F7A4E", "E7F7EF"), "medio": ("A15C04", "FDF2E2"), "malo": ("C8001F", "FDEAEE"),
    "info": (MUTED, PANEL),
}
_ESTADO_PLAN = {"critico": ("Crítico", "malo"), "atencion": ("Atención", "medio"), "mejora": ("Mejora", "bueno")}
_TIPO_HALLAZGO = {"positive": ("Favorable", "bueno"), "warning": ("Atención", "medio"),
                  "negative": ("Negativo", "malo"), "info": ("Informativo", "info")}
_SEVERIDAD = {"alta": "malo", "media": "medio", "baja": "info", "oportunidad": "bueno"}

FMT_ENTERO, FMT_DECIMAL = "#,##0", "#,##0.00"
FMT_PCT, FMT_PCT_SIGNO = "0.0%", "+0.0%;-0.0%;0.0%"
FMT_SIGNO = "+#,##0;-#,##0;0"
FMT_FECHA = "dd/mm/yyyy"

MAX_FILAS_EXCEL = 1_048_575
MAX_POR_DIMENSION = 500     # filas de «cómo va cada uno» (el resto se resume)
MAX_ITEMS_DINAMICA = 300    # una dimensión con más valores no sirve de eje en la dinámica
LARGO_CELDA = 32_000

_fino = Side(style="thin", color=LINEA)
_BORDE = Border(bottom=_fino)
_CAJA = Border(left=_fino, right=_fino, top=_fino, bottom=_fino)


# ── Utilidades de texto y formato ─────────────────────────────────────────

def _texto(valor) -> str:
    """Texto limpio para una celda: sin HTML, sin negritas markdown y sin los
    caracteres de control que Excel rechaza (y que harían fallar el guardado)."""
    t = clean_display_text(valor).replace("**", "")
    return ILLEGAL_CHARACTERS_RE.sub("", t)[:LARGO_CELDA]


def _label(schema: dict, columna) -> str:
    for item in schema.get("semantic", {}).get("columns", []):
        if item.get("column") == columna:
            return str(item.get("display_name") or columna)
    return str(columna)


def _compacto(valor) -> str:
    if valor is None or (isinstance(valor, float) and not math.isfinite(valor)):
        return "—"
    v = float(valor)
    a = abs(v)
    if a >= 1e9:
        return f"{v / 1e9:.1f}B"
    if a >= 1e6:
        return f"{v / 1e6:.1f}M"
    if a >= 1e3:
        return f"{v / 1e3:.1f}K"
    return f"{v:,.0f}" if float(v).is_integer() else f"{v:,.2f}"


def _formato_compacto(maximo: float) -> str:
    """Formato de celda que muestra 559.3M sin perder el valor exacto: quien
    lo necesite completo lo ve en la barra de fórmulas o cambia el formato."""
    a = abs(maximo or 0)
    if a >= 1e9:
        return '#,##0.0,,,"B"'
    if a >= 1e6:
        return '#,##0.0,,"M"'
    if a >= 1e5:
        return '#,##0,"K"'
    return FMT_ENTERO


def _formato_valores(valores) -> str:
    s = pd.to_numeric(pd.Series(list(valores)), errors="coerce").dropna()
    if s.empty:
        return FMT_ENTERO
    if float((s % 1 == 0).mean()) >= 0.98 or float(s.abs().median()) >= 1000:
        return FMT_ENTERO
    return FMT_DECIMAL


def _nombre_meta(columna: str) -> str:
    """«Meta» o «Presupuesto» tal cual; otro nombre, aclarado: «Meta (Cuota)»."""
    return columna if re.search(r"meta|presupuesto|objetivo|cuota|budget", columna, re.I) else f"Meta ({columna})"


def _nombre_hoja(wb: Workbook, base: str) -> str:
    limpio = re.sub(r"[\[\]:*?/\\']", " ", str(base)).strip() or "Hoja"
    limpio = re.sub(r"\s+", " ", limpio)[:31]
    nombre, i = limpio, 2
    usados = {ws.title.lower() for ws in wb.worksheets}
    while nombre.lower() in usados:
        sufijo = f" ({i})"
        nombre = limpio[:31 - len(sufijo)] + sufijo
        i += 1
    return nombre


def _num(v):
    """Número de Python para una celda, o None si no es un número finito."""
    if v is None or isinstance(v, bool):
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if math.isfinite(f) else None


def _alto_fila(textos_y_anchos, minimo: float = 18.0) -> float:
    """Alto aproximado para que el texto ajustado se lea completo: Excel no
    recalcula el alto de una fila escrita desde fuera al abrir el archivo."""
    lineas = 1
    for texto, ancho in textos_y_anchos:
        if not texto:
            continue
        por_linea = max(8, int(ancho * 1.15))
        n = sum(max(1, math.ceil(len(parte) / por_linea)) for parte in str(texto).split("\n"))
        lineas = max(lineas, n)
    return max(minimo, min(400.0, lineas * 15.0 + 4))


# ── Piezas de hoja ────────────────────────────────────────────────────────

def _hoja(wb: Workbook, nombre: str, titulo: str, subtitulo: str = "", color: str = TINTA, ancho: int = 10,
          volver: bool = True):
    ws = wb.create_sheet(_nombre_hoja(wb, nombre))
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = color
    ws.column_dimensions["A"].width = 2
    ws.row_dimensions[1].height = 8
    c = ws.cell(row=2, column=2, value=titulo)
    c.font = Font(size=16, bold=True, color=TINTA)
    ws.row_dimensions[2].height = 26
    if subtitulo:
        s = ws.cell(row=3, column=2, value=subtitulo)
        s.font = Font(size=10, italic=True, color=MUTED)
    if volver:
        v = ws.cell(row=2, column=ancho + 1, value="← Volver al resumen")
        v.hyperlink = "#'Resumen'!A1"
        v.font = Font(size=9, color=SERIE, underline="single")
        v.alignment = Alignment(horizontal="right")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _seccion(ws, fila: int, texto: str, nota: str = "", ancho: int = 10) -> int:
    """Título de un bloque con su nota «cómo leerlo». Devuelve la fila siguiente."""
    c = ws.cell(row=fila, column=2, value=texto)
    c.font = Font(size=12, bold=True, color=TINTA)
    for col in range(2, 2 + ancho):
        ws.cell(row=fila, column=col).border = Border(bottom=Side(style="medium", color=MARCA))
    fila += 1
    if nota:
        n = ws.cell(row=fila, column=2, value=_texto(nota))
        n.font = Font(size=9, italic=True, color=MUTED)
        n.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=1 + ancho)
        ws.row_dimensions[fila].height = _alto_fila([(nota, ancho * 11)], 15)
        fila += 1
    return fila + 1


def _tabla(ws, fila: int, encabezados: list[str], filas: list[list], formatos: dict | None = None,
           anchos: dict | None = None, envolver: set | None = None, total: list | None = None,
           columna: int = 2) -> tuple[int, int]:
    """Tabla con encabezado oscuro, filas cebra y bordes suaves.

    `formatos` e `envolver` van por índice de columna (0 = primera). Devuelve
    (fila del encabezado, última fila escrita)."""
    formatos, envolver = formatos or {}, envolver or set()
    for j, h in enumerate(encabezados):
        c = ws.cell(row=fila, column=columna + j, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=TINTA)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if anchos and j in anchos:
            # El ancho más grande que pida cualquier tabla de la hoja; el
            # predeterminado de openpyxl (13) no cuenta, o una columna «N.º»
            # nunca quedaría angosta.
            letra = get_column_letter(columna + j)
            fijados = ws.__dict__.setdefault("_anchos_tabla", {})
            fijados[letra] = max(fijados.get(letra, 0), anchos[j])
            ws.column_dimensions[letra].width = fijados[letra]
    ws.row_dimensions[fila].height = _alto_fila([(h, (anchos or {}).get(j, 12)) for j, h in enumerate(encabezados)], 20)
    cabecera = fila
    for i, valores in enumerate(filas):
        fila += 1
        textos = []
        for j, v in enumerate(valores):
            c = ws.cell(row=fila, column=columna + j, value=v)
            c.border = _BORDE
            if i % 2:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
            if j in formatos and isinstance(v, (int, float)):
                c.number_format = formatos[j]
            if j in envolver:
                c.alignment = Alignment(wrap_text=True, vertical="top")
                textos.append((v, (anchos or {}).get(j, 12)))
            else:
                c.alignment = Alignment(vertical="top")
        if textos:
            ws.row_dimensions[fila].height = _alto_fila(textos)
    if total:
        fila += 1
        for j, v in enumerate(total):
            c = ws.cell(row=fila, column=columna + j, value=v)
            c.font = Font(bold=True, color=TINTA)
            c.fill = PatternFill("solid", fgColor=PANEL)
            c.border = Border(top=Side(style="thin", color=TINTA))
            if j in formatos and isinstance(v, (int, float)):
                c.number_format = formatos[j]
    return cabecera, fila


def _pintar_estado(celda, tono: str) -> None:
    color, fondo = ESTADO.get(tono, ESTADO["info"])
    celda.font = Font(bold=True, color=color, size=10)
    celda.fill = PatternFill("solid", fgColor=fondo)
    celda.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)


def _colorear_variacion(ws, rango: str, menos_es_mejor: bool = False) -> None:
    """Verde si sube y rojo si baja (al revés si menos es mejor). El signo va
    en el número, así que el color nunca es la única pista."""
    sube, baja = ("C8001F", "0F7A4E") if menos_es_mejor else ("0F7A4E", "C8001F")
    ws.conditional_formatting.add(rango, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=sube, bold=True)))
    ws.conditional_formatting.add(rango, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=baja, bold=True)))


def _nota(ws, fila: int, texto: str, ancho: int = 10, color: str = MUTED, italica: bool = True) -> int:
    c = ws.cell(row=fila, column=2, value=_texto(texto))
    c.font = Font(size=9, italic=italica, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=1 + ancho)
    ws.row_dimensions[fila].height = _alto_fila([(texto, ancho * 11)], 15)
    return fila + 1


# ── Gráficos ──────────────────────────────────────────────────────────────

def _estilo_ejes(chart, formato_valor: str | None = None) -> None:
    # openpyxl 3.1 marca los ejes como borrados por defecto: sin esto, Excel
    # dibuja el gráfico sin etiquetas en los ejes.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    if formato_valor:
        chart.y_axis.number_format = formato_valor
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E8EBF1"))
    chart.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=LINEA))
    chart.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))


def _color_serie(serie, color: str, linea: bool = False) -> None:
    serie.graphicalProperties.line.solidFill = color
    if linea:
        serie.graphicalProperties.line.width = 28575  # 2.25 pt
        serie.marker.symbol = "circle"
        serie.marker.size = 7
        serie.marker.graphicalProperties = GraphicalProperties(solidFill=color)
        serie.marker.graphicalProperties.line.solidFill = "FFFFFF"
        serie.smooth = False
    else:
        serie.graphicalProperties.solidFill = color


def _grafico_linea(ws, titulo: str, cats: Reference, series: list[tuple[Reference, str]], formato: str,
                   ancho: float = 17, alto: float = 8.5):
    ch = LineChart()
    ch.title = titulo
    ch.height, ch.width = alto, ancho
    for ref, color in series:
        ch.add_data(ref, titles_from_data=True)
        _color_serie(ch.series[-1], color, linea=True)
    ch.set_categories(cats)
    _estilo_ejes(ch, formato)
    if len(series) == 1:
        ch.legend = None
    else:
        ch.legend.position = "r"  # abajo se encima con los meses y arriba con el título
    return ch


def _grafico_barras(ws, titulo: str, cats: Reference, series: list[tuple[Reference, str]], formato: str,
                    horizontal: bool = False, ancho: float = 17, alto: float = 8.5, etiquetas: bool = False,
                    min_cero: bool = False):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.title = titulo
    ch.height, ch.width = alto, ancho
    ch.gapWidth = 60
    for ref, color in series:
        ch.add_data(ref, titles_from_data=True)
        _color_serie(ch.series[-1], color)
    ch.set_categories(cats)
    _estilo_ejes(ch, formato)
    if horizontal:
        # El primero del ranking arriba, como en la tabla.
        ch.x_axis.scaling.orientation = "maxMin"
        ch.y_axis.crosses = "max"
    for serie in ch.series:
        serie.invertIfNegative = False  # si no, Excel dibuja huecas las barras negativas
    # Las etiquetas de categoría al borde del gráfico, no sobre el cero: con
    # valores negativos quedaban encima de las barras.
    ch.x_axis.tickLblPos = "low"
    if min_cero:
        ch.y_axis.scaling.min = 0
    if etiquetas:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        ch.dataLabels.showCatName = False
        ch.dataLabels.showSerName = False
        ch.dataLabels.showLegendKey = False
        ch.dataLabels.showPercent = False
        ch.dataLabels.numFmt = formato
    if len(series) == 1:
        ch.legend = None
    else:
        ch.legend.position = "b"
        ch.overlap = -10
    return ch


def _grafico_dona(titulo: str, cats: Reference, datos: Reference, n: int, ancho: float = 12, alto: float = 8.5):
    ch = DoughnutChart()
    ch.title = titulo
    ch.height, ch.width = alto, ancho
    ch.holeSize = 55
    ch.add_data(datos, titles_from_data=True)
    ch.set_categories(cats)
    serie = ch.series[0]
    for i in range(n):
        punto = DataPoint(idx=i)
        # "Otros" siempre gris; el resto en el orden fijo de la paleta.
        color = GRIS if i == n - 1 and n > len(CATEGORICA) - 1 else CATEGORICA[i % len(CATEGORICA)]
        punto.graphicalProperties.solidFill = color
        punto.graphicalProperties.line.solidFill = "FFFFFF"
        serie.dPt.append(punto)
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showPercent = True
    ch.dataLabels.showVal = False
    ch.dataLabels.showCatName = False
    ch.dataLabels.showSerName = False
    ch.dataLabels.showLeaderLines = False
    ch.legend.position = "r"
    return ch


# ── Contexto del análisis ─────────────────────────────────────────────────

def _contexto(df: pd.DataFrame, schema: dict, dashboard: dict, filename: str, sheet: str, alcance: str) -> dict:
    metrics = metric_candidates(df, schema)
    primary = dashboard.get("primary_metric") or (metrics[0] if metrics else None)
    if primary not in df.columns:
        primary = None
    calculo = calculo_automatico(df, schema, primary) if primary else "Conteo"
    ctx = {
        "archivo": str(filename), "hoja": str(sheet), "alcance": alcance or "Hoja completa",
        "registros": len(df), "primary": primary, "calculo": calculo,
        "metrica": _label(schema, primary) if primary else "Cantidad de registros",
        "fecha_col": columna_fecha(df, schema), "desde": None, "hasta": None,
        "generado": datetime.now().strftime("%d/%m/%Y %H:%M"),
    }
    if ctx["fecha_col"]:
        fechas = a_datetime(df[ctx["fecha_col"]], errors="coerce").dropna()
        if not fechas.empty:
            ctx["desde"], ctx["hasta"] = fechas.min(), fechas.max()
    ctx["periodo"] = (f"del {fecha_larga(ctx['desde'])} al {fecha_larga(ctx['hasta'])}"
                      if ctx["desde"] is not None else "Sin fechas en el archivo")
    valores = numeric_valid(df[primary]).dropna() if primary else pd.Series(dtype=float)
    ctx["formato"] = _formato_valores(valores) if primary else FMT_ENTERO
    return ctx


def _valor_serie(df: pd.DataFrame, primary) -> pd.Series:
    return numeric_valid(df[primary]) if primary else pd.Series(1.0, index=df.index)


def _agregar(serie_agrupada, calculo: str):
    if calculo == "Promedio":
        return serie_agrupada.mean()
    if calculo == "Conteo":
        return serie_agrupada.size()
    return serie_agrupada.sum(min_count=1)


# ── Hojas ─────────────────────────────────────────────────────────────────

def _hoja_tendencia(wb, df, schema, ctx, dashboard) -> dict | None:
    """Mes a mes: la métrica principal, su variación, el acumulado, la meta y
    las demás métricas. Devuelve dónde quedó la tabla para el gráfico de la
    portada."""
    col = ctx["fecha_col"]
    if not col:
        return None
    periodos = a_datetime(df[col], errors="coerce").dt.to_period("M")
    primary, calculo = ctx["primary"], ctx["calculo"]
    base = pd.DataFrame({"_p": periodos, "_v": _valor_serie(df, primary)}).dropna(subset=["_p"])
    if calculo != "Conteo":
        # Un mes sin dato no es un mes en cero (ver dashboard_engine.growth).
        serie = _agregar(base.dropna(subset=["_v"]).groupby("_p")["_v"], calculo)
    else:
        serie = base.groupby("_p")["_v"].size()
    serie = serie.sort_index().dropna()
    if len(serie) < 2:
        return None
    registros = base.groupby("_p").size().reindex(serie.index).fillna(0)

    meta_col = None
    try:
        from core.performance import columna_meta
        meta_col = columna_meta(df, schema, primary) if primary else None
    except Exception:
        meta_col = None
    metas = None
    if meta_col is not None and meta_col in df.columns:
        m = pd.DataFrame({"_p": periodos, "_m": numeric_valid(df[meta_col])}).dropna()
        metas = _agregar(m.groupby("_p")["_m"], "Promedio" if calculo == "Promedio" else "Suma").reindex(serie.index)

    otras = [c for c in metric_candidates(df, schema) if c not in {primary, meta_col}][:3]
    sumable = calculo in {"Suma", "Conteo"}

    ws = _hoja(wb, "Tendencia", f"Evolución mes a mes · {ctx['metrica']}",
               f"{'Suma' if sumable else 'Promedio'} de «{ctx['metrica']}» por mes, {ctx['periodo']}. "
               f"Fuente: columna «{col}».", color=SERIE, ancho=10)
    fila = 5
    parcial = False
    try:
        parcial = ultimo_incompleto(df, schema, primary, calculo, "Mes")
    except Exception:
        pass
    nota = ("Variación = cambio frente al mes anterior. Verde sube, rojo baja; el signo acompaña al número."
            + (" El acumulado suma los meses desde el primero." if sumable else "")
            + (" ⚠ El último mes parece incompleto (vale mucho menos que los anteriores o tiene menos días): "
               "no lo leas como una caída sin confirmarlo." if parcial else ""))
    fila = _seccion(ws, fila, "Tabla mensual", nota)

    encabezados = ["Mes", ctx["metrica"], "Variación", "Variación %"]
    formatos = {1: ctx["formato"], 2: FMT_SIGNO, 3: FMT_PCT_SIGNO}
    if sumable:
        encabezados.append("Acumulado")
        formatos[len(encabezados) - 1] = ctx["formato"]
    idx_meta = None
    if metas is not None and metas.notna().sum() >= 2:
        idx_meta = len(encabezados)
        encabezados += [_nombre_meta(_label(schema, meta_col)), "Cumplimiento"]
        formatos[idx_meta], formatos[idx_meta + 1] = ctx["formato"], FMT_PCT
    encabezados.append("Registros")
    formatos[len(encabezados) - 1] = FMT_ENTERO
    inicio_otras = len(encabezados)
    for c in otras:
        encabezados.append(_label(schema, c))
        formatos[len(encabezados) - 1] = _formato_valores(numeric_valid(df[c]).dropna())
    otras_series = {}
    for c in otras:
        o = pd.DataFrame({"_p": periodos, "_v": numeric_valid(df[c])}).dropna()
        otras_series[c] = _agregar(o.groupby("_p")["_v"], calculo_automatico(df, schema, c)).reindex(serie.index)

    filas, acumulado, previo = [], 0.0, None
    for p, v in serie.items():
        v = float(v)
        acumulado += v
        var = v - previo if previo is not None else None
        var_pct = (var / abs(previo)) if (previo not in (None, 0) and var is not None) else None
        fila_v = [mes(p.to_timestamp(), True), v, var, var_pct]
        if sumable:
            fila_v.append(acumulado)
        if idx_meta is not None:
            mt = _num(metas.get(p))
            fila_v += [mt, (v / mt) if mt else None]
        fila_v.append(int(registros.get(p, 0)))
        for c in otras:
            fila_v.append(_num(otras_series[c].get(p)))
        filas.append(fila_v)
        previo = v
    anchos = {0: 12, 1: 16, 2: 14, 3: 12, 4: 16}
    anchos.update({j: 15 for j in range(5, len(encabezados))})
    total = None
    if sumable:
        total = ["Total", float(serie.sum()), None, None, None]
        if idx_meta is not None:
            mt = float(metas.sum())
            total += [mt, float(serie.reindex(metas.dropna().index).sum()) / mt if mt else None]
        total.append(int(registros.sum()))
        total += [None] * len(otras)
    cab, ultima = _tabla(ws, fila, encabezados, filas, formatos, anchos, total=total)
    primera_dato, ultima_dato = cab + 1, cab + len(filas)
    _colorear_variacion(ws, f"D{primera_dato}:E{ultima_dato}")
    if idx_meta is not None:
        letra = get_column_letter(2 + idx_meta + 1)
        _cumplimiento_condicional(ws, f"{letra}{primera_dato}:{letra}{ultima_dato}")

    # Gráficos a la derecha de la tabla.
    ancla = get_column_letter(2 + len(encabezados) + 1)
    cats = Reference(ws, min_col=2, min_row=primera_dato, max_row=ultima_dato)
    series = [(Reference(ws, min_col=3, min_row=cab, max_row=ultima_dato), SERIE)]
    if idx_meta is not None:
        series.append((Reference(ws, min_col=2 + idx_meta, min_row=cab, max_row=ultima_dato), GRIS))
    formato_eje = _formato_compacto(float(serie.abs().max()))
    titulo = f"{ctx['metrica']} por mes" + (" frente a la meta" if idx_meta is not None else "")
    ws.add_chart(_grafico_linea(ws, titulo, cats, series, formato_eje), f"{ancla}5")
    var_ref = Reference(ws, min_col=5, min_row=cab, max_row=ultima_dato)
    ws.add_chart(_grafico_barras(ws, "Variación % frente al mes anterior", cats, [(var_ref, SERIE)], "0%"),
                 f"{ancla}23")
    if inicio_otras < len(encabezados):
        _nota(ws, ultima + 2, "Las otras métricas se suman o se promedian según lo que son (un precio o un "
                              "porcentaje se promedia, no se suma).")
    return {"hoja": ws.title, "cab": cab, "ultima": ultima_dato, "formato_eje": formato_eje, "titulo": titulo,
            "serie_meta": idx_meta is not None, "col_meta": 2 + idx_meta if idx_meta is not None else None}


def _cumplimiento_condicional(ws, rango: str) -> None:
    corte = CERCA_DE_META / 100
    ws.conditional_formatting.add(rango, CellIsRule(operator="greaterThanOrEqual", formula=["1"],
                                                    font=Font(color=ESTADO["bueno"][0], bold=True),
                                                    fill=PatternFill("solid", fgColor=ESTADO["bueno"][1])))
    ws.conditional_formatting.add(rango, CellIsRule(operator="between", formula=[str(corte), "0.99999"],
                                                    font=Font(color=ESTADO["medio"][0], bold=True),
                                                    fill=PatternFill("solid", fgColor=ESTADO["medio"][1])))
    ws.conditional_formatting.add(rango, CellIsRule(operator="lessThan", formula=[str(corte)],
                                                    font=Font(color=ESTADO["malo"][0], bold=True),
                                                    fill=PatternFill("solid", fgColor=ESTADO["malo"][1])))


def _hoja_por_dimension(wb, df, schema, ctx, dimension) -> dict | None:
    """Cómo va cada uno en una dimensión. Usa core/cuadro_comparativo con
    TODOS los valores elegidos, así posición, promedio, participación y
    cumplimiento son del grupo completo (la regla del dominio)."""
    primary = ctx["primary"]
    previo = cuadro_comparativo(df, schema, dimension, primary)
    if not previo:
        return None
    todos = previo["opciones"]
    cuadro = cuadro_comparativo(df, schema, dimension, primary, seleccion=todos[:MAX_POR_DIMENSION])
    if not cuadro or not cuadro["filas"]:
        return None
    dim_label = _label(schema, dimension)
    n_total = cuadro["total_grupo"]
    usa_meta = cuadro["base"] == "meta"
    metrica = "Cantidad de registros" if cuadro["conteo"] else ctx["metrica"]
    formato = FMT_ENTERO if cuadro["conteo"] else ctx["formato"]

    ws = _hoja(wb, f"Por {dim_label}", f"Cómo va cada {dim_label.lower()}",
               f"{metrica} de los {n_total:,} valores de «{dim_label}» · {ctx['periodo']}",
               color=CATEGORICA[1], ancho=12)
    fila = 5
    referencia = (f"Cada uno contra su propia meta: {CERCA_DE_META:.0f}% o más es «cerca de la meta». "
                  if usa_meta else
                  f"Sin meta, cada uno se compara con el promedio de los {n_total:,} "
                  f"({_compacto(cuadro['promedio'])}). ")
    fila = _seccion(ws, fila, "Lectura", referencia + "Posición, participación y promedio se calculan sobre "
                                                     f"los {n_total:,}, no sobre una selección.", ancho=12)
    for frase in (cuadro.get("lectura") or [])[:6]:
        fila = _nota(ws, fila, f"•  {frase}", ancho=12, color=TINTA, italica=False)
    if cuadro.get("aviso"):
        fila = _nota(ws, fila, f"⚠ {cuadro['aviso']}", ancho=12, color=ESTADO["medio"][0], italica=False)
    fila += 1

    periodo_txt = (f"Var. {cuadro['periodo_label']} vs {cuadro['periodo_anterior_label']}"
                   if cuadro.get("periodo_label") else None)
    encabezados = ["Posición", dim_label, metrica]
    formatos = {0: FMT_ENTERO, 2: formato}
    anchos = {0: 9, 1: 26, 2: 16}
    if usa_meta:
        encabezados += [_nombre_meta(_label(schema, cuadro["meta_col"])), "Cumplimiento"]
        formatos.update({3: formato, 4: FMT_PCT})
        anchos.update({3: 16, 4: 13})
    i_part = None
    if cuadro["aditiva"]:
        i_part = len(encabezados)
        encabezados.append(f"Participación (de {n_total:,})")
        formatos[i_part] = FMT_PCT
        anchos[i_part] = 14
    i_vs = len(encabezados)
    encabezados.append(f"vs promedio de los {n_total:,}")
    formatos[i_vs] = FMT_PCT_SIGNO
    anchos[i_vs] = 14
    i_var = None
    if periodo_txt:
        i_var = len(encabezados)
        encabezados.append(periodo_txt)
        formatos[i_var] = FMT_PCT_SIGNO
        anchos[i_var] = 16
    i_reg = len(encabezados)
    encabezados.append("Registros")
    formatos[i_reg] = FMT_ENTERO
    anchos[i_reg] = 11
    if cuadro["aditiva"] and not cuadro["conteo"]:
        encabezados.append("Promedio por registro")
        formatos[len(encabezados) - 1] = formato
        anchos[len(encabezados) - 1] = 14
    i_estado = len(encabezados)
    encabezados.append("Estado")
    anchos[i_estado] = 20

    def _pct(x):
        return None if x is None else x / 100

    filas = []
    for f in cuadro["filas"]:
        v = [f.get("posicion_general") or f["posicion"], f["nombre"], f["valor"]]
        if usa_meta:
            v += [f.get("meta"), _pct(f.get("cumplimiento"))]
        if i_part is not None:
            v.append(_pct(f.get("participacion")))
        v.append(_pct(f.get("vs_promedio")))
        if i_var is not None:
            v.append(_pct(f.get("variacion")))
        v.append(f["registros"])
        if cuadro["aditiva"] and not cuadro["conteo"]:
            v.append(f.get("por_registro"))
        v.append(f["estado"])
        filas.append(v)
    total = None
    if cuadro["aditiva"]:
        total = [None, f"Total de los {n_total:,}", float(sum(f["valor"] for f in cuadro["filas"]))
                 if len(cuadro["filas"]) == n_total else None]
        if usa_meta:
            total += [None, _pct(cuadro.get("cumplimiento_grupo"))]
        total += [None] * (len(encabezados) - len(total))
    cab, ultima = _tabla(ws, fila, encabezados, filas, formatos, anchos, total=total)
    primera_dato, ultima_dato = cab + 1, cab + len(filas)
    for i, f in enumerate(cuadro["filas"]):
        _pintar_estado(ws.cell(row=primera_dato + i, column=2 + i_estado), f.get("tono", "info"))
    letra_valor = get_column_letter(4)
    ws.conditional_formatting.add(f"{letra_valor}{primera_dato}:{letra_valor}{ultima_dato}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max", color=SERIE_SUAVE))
    if usa_meta:
        _cumplimiento_condicional(ws, f"F{primera_dato}:F{ultima_dato}")
    for j in (i_vs, i_var):
        if j is not None:
            letra = get_column_letter(2 + j)
            _colorear_variacion(ws, f"{letra}{primera_dato}:{letra}{ultima_dato}", cuadro["menos_es_mejor"])
    ws.auto_filter.ref = f"B{cab}:{get_column_letter(1 + len(encabezados))}{ultima_dato}"
    if n_total > len(filas):
        _nota(ws, ultima + 2, f"Se listan los {len(filas):,} primeros de {n_total:,}; el total y las referencias "
                              f"sí incluyen a todos.", ancho=12)

    # Gráfico: los 15 primeros, por cumplimiento si hay meta.
    ancla_col = get_column_letter(2 + len(encabezados) + 1)
    top = min(15, len(filas))
    cats = Reference(ws, min_col=3, min_row=primera_dato, max_row=primera_dato + top - 1)
    if usa_meta:
        ref = Reference(ws, min_col=6, min_row=cab, max_row=primera_dato + top - 1)
        titulo, fmt = f"Cumplimiento de meta · {min(top, n_total)} primeros de {n_total:,}", "0%"
    else:
        ref = Reference(ws, min_col=4, min_row=cab, max_row=primera_dato + top - 1)
        titulo, fmt = f"{metrica} · {min(top, n_total)} primeros de {n_total:,}", _formato_compacto(filas[0][2] or 0)
    ch = _grafico_barras(ws, titulo, cats, [(ref, SERIE)], fmt, horizontal=True, alto=max(7.5, top * 0.62),
                         etiquetas=True, min_cero=all((f[2] or 0) >= 0 for f in filas))
    ws.add_chart(ch, f"{ancla_col}{cab}")

    # Participación: los 6 primeros + «Otros», en una tabla auxiliar visible.
    if cuadro["aditiva"] and not cuadro["conteo"] and len(cuadro["filas"]) >= 3:
        ordenadas = sorted(cuadro["filas"], key=lambda f: -f["valor"])
        top6 = ordenadas[:6]
        resto = float(sum(f["valor"] for f in ordenadas[6:]))
        filas_p = [[f["nombre"], f["valor"]] for f in top6] + ([["Otros", resto]] if resto > 0 else [])
        if all((v or 0) >= 0 for _, v in filas_p):
            fila_p = ultima + 4
            fila_p = _seccion(ws, fila_p, f"Participación en el total · {metrica}", ancho=3)
            cab_p, fin_p = _tabla(ws, fila_p, [dim_label, metrica], filas_p, {1: formato})
            ws.add_chart(_grafico_dona(f"Participación por {dim_label.lower()}",
                                       Reference(ws, min_col=2, min_row=cab_p + 1, max_row=fin_p),
                                       Reference(ws, min_col=3, min_row=cab_p, max_row=fin_p), len(filas_p)),
                         f"E{cab_p}")
    return {"hoja": ws.title, "primera": primera_dato, "top": top, "usa_meta": usa_meta, "titulo": titulo,
            "formato": fmt, "dimension": dim_label, "min_cero": all((f[2] or 0) >= 0 for f in filas)}


def _hoja_matriz(wb, df, schema, ctx, dims) -> None:
    """Cruces tipo tabla dinámica ya calculados, con mapa de calor: dimensión
    × mes y, si hay dos dimensiones, una contra la otra."""
    cruces = []
    if ctx["fecha_col"] and dims:
        cruces.append((dims[0], PERIODO, f"{_label(schema, dims[0])} × mes"))
    if len(dims) >= 2:
        cruces.append((dims[0], dims[1], f"{_label(schema, dims[0])} × {_label(schema, dims[1])}"))
    if ctx["fecha_col"] and len(dims) >= 2:
        cruces.append((dims[1], PERIODO, f"{_label(schema, dims[1])} × mes"))
    resultados = []
    for filas, columnas, titulo in cruces:
        try:
            r = tabla_cruzada(df, schema, filas, columnas, ctx["primary"], max_filas=40, max_columnas=24)
        except Exception:
            r = None
        if r is not None:
            resultados.append((filas, titulo, r))
    if not resultados:
        return
    ws = _hoja(wb, "Matriz", "Matrices cruzadas",
               f"{ctx['metrica']} ({resultados[0][2]['calculo'].lower()}) cruzada por dimensiones y por mes. "
               "Color más intenso = valor más alto.", color=CATEGORICA[2], ancho=14)
    fila = 5
    for dim_filas, titulo, r in resultados:
        matriz = r["matriz"]
        sumable = r["sumable"]
        fila = _seccion(ws, fila, titulo, " · ".join(_texto(h) for h in r.get("hallazgos", [])[:3]), ancho=14)
        encabezados = [_label(schema, dim_filas)] + [str(c) for c in matriz.columns] + (["Total"] if sumable else [])
        # Una combinación sin actividad se muestra como «–»: una matriz llena
        # de ceros esconde las cifras que importan.
        sin_cero = ctx["formato"] + ';-' + ctx["formato"] + ';"–"'
        formatos = {j: sin_cero for j in range(1, len(encabezados))}
        filas = []
        for nombre, valores in matriz.iterrows():
            v = [str(nombre)] + [_num(x) for x in valores.tolist()]
            if sumable:
                v.append(_num(valores.sum()))
            filas.append(v)
        total = None
        if sumable:
            total = ["Total"] + [_num(matriz[c].sum()) for c in matriz.columns] + [_num(matriz.values.sum())]
        anchos = {0: 24}
        anchos.update({j: 12 for j in range(1, len(encabezados))})
        cab, ultima = _tabla(ws, fila, encabezados, filas, formatos, anchos, total=total)
        fin_col = get_column_letter(1 + len(matriz.columns) + 1)
        ws.conditional_formatting.add(
            f"C{cab + 1}:{fin_col}{cab + len(filas)}",
            ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="8DB9EC"))
        if sumable:
            letra_t = get_column_letter(2 + len(encabezados) - 1)
            ws.conditional_formatting.add(f"{letra_t}{cab + 1}:{letra_t}{cab + len(filas)}",
                                          DataBarRule(start_type="num", start_value=0, end_type="max", color=SERIE_SUAVE))
        fila = ultima + 3


def _hoja_plan(wb, plan: dict, ctx) -> dict | None:
    planes = (plan or {}).get("planes") or []
    if not planes:
        return None
    ws = _hoja(wb, "Plan de acción", "Plan de acción",
               _texto(plan.get("resumen") or "Frentes a atender, en orden de urgencia."), color=MARCA, ancho=10)
    fila = _seccion(ws, 5, "Frentes en orden de atención",
                    "Crítico = actuar ya · Atención = vigilar y corregir · Mejora = oportunidad. «Valor estimado» es "
                    "lo que vale cerrar la brecha en un mes, calculado con la historia de cada caso.", ancho=10)
    encabezados = ["N.º", "Estado", "Frente", "Situación", "Por qué importa", "Qué hacer", "Cómo se mide",
                   "Control semanal", "Alarma", "Valor estimado"]
    anchos = {0: 5, 1: 11, 2: 24, 3: 38, 4: 32, 5: 58, 6: 30, 7: 34, 8: 34, 9: 14}
    filas = []
    for i, p in enumerate(planes, 1):
        pasos = "\n".join(f"{k}. {_texto(x)}" for k, x in enumerate(p.get("pasos") or [], 1))
        filas.append([i, _ESTADO_PLAN.get(p.get("estado"), (str(p.get("estado") or ""), "info"))[0],
                      _texto(p.get("titulo")), _texto(p.get("situacion")), _texto(p.get("por_que")), pasos,
                      _texto(p.get("medir")), _texto(p.get("control")), _texto(p.get("alarma")),
                      _num(p.get("impacto"))])
    cab, ultima = _tabla(ws, fila, encabezados, filas, {9: ctx["formato"]}, anchos,
                         envolver={2, 3, 4, 5, 6, 7, 8})
    for i, p in enumerate(planes):
        _pintar_estado(ws.cell(row=cab + 1 + i, column=3), _ESTADO_PLAN.get(p.get("estado"), ("", "info"))[1])
        ws.cell(row=cab + 1 + i, column=4).font = Font(bold=True, color=TINTA)

    casos = [(p, c) for p in planes for c in (p.get("casos") or [])]
    if casos:
        fila = _seccion(ws, ultima + 3, "Metas por caso",
                        "Objetivo del próximo mes para cada caso, con su brecha, el ritmo semanal para lograrlo y si "
                        "es factible según su propia historia.", ancho=10)
        enc = ["Frente", "Caso", "Actual", "Objetivo", "Brecha", "Por semana", "Referencia", "Factibilidad", "Lectura"]
        filas = [[_texto(p.get("titulo")), _texto(c.get("nombre")), _num(c.get("actual")), _num(c.get("objetivo")),
                  _num(c.get("brecha")), _num(c.get("semanal")), _texto(c.get("referencia")),
                  _texto(c.get("factibilidad_txt") or c.get("factibilidad")), _texto(c.get("lectura"))]
                 for p, c in casos]
        f = ctx["formato"]
        _tabla(ws, fila, enc, filas, {2: f, 3: f, 4: f, 5: f}, envolver={0, 6, 7, 8},
               anchos={0: 24, 1: 22, 2: 14, 3: 14, 4: 14, 5: 14, 6: 26, 7: 34, 8: 58})
    return {"hoja": ws.title}


def _hoja_hallazgos(wb, dashboard, ctx) -> None:
    alertas = [a for a in (dashboard.get("alerts") or []) if isinstance(a, dict)]
    hallazgos = [h for h in (dashboard.get("insights") or []) if isinstance(h, dict)]
    if not alertas and not hallazgos:
        return
    ws = _hoja(wb, "Hallazgos", "Alertas y hallazgos",
               "Lo que detectó el análisis automático, con la evidencia y qué revisar.", color=CATEGORICA[3], ancho=8)
    fila = 5
    if alertas:
        fila = _seccion(ws, fila, "Alertas", "Ordenadas por urgencia. «Impacto» es el monto en juego cuando el "
                                             "análisis lo pudo calcular.", ancho=8)
        filas = [[_texto(a.get("severity") or "Media"), _texto(a.get("title")), _texto(a.get("text") or a.get("message")),
                  _texto(a.get("action")), _texto(a.get("implication")), _num(a.get("impacto"))] for a in alertas]
        cab, fila = _tabla(ws, fila, ["Severidad", "Alerta", "Detalle", "Qué hacer", "Implicación", "Impacto"], filas,
                           {5: ctx["formato"]}, {0: 12, 1: 26, 2: 50, 3: 42, 4: 42, 5: 14}, envolver={1, 2, 3, 4})
        for i, a in enumerate(alertas):
            _pintar_estado(ws.cell(row=cab + 1 + i, column=2),
                           _SEVERIDAD.get(str(a.get("severity") or "").lower(), "medio"))
        fila += 3
    if hallazgos:
        fila = _seccion(ws, fila, "Hallazgos", "Favorable, atención o informativo, con los nombres y cifras que "
                                               "los sustentan.", ancho=8)
        filas = []
        for h in hallazgos:
            evidencia = "\n".join(
                f"{_texto(e.get('nombre'))}: {_texto(e.get('valor'))}" + (f" ({_texto(e.get('detalle'))})" if e.get("detalle") else "")
                for e in (h.get("evidence") or []) if isinstance(e, dict))
            filas.append([_TIPO_HALLAZGO.get(h.get("kind"), ("Informativo", "info"))[0], _texto(h.get("title")),
                          _texto(h.get("finding") or h.get("text")), evidencia, _texto(h.get("implication")),
                          _texto(h.get("action")), _texto(h.get("confidence") or "Media")])
        cab, _ = _tabla(ws, fila, ["Tipo", "Hallazgo", "Detalle", "Evidencia", "Implicación", "Qué hacer", "Confianza"],
                        filas, anchos={0: 12, 1: 26, 2: 50, 3: 34, 4: 42, 5: 42, 6: 11}, envolver={1, 2, 3, 4, 5})
        for i, h in enumerate(hallazgos):
            _pintar_estado(ws.cell(row=cab + 1 + i, column=2), _TIPO_HALLAZGO.get(h.get("kind"), ("", "info"))[1])


def _hoja_estadistica(wb, df, schema, dashboard) -> None:
    ws = _hoja(wb, "Estadística", "Estadística y calidad del dato",
               "Para verificar las cifras: valores típicos, dispersión y qué tan completas vienen las columnas.",
               color=GRIS, ancho=10)
    fila = 5
    stats = dashboard.get("statistics")
    if isinstance(stats, pd.DataFrame) and not stats.empty:
        fila = _seccion(ws, fila, "Columnas numéricas",
                        "La mediana es el caso típico (la mitad está por encima y la mitad por debajo); si se aleja "
                        "mucho del promedio, hay valores extremos que lo arrastran.")
        filas = []
        for _, r in stats.iterrows():
            col = r.get("Columna")
            suma = _num(numeric_valid(df[col]).sum()) if col in df.columns else None
            filas.append([_label(schema, col), _num(r.get("N")), suma, _num(r.get("Media")), _num(r.get("Mediana")),
                          _num(r.get("Std")), _num(r.get("Min")), _num(r.get("Q1")), _num(r.get("Q3")), _num(r.get("Max"))])
        fmt = {j: FMT_DECIMAL for j in range(2, 10)}
        fmt[1] = FMT_ENTERO
        _, fila = _tabla(ws, fila, ["Columna", "Con dato", "Suma", "Promedio", "Mediana", "Desviación", "Mínimo",
                                    "Percentil 25", "Percentil 75", "Máximo"], filas, fmt,
                         {0: 26, **{j: 14 for j in range(1, 10)}})
        fila += 3
    tipos = {}
    for clave, nombre in (("dates", "Fecha"), ("metrics", "Numérica"), ("categorical", "Categoría"),
                          ("ids", "Identificador"), ("geography", "Geografía"), ("text", "Texto"),
                          ("emails", "Correo")):
        for c in schema.get(clave, []) or []:
            tipos.setdefault(c, nombre)
    fila = _seccion(ws, fila, "Calidad por columna", "Vacíos = celdas sin dato. Un identificador (cédula, NIT, código) "
                                                     "no se suma aunque sea un número.")
    filas = []
    n = max(len(df), 1)
    for c in df.columns:
        vacios = int(df[c].isna().sum() + (df[c].astype(str).str.strip() == "").sum()) if len(df) else 0
        vacios = min(vacios, len(df))
        ejemplo = df[c].dropna()
        ejemplo = _texto(ejemplo.iloc[0]) if not ejemplo.empty else ""
        filas.append([_texto(c), tipos.get(c, "Otro"), vacios, vacios / n, int(df[c].nunique(dropna=True)), ejemplo[:60]])
    cab, fin = _tabla(ws, fila, ["Columna", "Tipo detectado", "Vacíos", "% vacío", "Valores distintos", "Ejemplo"],
                      filas, {2: FMT_ENTERO, 3: FMT_PCT, 4: FMT_ENTERO}, {0: 26, 1: 16, 2: 11, 3: 11, 4: 14, 5: 30})
    ws.conditional_formatting.add(f"E{cab + 1}:E{fin}",
                                  DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="F6C9CF"))


# ── Datos y tablas dinámicas ──────────────────────────────────────────────

def _encabezados_unicos(columnas) -> list[str]:
    usados, salida = set(), []
    for i, c in enumerate(columnas, 1):
        base = re.sub(r"\s+", " ", ILLEGAL_CHARACTERS_RE.sub("", str(c))).strip()[:250] or f"Columna {i}"
        nombre, k = base, 2
        while nombre.lower() in usados:
            nombre, k = f"{base} ({k})", k + 1
        usados.add(nombre.lower())
        salida.append(nombre)
    return salida


def _columna_para_celdas(serie: pd.Series) -> tuple[list, str | None]:
    """Valores listos para openpyxl y el formato numérico de la columna."""
    if pd.api.types.is_datetime64_any_dtype(serie):
        s = serie
        if getattr(s.dt, "tz", None) is not None:
            s = s.dt.tz_localize(None)
        valores = [None if pd.isna(v) else v.to_pydatetime() for v in s]
        solo_fecha = all(v is None or (v.hour == 0 and v.minute == 0 and v.second == 0) for v in valores[:2000])
        return valores, FMT_FECHA if solo_fecha else "dd/mm/yyyy hh:mm"
    if pd.api.types.is_bool_dtype(serie):
        return [None if pd.isna(v) else ("Sí" if v else "No") for v in serie], None
    if pd.api.types.is_numeric_dtype(serie):
        valores = [_num(v) for v in serie]
        return valores, _formato_valores(v for v in valores if v is not None)
    salida = []
    for v in serie:
        if v is None or (isinstance(v, float) and math.isnan(v)) or v is pd.NaT:
            salida.append(None)
        elif isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool):
            salida.append(_num(v))
        elif isinstance(v, (datetime, date, pd.Timestamp)):
            try:
                t = pd.Timestamp(v)
                salida.append((t.tz_localize(None) if t.tzinfo else t).to_pydatetime())
            except Exception:
                salida.append(_texto(v))
        else:
            salida.append(ILLEGAL_CHARACTERS_RE.sub("", str(v))[:LARGO_CELDA])
    return salida, None


def _hoja_datos(wb, df: pd.DataFrame, ctx, nombre: str = "Datos", titulo_tabla: str = "TablaDatos",
                periodo_col: str | None = None) -> dict:
    """Los datos como tabla de Excel (filtros, bandas y encabezado fijo). Si
    hay fecha se agrega «Periodo (mes)» al final: es el eje por mes de las
    tablas dinámicas y sirve para filtrar por mes sin fórmulas."""
    ws = wb.create_sheet(_nombre_hoja(wb, nombre))
    ws.sheet_properties.tabColor = TINTA
    datos = df.iloc[:MAX_FILAS_EXCEL - 1]
    encabezados = _encabezados_unicos(datos.columns)
    columnas = []
    formatos = []
    for c in datos.columns:
        valores, fmt = _columna_para_celdas(datos[c])
        columnas.append(valores)
        formatos.append(fmt)
    extra = None
    if periodo_col and periodo_col in datos.columns:
        per = a_datetime(datos[periodo_col], errors="coerce").dt.strftime("%Y-%m")
        extra = _encabezados_unicos(encabezados + ["Periodo (mes)"])[-1]
        encabezados.append(extra)
        columnas.append([None if pd.isna(v) else str(v) for v in per])
        formatos.append(None)
    ws.append(encabezados)
    for valores in zip(*columnas):
        ws.append(list(valores))
    n = len(datos)
    for j, fmt in enumerate(formatos, 1):
        if fmt:
            for (celda,) in ws.iter_rows(min_row=2, max_row=n + 1, min_col=j, max_col=j):
                celda.number_format = fmt
        else:
            # Un texto que empieza con "=" no es una fórmula: sin esto Excel
            # lo ejecutaría (o marcaría el archivo como dañado).
            for k, v in enumerate(columnas[j - 1]):
                if isinstance(v, str) and v.startswith("="):
                    ws.cell(row=k + 2, column=j).data_type = "s"
        muestra = [len(str(v)) for v in columnas[j - 1][:300] if v is not None]
        ancho = max([len(encabezados[j - 1]) + 2] + [min(m, 50) + 2 for m in muestra]) if muestra else len(encabezados[j - 1]) + 2
        ws.column_dimensions[get_column_letter(j)].width = max(9, min(ancho, 52))
    ultima_col = get_column_letter(len(encabezados))
    ref = f"A1:{ultima_col}{max(n + 1, 2)}"
    if n == 0:
        ws.append([None] * len(encabezados))
    tabla = Table(displayName=titulo_tabla, ref=ref)
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tabla)
    # Solo aquí se inmoviliza el encabezado: en las hojas de informe la línea
    # partía los gráficos y escondía filas si la hoja quedaba desplazada.
    ws.freeze_panes = "A2"
    if len(df) > n:
        ws.cell(row=1, column=len(encabezados) + 2,
                value=f"Se incluyen las primeras {n:,} filas de {len(df):,} (límite de Excel).")
    return {"ws": ws, "encabezados": encabezados, "columnas": columnas, "n": n, "ref": ref, "periodo": extra}


def _items_eje(valores: list) -> tuple[list, list]:
    """Valores distintos de un campo de la dinámica, tipados como en la celda."""
    vistos, items = set(), []
    for v in valores:
        clave = ("n", float(v)) if isinstance(v, (int, float)) else ("m", None) if v is None else ("s", str(v))
        if clave not in vistos:
            vistos.add(clave)
            items.append(clave)
    items.sort(key=lambda x: (x[0] == "m", x[0], x[1] if x[1] is not None else 0))
    shared = [Number(v=x[1]) if x[0] == "n" else Missing() if x[0] == "m" else Text(v=x[1]) for x in items]
    return items, shared


def _cache_dinamica(datos: dict, ejes: set[int]) -> CacheDefinition:
    campos = []
    for j, nombre in enumerate(datos["encabezados"]):
        valores = datos["columnas"][j]
        if j in ejes:
            items, shared = _items_eje(valores)
            si = SharedItems(_fields=shared, count=len(shared))
            if any(x[0] == "m" for x in items):
                si.containsBlank = True
            if all(x[0] == "n" for x in items if x[0] != "m"):
                si.containsSemiMixedTypes, si.containsString, si.containsNumber = False, False, True
            campos.append(CacheField(name=nombre, sharedItems=si))
        else:
            campos.append(CacheField(name=nombre, sharedItems=SharedItems()))
    hoja = datos["ws"].title
    return CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=WorksheetSource(ref=datos["ref"], sheet=hoja)),
        cacheFields=campos, refreshOnLoad=True, recordCount=0)


def _definicion_dinamica(nombre: str, cache: CacheDefinition, datos: dict, filas: list[int], columnas: list[int],
                         paginas: list[int], valores: list[tuple[int, str, str, int]], fila_inicio: int,
                         encabezado_filas: str) -> TableDefinition:
    ejes = {j: "axisRow" for j in filas}
    ejes.update({j: "axisCol" for j in columnas})
    ejes.update({j: "axisPage" for j in paginas})
    con_datos = {v[0] for v in valores}
    campos = []
    for j, _ in enumerate(datos["encabezados"]):
        if j in ejes:
            n_items = len(cache.cacheFields[j].sharedItems._fields)
            campos.append(PivotField(axis=ejes[j], showAll=False, dataField=j in con_datos or None,
                                     items=[FieldItem(x=i) for i in range(n_items)] + [FieldItem(t="default")]))
        else:
            campos.append(PivotField(dataField=j in con_datos or None, showAll=False))
    col_fields = [RowColField(x=j) for j in columnas]
    if len(valores) > 1:
        col_fields.append(RowColField(x=-2))  # «Valores» como columna cuando hay más de una medida
    td = TableDefinition(
        name=nombre, cacheId=1, dataCaption="Valores", updatedVersion=6, minRefreshableVersion=3, createdVersion=6,
        applyNumberFormats=False, applyBorderFormats=False, applyFontFormats=False, applyPatternFormats=False,
        applyAlignmentFormats=False, applyWidthHeightFormats=True, useAutoFormatting=True, itemPrintTitles=True,
        indent=0, outline=True, outlineData=True, multipleFieldFilters=False, rowHeaderCaption=encabezado_filas,
        location=Location(ref=f"B{fila_inicio}:D{fila_inicio + 3}", firstHeaderRow=1, firstDataRow=2, firstDataCol=1,
                          rowPageCount=len(paginas) or None, colPageCount=1 if paginas else None),
        pivotFields=campos, rowFields=[RowColField(x=j) for j in filas], colFields=col_fields,
        pageFields=[PageField(fld=j, hier=-1) for j in paginas],
        dataFields=[DataField(name=titulo, fld=j, subtotal=sub, baseField=0, baseItem=0, numFmtId=fmt)
                    for j, sub, titulo, fmt in valores],
        pivotTableStyleInfo=PivotTableStyle(name="PivotStyleLight16", showRowHeaders=True, showColHeaders=True,
                                            showRowStripes=False, showColStripes=False, showLastColumn=True))
    td.cache = cache
    return td


def _hojas_dinamicas(wb, datos: dict, dims: list, ctx, schema) -> list[str]:
    """Tablas dinámicas reales de Excel sobre la hoja Datos. Se recalculan al
    abrir el archivo (refreshOnLoad) y el usuario puede reorganizarlas."""
    if datos["n"] == 0:
        return []
    nombres = datos["encabezados"]
    originales = ctx["_columnas_datos"]
    pos = {c: j for j, c in enumerate(originales)}
    ejes_validos = []
    for d in dims:
        j = pos.get(d)
        if j is None:
            continue
        distintos = len({v for v in datos["columnas"][j] if v is not None})
        if 2 <= distintos <= MAX_ITEMS_DINAMICA:
            ejes_validos.append(j)
    if not ejes_validos:
        return []
    j_periodo = len(nombres) - 1 if datos.get("periodo") else None
    primary = ctx["primary"]
    j_valor = pos.get(primary) if primary else None
    if j_valor is not None:
        sub = "average" if ctx["calculo"] == "Promedio" else "sum"
        medida = (j_valor, sub, f"{'Promedio' if sub == 'average' else 'Suma'} de {nombres[j_valor]}",
                  4 if sub == "average" else 3)
    else:
        medida = (ejes_validos[0], "count", "Cantidad de registros", 3)

    ejes = set(ejes_validos[:4]) | ({j_periodo} if j_periodo is not None else set())
    cache = _cache_dinamica(datos, ejes)
    hojas = []

    # 1) Dimensión principal × mes, con las demás dimensiones como filtros.
    # La definición se arma antes que la hoja: si falla, no queda una hoja vacía.
    paginas = ejes_validos[1:4]
    fila = 6 + len(paginas)
    td = _definicion_dinamica("DinamicaPorMes", cache, datos, [ejes_validos[0]],
                              [j_periodo] if j_periodo is not None else [], paginas, [medida], fila,
                              nombres[ejes_validos[0]])
    por_mes = j_periodo is not None
    ws = _hoja(wb, "Dinámica por mes" if por_mes else "Tabla dinámica",
               f"Tabla dinámica · {nombres[ejes_validos[0]]}" + (" por mes" if por_mes else ""),
               "Tabla dinámica real de Excel: se actualiza al abrir. Usa los filtros de arriba o arrastra campos en "
               "«Campos de tabla dinámica» para reorganizarla. Fuente: hoja «" + datos["ws"].title + "».",
               color=CATEGORICA[6], ancho=8)
    ws._pivots.append(td)
    ws.column_dimensions["B"].width = 28
    for k in range(3, 30):
        ws.column_dimensions[get_column_letter(k)].width = 13
    hojas.append(ws.title)

    # 2) Detalle: principal > segunda dimensión, con la medida y el conteo.
    if len(ejes_validos) >= 2:
        medidas = [medida]
        if medida[1] != "count":
            medidas.append((ejes_validos[0], "count", "N.º de registros", 3))
        paginas2 = [j for j in ([j_periodo] if j_periodo is not None else []) + ejes_validos[2:4]][:3]
        td2 = _definicion_dinamica("DinamicaDetalle", cache, datos, ejes_validos[:2], [], paginas2, medidas,
                                   6 + len(paginas2), f"{nombres[ejes_validos[0]]} / {nombres[ejes_validos[1]]}")
        ws2 = _hoja(wb, "Dinámica detalle", f"Tabla dinámica · {nombres[ejes_validos[0]]} y {nombres[ejes_validos[1]]}",
                    "Cada grupo abierto por la segunda dimensión, con su total y cuántos registros tiene. "
                    "Clic en los signos +/− para abrir o cerrar un grupo.", color=CATEGORICA[6], ancho=6)
        ws2._pivots.append(td2)
        ws2.column_dimensions["B"].width = 32
        for k in range(3, 8):
            ws2.column_dimensions[get_column_letter(k)].width = 18
        hojas.append(ws2.title)
    return hojas


# ── Portada ───────────────────────────────────────────────────────────────

def _kpi(ws, col: int, titulo: str, valor, formato: str | None, detalle: str, color_valor: str = TINTA) -> None:
    fin = col + 2
    for f in range(7, 11):
        for c in range(col, fin + 1):
            celda = ws.cell(row=f, column=c)
            celda.fill = PatternFill("solid", fgColor=PANEL)
    ws.merge_cells(start_row=7, start_column=col, end_row=7, end_column=fin)
    ws.merge_cells(start_row=8, start_column=col, end_row=9, end_column=fin)
    ws.merge_cells(start_row=10, start_column=col, end_row=10, end_column=fin)
    t = ws.cell(row=7, column=col, value=_texto(titulo).upper())
    t.font = Font(size=8.5, bold=True, color=MUTED)
    t.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    v = ws.cell(row=8, column=col, value=valor)
    v.font = Font(size=22, bold=True, color=color_valor)
    v.alignment = Alignment(horizontal="left", vertical="center", indent=1, shrink_to_fit=True)
    if formato and isinstance(valor, (int, float)):
        v.number_format = formato
    d = ws.cell(row=10, column=col, value=_texto(detalle))
    d.font = Font(size=8.5, color=MUTED)
    d.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    for c in range(col, fin + 1):
        ws.cell(row=7, column=c).border = Border(top=Side(style="medium", color=SERIE))


def _hoja_resumen(ws, df, schema, dashboard, ctx, plan, tendencia, ranking, indice) -> None:
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = MARCA
    ws.column_dimensions["A"].width = 2
    for c in range(2, 14):
        ws.column_dimensions[get_column_letter(c)].width = 13.5
    ws.column_dimensions["N"].width = 2
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Banda de título.
    for f in (2, 3, 4):
        for c in range(2, 14):
            ws.cell(row=f, column=c).fill = PatternFill("solid", fgColor=TINTA)
    for c in range(2, 14):
        ws.cell(row=5, column=c).fill = PatternFill("solid", fgColor=MARCA)
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 4
    k = ws.cell(row=2, column=2, value="INFORME EJECUTIVO · PANEL ANALÍTICO UNIVERSAL")
    k.font = Font(size=8.5, bold=True, color="F6C9CF")
    k.alignment = Alignment(indent=1, vertical="bottom")
    t = ws.cell(row=3, column=2, value=_texto(ctx["hoja"]))
    t.font = Font(size=22, bold=True, color="FFFFFF")
    t.alignment = Alignment(indent=1, vertical="center")
    s = ws.cell(row=4, column=2, value=f"{ctx['archivo']}  ·  {ctx['periodo']}  ·  {ctx['registros']:,} registros")
    s.font = Font(size=10, color="D8DCE6")
    s.alignment = Alignment(indent=1, vertical="top")
    ws.row_dimensions[6].height = 10

    # Indicadores.
    kpis = [k for k in (dashboard.get("kpis") or []) if isinstance(k, dict)]
    principal = next((k for k in kpis if k.get("kind") == "primary"), None)
    ex = dashboard.get("executive") or {}
    ger = dashboard.get("gerencia") or {}
    if principal is not None:
        valor = _num(principal.get("raw", principal.get("value")))
        etiqueta = f"{principal.get('label', 'Total')} · {ctx['metrica']}"
        detalle = (f"{'Suma' if principal.get('label') == 'Total' else 'Promedio'} de los {ctx['registros']:,} registros")
        _kpi(ws, 2, etiqueta, valor, _formato_compacto(valor or 0), detalle)
    else:
        _kpi(ws, 2, "Registros analizados", ctx["registros"], FMT_ENTERO, ctx["alcance"])
    cambio = ex.get("change")
    if cambio is not None:
        meses = (f"{ger.get('mes_b')} vs {ger.get('mes_a')}" if ger.get("mes_a") else "Último periodo vs el anterior")
        tono = {"positive": ESTADO["bueno"][0], "negative": ESTADO["malo"][0]}.get(ex.get("status"), TINTA)
        _kpi(ws, 5, "Cambio del último mes", float(cambio) / 100, FMT_PCT_SIGNO,
             f"{meses} · {_compacto(ex.get('current_period'))} frente a {_compacto(ex.get('previous'))}", tono)
    else:
        _kpi(ws, 5, "Periodo analizado", ctx["periodo"] if ctx["desde"] is None else mes(ctx["hasta"]), None,
             "Sin dos meses comparables en el archivo")
    _kpi(ws, 8, "Registros analizados", ctx["registros"], FMT_ENTERO,
         "Con los filtros aplicados" if "sin filtros" not in ctx["alcance"].lower() else "Vista completa, sin filtros")
    lider = next((k for k in kpis if k.get("kind") == "leader"), None)
    planes = (plan or {}).get("planes") or []
    if lider is not None:
        _kpi(ws, 11, str(lider.get("label") or "Líder"), _texto(lider.get("value")), None,
             " · ".join(_texto(x) for x in (lider.get("texto"), lider.get("detalle")) if x))
    else:
        criticos = sum(1 for p in planes if p.get("estado") == "critico")
        _kpi(ws, 11, "Frentes críticos", criticos, FMT_ENTERO, f"de {len(planes)} frentes en el plan",
             ESTADO["malo"][0] if criticos else ESTADO["bueno"][0])
    ws.row_dimensions[8].height = 22
    ws.row_dimensions[9].height = 22
    ws.row_dimensions[10].height = 30

    fila = 12
    # Veredicto y por qué.
    titular = ger.get("titular") or ex.get("headline")
    if titular:
        fila = _seccion(ws, fila, "Cómo va", ancho=12)
        c = ws.cell(row=fila, column=2, value=_texto(titular))
        tono = {"positive": "bueno", "negative": "malo"}.get(ex.get("status"), "info")
        c.font = Font(size=13, bold=True, color=ESTADO[tono][0])
        c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        c.fill = PatternFill("solid", fgColor=ESTADO[tono][1])
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=13)
        ws.row_dimensions[fila].height = _alto_fila([(titular, 150)], 26)
        fila += 1
        frases = list(ger.get("frases") or [])
        if not frases and ex.get("detail"):
            frases = [ex["detail"]]
        signif = (ger.get("significancia") or {}).get("texto")
        if signif:
            frases.append(signif)
        for frase in frases[:6]:
            fila = _nota(ws, fila, f"•  {frase}", ancho=12, color=TINTA, italica=False)
        fila += 1

    # Prioridades.
    if planes:
        fila = _seccion(ws, fila, "Prioridades", "Los frentes más urgentes y su primer paso. El detalle completo, con "
                                                 "metas y control semanal, está en la hoja «Plan de acción».", ancho=12)
        for c in range(2, 14):
            celda = ws.cell(row=fila, column=c)
            celda.fill = PatternFill("solid", fgColor=TINTA)
            celda.font = Font(bold=True, color="FFFFFF", size=10)
        for col, texto in ((2, "Estado"), (3, "Frente"), (6, "Primer paso"), (12, "Valor estimado")):
            ws.cell(row=fila, column=col, value=texto)
        ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
        ws.merge_cells(start_row=fila, start_column=6, end_row=fila, end_column=11)
        ws.merge_cells(start_row=fila, start_column=12, end_row=fila, end_column=13)
        fila += 1
        for p in planes[:3]:
            estado, tono = _ESTADO_PLAN.get(p.get("estado"), (str(p.get("estado") or ""), "info"))
            _pintar_estado(ws.cell(row=fila, column=2, value=estado), tono)
            frente = ws.cell(row=fila, column=3, value=_texto(p.get("titulo")))
            frente.font = Font(bold=True, color=TINTA)
            frente.alignment = Alignment(wrap_text=True, vertical="top")
            paso = _texto((p.get("pasos") or [p.get("situacion") or ""])[0])
            pc = ws.cell(row=fila, column=6, value=paso)
            pc.alignment = Alignment(wrap_text=True, vertical="top")
            vc = ws.cell(row=fila, column=12, value=_num(p.get("impacto")))
            vc.number_format = _formato_compacto(_num(p.get("impacto")) or 0)
            vc.alignment = Alignment(horizontal="right", vertical="top")
            vc.font = Font(bold=True, color=TINTA)
            ws.merge_cells(start_row=fila, start_column=3, end_row=fila, end_column=5)
            ws.merge_cells(start_row=fila, start_column=6, end_row=fila, end_column=11)
            ws.merge_cells(start_row=fila, start_column=12, end_row=fila, end_column=13)
            for c in range(2, 14):
                ws.cell(row=fila, column=c).border = _BORDE
            ws.row_dimensions[fila].height = _alto_fila([(p.get("titulo"), 40), (paso, 80)], 20)
            fila += 1
        fila += 1

    # Gráficos principales.
    if tendencia or ranking:
        fila = _seccion(ws, fila, "Gráficos principales", "Los datos de cada gráfico están en su hoja; ahí también "
                                                          "están las tablas completas.", ancho=12)
        if tendencia:
            ht = ws.parent[tendencia["hoja"]]
            cats = Reference(ht, min_col=2, min_row=tendencia["cab"] + 1, max_row=tendencia["ultima"])
            series = [(Reference(ht, min_col=3, min_row=tendencia["cab"], max_row=tendencia["ultima"]), SERIE)]
            if tendencia.get("col_meta"):
                series.append((Reference(ht, min_col=tendencia["col_meta"], min_row=tendencia["cab"],
                                         max_row=tendencia["ultima"]), GRIS))
            ws.add_chart(_grafico_linea(ht, tendencia["titulo"], cats, series, tendencia["formato_eje"],
                                        ancho=16.5 if ranking else 30, alto=8), f"B{fila}")
        if ranking:
            hr = ws.parent[ranking["hoja"]]
            top = ranking["top"]
            cats = Reference(hr, min_col=3, min_row=ranking["primera"], max_row=ranking["primera"] + top - 1)
            col = 6 if ranking["usa_meta"] else 4
            ref = Reference(hr, min_col=col, min_row=ranking["primera"] - 1, max_row=ranking["primera"] + top - 1)
            ch = _grafico_barras(hr, ranking["titulo"], cats, [(ref, SERIE)], ranking["formato"], horizontal=True,
                                 ancho=16.5 if tendencia else 30, alto=8, etiquetas=True, min_cero=ranking["min_cero"])
            ws.add_chart(ch, f"{'H' if tendencia else 'B'}{fila}")
        fila += 17

    # Sobre este informe.
    fila = _seccion(ws, fila, "Sobre este informe", ancho=12)
    comparacion = (f"{ger.get('mes_b')} frente a {ger.get('mes_a')}" if ger.get("mes_a") else "—")
    ficha = [("Archivo", ctx["archivo"]), ("Hoja", ctx["hoja"]), ("Alcance / filtros", ctx["alcance"]),
             ("Periodo", ctx["periodo"]), ("Registros", f"{ctx['registros']:,}"),
             ("Métrica principal", f"{ctx['metrica']} ({ctx['calculo'].lower()})"),
             ("Comparación", comparacion), ("Generado", ctx["generado"])]
    for etiqueta, valor in ficha:
        e = ws.cell(row=fila, column=2, value=etiqueta)
        e.font = Font(bold=True, color=MUTED, size=9.5)
        v = ws.cell(row=fila, column=4, value=_texto(valor))
        v.font = Font(color=TINTA, size=10)
        v.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=3)
        ws.merge_cells(start_row=fila, start_column=4, end_row=fila, end_column=13)
        ws.row_dimensions[fila].height = _alto_fila([(valor, 110)], 16)
        for c in range(2, 14):
            ws.cell(row=fila, column=c).border = _BORDE
        fila += 1
    fila += 1

    # Contenido del libro.
    fila = _seccion(ws, fila, "Contenido del libro", "Clic en el nombre de una hoja para ir a ella.", ancho=12)
    for nombre, descripcion in indice:
        e = ws.cell(row=fila, column=2, value=nombre)
        e.hyperlink = f"#'{nombre}'!A1"
        e.font = Font(color=SERIE, underline="single", bold=True)
        d = ws.cell(row=fila, column=5, value=descripcion)
        d.font = Font(color=MUTED, size=9.5)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        ws.merge_cells(start_row=fila, start_column=5, end_row=fila, end_column=13)
        fila += 1


# ── Punto de entrada ──────────────────────────────────────────────────────

def _formato_largo(df: pd.DataFrame, schema: dict, filename: str, sheet: str):
    """Una tabla ancha (Enero…Diciembre como columnas, sin fecha) se pasa a
    formato largo —una fila por elemento y mes— para que la tendencia, el
    ranking y las dinámicas funcionen igual que con un registro por fecha."""
    meses = month_columns(df)
    if len(meses) < 2 or columna_fecha(df, schema):
        return None
    try:
        from core.dates import extract_year_hint
        anio = extract_year_hint(filename, sheet) or datetime.now().year
    except Exception:
        anio = datetime.now().year
    dims = [c for c in dimension_candidates(df, schema) if c in df.columns][:4]
    cols = {c: n for n, c in meses}
    largo = df[dims + list(cols)].melt(id_vars=dims, value_vars=list(cols), var_name="Mes", value_name="Valor")
    largo["Fecha"] = [pd.Timestamp(int(anio), int(cols[m]), 1) for m in largo["Mes"]]
    largo["Valor"] = numeric_valid(largo["Valor"])
    largo = largo.dropna(subset=["Valor"])[["Fecha"] + dims + ["Valor"]]
    if largo.empty:
        return None
    esquema = {
        "dates": ["Fecha"], "metrics": ["Valor"], "categorical": dims, "ids": [], "text": [], "geography": [],
        "semantic": {"columns": [{"column": "Valor", "semantic_type": "quantity", "display_name": "Valor"}]
                     + [c for c in schema.get("semantic", {}).get("columns", []) if c.get("column") in dims],
                     "metrics": ["Valor"], "dimensions": dims},
    }
    return largo.reset_index(drop=True), esquema


def build_excel_report(df: pd.DataFrame, schema: dict, dashboard: dict, filename: str, sheet: str,
                       scope_label: str = "") -> bytes:
    """El libro .xlsx completo, listo para enviar. Nunca falla por una
    sección: si una no se puede armar, se omite y las demás siguen."""
    schema = schema or {}
    dashboard = dashboard or {}
    wb = Workbook()
    portada = wb.active
    portada.title = "Resumen"

    analisis_df, analisis_schema = df, schema
    largo = None
    try:
        largo = _formato_largo(df, schema, filename, sheet)
    except Exception:
        largo = None
    if largo is not None:
        analisis_df, analisis_schema = largo

    ctx = _contexto(analisis_df, analisis_schema, dashboard if largo is None else {}, filename, sheet, scope_label)
    ctx["registros"] = len(df)

    def _seguro(fn, defecto=None):
        try:
            return fn()
        except Exception:
            return defecto

    plan = {}
    try:
        from ui.report_secciones import calcular_planes
        plan = calcular_planes(df, schema, dashboard) or {}
    except Exception:
        plan = {}

    indice: list[tuple[str, str]] = []
    info_plan = _seguro(lambda: _hoja_plan(wb, plan, ctx))
    if info_plan:
        indice.append((info_plan["hoja"], "Frentes en orden de urgencia, con pasos, cómo se mide y su valor."))
    tendencia = _seguro(lambda: _hoja_tendencia(wb, analisis_df, analisis_schema, ctx, dashboard))
    if tendencia:
        indice.append((tendencia["hoja"], "Mes a mes, con variación, acumulado y meta si existe."))
    dims_cmp = _seguro(lambda: opciones_de_comparacion(analisis_df, analisis_schema), []) or []
    ranking = None
    for dim in dims_cmp[:2]:
        info = _seguro(lambda d=dim: _hoja_por_dimension(wb, analisis_df, analisis_schema, ctx, d))
        if info:
            ranking = ranking or info
            indice.append((info["hoja"], f"Cómo va cada {info['dimension'].lower()}: posición, participación, "
                                         "meta, variación y estado."))
    dims = [d for d in (_seguro(lambda: dimension_candidates(analisis_df, analisis_schema), []) or [])
            if d in analisis_df.columns]
    antes = len(wb.worksheets)
    _seguro(lambda: _hoja_matriz(wb, analisis_df, analisis_schema, ctx, dims))
    if len(wb.worksheets) > antes:
        indice.append((wb.worksheets[-1].title, "Cruces tipo tabla dinámica con mapa de calor."))

    hojas_pivote_pos = len(wb.worksheets)
    antes = len(wb.worksheets)
    _seguro(lambda: _hoja_hallazgos(wb, dashboard, ctx))
    if len(wb.worksheets) > antes:
        indice.append((wb.worksheets[-1].title, "Alertas y hallazgos del análisis, con evidencia."))
    antes = len(wb.worksheets)
    _seguro(lambda: _hoja_estadistica(wb, df, schema, dashboard))
    if len(wb.worksheets) > antes:
        indice.append((wb.worksheets[-1].title, "Estadística descriptiva y calidad de cada columna."))

    # Datos al final; las dinámicas se leen de aquí.
    datos = _hoja_datos(wb, df, ctx, periodo_col=ctx["fecha_col"] if largo is None else None)
    indice.append((datos["ws"].title, f"Los {min(len(df), MAX_FILAS_EXCEL - 1):,} registros como tabla de Excel "
                                      "con filtros."))
    fuente_pivote = datos
    ctx["_columnas_datos"] = list(df.columns)
    if largo is not None:
        fuente_pivote = _hoja_datos(wb, analisis_df, ctx, nombre="Datos por mes", titulo_tabla="TablaDatosMes",
                                    periodo_col="Fecha")
        ctx["_columnas_datos"] = list(analisis_df.columns)
        indice.append((fuente_pivote["ws"].title, "Los mismos datos, una fila por elemento y mes."))
    hojas_pivote = _seguro(lambda: _hojas_dinamicas(wb, fuente_pivote, dims, ctx, analisis_schema), []) or []
    # Las dinámicas van después de la matriz, antes del soporte.
    for k, nombre in enumerate(hojas_pivote):
        hoja = wb[nombre]
        wb.move_sheet(hoja, offset=(hojas_pivote_pos + k) - wb.worksheets.index(hoja))
    indice += [(nombre, "Tabla dinámica real de Excel: filtra y reorganiza arrastrando campos.")
               for nombre in hojas_pivote]
    indice.sort(key=lambda x: wb.worksheets.index(wb[x[0]]))

    try:
        _hoja_resumen(portada, df, schema, dashboard, ctx, plan, tendencia, ranking, indice)
    except Exception:
        portada.cell(row=2, column=2, value=f"Informe · {ctx['hoja']}").font = Font(size=16, bold=True)
    wb.active = 0

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def nombre_archivo(sheet: str) -> str:
    base = re.sub(r"[^A-Za-z0-9_-]+", "_", str(sheet)).strip("_") or "datos"
    return f"informe_{base}_{datetime.now():%Y-%m-%d}.xlsx"
